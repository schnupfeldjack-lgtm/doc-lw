# -*- coding: utf-8 -*-
"""最终核验：模板明文规定 vs 三篇论文实测值"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'
FILES = [
    ('论文一', '再生混凝土在装配式建筑中的应用评价——以住宅项目为例'),
    ('论文二', '装配式施工质量管理问题及优化研究——以市政项目为例'),
    ('论文三', 'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例'),
]
TW = 12700.0
NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def ea_of(r):
    try:
        rpr = r._element.rPr
        if rpr is not None and rpr.rFonts is not None:
            return rpr.rFonts.get(NS + 'eastAsia')
    except Exception:
        pass
    return None


def inf(para):
    r = para.runs[0] if para.runs else None
    pf = para.paragraph_format
    sz = r.font.size.pt if (r and r.font.size) else None
    ls = pf.line_spacing
    if isinstance(ls, (int, float)) and ls and ls > 10000:
        ls = round(ls / TW, 1)
    return dict(sz=sz, ea=ea_of(r) if r else None,
                b=(r.font.bold if r else None),
                al=para.alignment, ls=ls,
                fi=round(pf.first_line_indent / TW, 1) if pf.first_line_indent else None,
                li=round(pf.left_indent / TW, 1) if pf.left_indent else None)


def pick(paras, pred):
    return [p for p in paras if pred(p)]


def body(p):
    t = p.text.strip()
    return '\t' not in p.text and not re.match(r'^\[\d+\]', t)


CHECKS = [
    ('摘要标签', lambda p: p.text.strip().startswith('摘 要：'),
     lambda i: i['sz'] == 12 and i['ea'] == '宋体' and i['b'] is True,
     '小四号宋体、加粗'),
    ('关键词标签', lambda p: p.text.strip().startswith('关键词：'),
     lambda i: i['sz'] == 12 and i['b'] is True, '小四号宋体、加粗'),
    ('Abstract标签', lambda p: p.text.strip().startswith('Abstract'),
     lambda i: str(i['ea']) == 'Times New Roman' and i['b'] is True,
     '小四号Times New Roman、加粗'),
    ('Key words标签', lambda p: p.text.strip().startswith('Key words'),
     lambda i: str(i['ea']) == 'Times New Roman' and i['b'] is True,
     '小四号Times New Roman、加粗'),
    ('目录标题', lambda p: p.text.strip().startswith('目') and '录' in p.text[:4],
     lambda i: i['sz'] == 18 and i['ea'] == '黑体' and i['b'] is True
     and i['al'] == WD_ALIGN_PARAGRAPH.CENTER, '黑体小二号加粗居中'),
    ('目录条目', lambda p: '\t' in p.text and re.match(r'^\d', p.text.strip()),
     lambda i: i['sz'] == 12 and i['ea'] == '宋体' and i['ls'] == 28.0,
     '宋体小四号，行距固定值28磅'),
    ('一级标题', lambda p: body(p) and re.match(r'^\d+\s+\S', p.text.strip())
     and not re.match(r'^\d+[.．]\d', p.text.strip()) and len(p.text) < 40,
     lambda i: i['sz'] == 15 and i['ea'] == '黑体' and i['b'] is True,
     '小3号黑体、加粗，段前0.5行段后0.5行'),
    ('二级标题', lambda p: body(p) and re.match(r'^\d+[.．]\d+\s+\S', p.text.strip())
     and len(p.text) < 40,
     lambda i: i['sz'] == 14 and i['ea'] == '黑体' and i['b'] is True,
     '4号黑体、加粗'),
    ('三级标题', lambda p: body(p) and re.match(r'^\d+[.．]\d+[.．]\d+\s+\S', p.text.strip())
     and len(p.text) < 40,
     lambda i: i['sz'] == 12 and i['ea'] == '黑体' and i['b'] is not True,
     '小4号黑体、不加粗'),
    ('正文段落', lambda p: body(p) and len(p.text) > 120
     and not p.text.strip().startswith(
         ('摘 要', '关键词', 'Abstract', 'Key words', '本文以', '本论文是在')),
     lambda i: i['sz'] == 12 and i['ls'] == 1.5 and i['fi'] == 24.0,
     '小4号宋体、1.5倍行距、首行缩进2字符'),
    ('图表题', lambda p: re.match(r'^(图|表)\s*\d+[–\-—]\d+\s+\S', p.text.strip())
     and len(p.text) < 45,
     lambda i: i['sz'] == 10.5 and i['ea'] == '宋体',
     '5号宋体'),
    ('结论标题', lambda p: p.text.strip() == '结  论',
     lambda i: i['sz'] == 14 and i['ea'] == '宋体' and i['b'] is True
     and i['al'] == WD_ALIGN_PARAGRAPH.CENTER, '四号宋体、加粗、居中'),
    ('结论正文', lambda p: body(p) and len(p.text) > 100 and '本文以' in p.text[:30],
     lambda i: i['sz'] == 12 and i['ls'] == 22.0, '小四号宋体，固定值22磅'),
    ('参考文献标题', lambda p: p.text.strip() == '参 考 文 献',
     lambda i: i['sz'] == 14 and i['ea'] == '宋体' and i['b'] is True
     and i['al'] == WD_ALIGN_PARAGRAPH.CENTER, '四号宋体、加粗、居中'),
    ('文献条目', lambda p: re.match(r'^\[\d+\]', p.text.strip()),
     lambda i: i['sz'] == 10.5 and i['ls'] == 22.0, '五号宋体，固定值22磅'),
    ('致谢标题', lambda p: p.text.strip() == '致  谢',
     lambda i: i['sz'] == 14 and i['ea'] == '宋体' and i['b'] is True
     and i['al'] in (None, WD_ALIGN_PARAGRAPH.LEFT), '四号宋体、加粗、左对齐'),
]

for tag, nm in FILES:
    d = Document(f'{W}\\{nm}\\{nm}.docx')
    paras = d.paragraphs
    print('=' * 74)
    print(tag, nm[:22])
    print('=' * 74)
    ng = 0
    for name, pred, chk, spec in CHECKS:
        hits = pick(paras, pred)
        if not hits:
            print(f'  ?? {name:<8} 未找到样本      规定：{spec}')
            ng += 1
            continue
        oks = [chk(inf(h)) for h in hits]
        i0 = inf(hits[0])
        if all(oks):
            print(f'  OK {name:<8} 共{len(hits):>2}处  sz={i0["sz"]} {i0["ea"]} '
                  f'b={i0["b"]} ls={i0["ls"]}  [{spec}]')
        else:
            ng += 1
            print(f'  NG {name:<8} 共{len(hits):>2}处  sz={i0["sz"]} {i0["ea"]} '
                  f'b={i0["b"]} al={i0["al"]} ls={i0["ls"]} fi={i0["fi"]}  [{spec}]')
    # 角标上标
    tot = sup = 0
    for p in paras:
        if re.match(r'^\[\d+\]', p.text.strip()):
            continue
        for m in re.finditer(r'\[\d+\]', p.text):
            tot += 1
            pos = 0
            for r in p.runs:
                if pos <= m.start() < pos + len(r.text):
                    rpr = r._element.rPr
                    v = rpr.find(qn('w:vertAlign')) if rpr is not None else None
                    if v is not None and v.get(qn('w:val')) == 'superscript':
                        sup += 1
                    break
                pos += len(r.text)
    print(f'  {"OK" if tot == sup else "NG"} 角标上标   {sup}/{tot} 个为右上角上标  '
          f'[模板：句子右上角方括弧角注]')
    if tot != sup:
        ng += 1
    # 彩色字体
    col = 0
    for p in paras:
        for r in p.runs:
            if r.font.color and r.font.color.rgb and str(r.font.color.rgb) != '000000':
                c = str(r.font.color.rgb)
                if max(int(c[i:i + 2], 16) for i in (0, 2, 4)) - min(
                        int(c[i:i + 2], 16) for i in (0, 2, 4)) > 40:
                    col += 1
    print(f'  {"OK" if col == 0 else "NG"} 字体颜色   彩色残留 {col} 处  '
          f'[模板：全文字体统一设置成黑色]')
    if col:
        ng += 1
    # 表格填写内容字号（模板 T0/T1/T2 示例 = 小四 12pt）
    bad_sz = []
    for ti in (0, 1, 2):
        for row in d.tables[ti].rows:
            for c in row.cells:
                for para in c.paragraphs:
                    for r in para.runs:
                        if r.text.strip() and r.font.size and abs(r.font.size.pt - 12.0) > 0.01:
                            bad_sz.append((ti, round(r.font.size.pt, 1), r.text[:16]))
    print(f'  {"OK" if not bad_sz else "NG"} 前期表格字号 共检查3张表，非12pt {len(bad_sz)} 处  [模板：小四号]')
    if bad_sz:
        ng += 1
        print(f'       {bad_sz[:3]}')
    # 封面年月数字（三号楷体_GB2312）
    bad_cov = []
    for row in d.tables[4].rows:
        for ci, c in enumerate(row.cells):
            if ci not in (0, 2):
                continue
            for para in c.paragraphs:
                for r in para.runs:
                    if not r.text.strip():
                        continue
                    sz = r.font.size.pt if r.font.size else None
                    ea = ea_of(r) if r is not None else None
                    if sz != 16 or ea != '楷体_GB2312':
                        bad_cov.append((ci, sz, ea, r.text[:8]))
    print(f'  {"OK" if not bad_cov else "NG"} 封面年月填写  {len(bad_cov)} 处不符  [模板重要提示：三号楷体_GB2312]')
    if bad_cov:
        ng += 1
        print(f'       {bad_cov[:3]}')
    # 关键词分隔符
    for p in paras:
        t = p.text.strip()
        if t.startswith('关键词：'):
            okk = ('；' not in t) and ('，' in t)
            desc = '用逗号分隔' if okk else '含分号'
            print(f'  {"OK" if okk else "NG"} 中文关键词分隔符  {desc}  [模板明文：×××，×××，×××]')
            if not okk:
                ng += 1
        if t.startswith('Key words'):
            okk = (';' not in t) and (', ' in t)
            desc = '用逗号+空格' if okk else '含分号'
            print(f'  {"OK" if okk else "NG"} 英文关键词分隔符  {desc}  [模板明文：×××, ×××, ×××]')
            if not okk:
                ng += 1
    # 摘要行距
    for p in paras:
        if p.text.strip().startswith('摘 要：'):
            pPr = p._element.find(qn('w:pPr'))
            s = pPr.find(qn('w:spacing')) if pPr is not None else None
            lv = s.get(qn('w:line')) if s is not None else None
            lr = s.get(qn('w:lineRule')) if s is not None else None
            ok = (lv == '560' and lr == 'exact')
            print(f'  {"OK" if ok else "NG"} 摘要行距      line={lv} rule={lr}  [模板母段落：560/exact=固定28磅]')
            if not ok:
                ng += 1
    # 致谢正文首行缩进
    for p in paras:
        if p.text.strip().startswith('本论文是在'):
            pPr = p._element.find(qn('w:pPr'))
            ind = pPr.find(qn('w:ind')) if pPr is not None else None
            fl = ind.get(qn('w:firstLine')) if ind is not None else None
            ok = (fl == '573')
            print(f'  {"OK" if ok else "NG"} 致谢正文缩进    firstLine={fl}  [模板母段落：573=28.65磅]')
            if not ok:
                ng += 1
    print(f'  --> 本篇不符项：{ng}')
    print()
