# -*- coding: utf-8 -*-
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
tpl = docx.Document(os.path.join(BASE,"炎黄职业技术学院毕业论文模板(1).docx"))
sub = '再生混凝土在装配式建筑中的应用评价——以住宅项目为例'
d = docx.Document(os.path.join(BASE, sub, sub+".docx"))
print("###### 模板 段30~58（节0末/节1/节2头）######")
for i in range(30, 58):
    t = tpl.paragraphs[i].text.strip()
    if t: print(f"  TPL[{i}] {t[:60]}")
print("\n###### 论文 段33~48（节0末/节1头）######")
for i in range(33, 48):
    t = d.paragraphs[i].text.strip()
    print(f"  [{i}] {t[:60]!r}")
print("\n###### 论文段 68~74（节1末/节2头）######")
for i in range(68, 75):
    t = d.paragraphs[i].text.strip()
    print(f"  [{i}] {t[:60]!r}")
