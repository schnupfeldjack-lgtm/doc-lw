# -*- coding: utf-8 -*-
"""查三级标题实际用的点号"""
import re
from docx import Document
BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
for path, nm in [
    (BASE + r"\再生混凝土在装配式建筑中的应用评价——以住宅项目为例\再生混凝土在装配式建筑中的应用评价——以住宅项目为例.docx", "论文一"),
    (BASE + r"\装配式施工质量管理问题及优化研究——以市政项目为例\装配式施工质量管理问题及优化研究——以市政项目为例.docx", "论文二"),
    (BASE + r"\BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例\BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例.docx", "论文三"),
]:
    d = Document(path)
    h3 = [p.text for p in d.paragraphs if re.match(r'^\d+\W\d+\W\d+', p.text.strip()) and '\t' not in p.text]
    print(f"{nm} 三级标题示例: {h3[:3]}")
    # TOC 一级/二级
    toc1 = [p.text for p in d.paragraphs if re.match(r'^\d+\s', p.text.strip()) and '\t' in p.text]
    toc2 = [p.text for p in d.paragraphs if re.match(r'^\d+\W\d+', p.text.strip()) and '\t' in p.text]
    print(f"  TOC一级示例: {toc1[:3]}")
    print(f"  TOC二级示例: {toc2[:3]}")