# -*- coding: utf-8 -*-
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn
BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']

def runinfo(r):
    rPr = r._element.find(qn('w:rPr'))
    if rPr is None: return {}
    d={}
    e=rPr.find(qn('w:sz'))
    if e is not None: d['sz']=int(e.get(qn('w:val')))/2
    e=rPr.find(qn('w:rFonts'))
    if e is not None: d['ea']=e.get(qn('w:eastAsia'))
    return d

for sub in FILES:
    d = docx.Document(os.path.join(BASE, sub, sub+".docx"))
    print(f"\n===== {sub[:20]} =====")
    # 封面表3 (5行x4列) 的填写内容
    t = d.tables[3]
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            tx = c.text.strip()
            if not tx: continue
            ri_ = runinfo(c.paragraphs[0].runs[0]) if c.paragraphs[0].runs else {}
            print(f"   表3 行{ri}列{ci}: {tx[:34]!r}  {ri_}")
