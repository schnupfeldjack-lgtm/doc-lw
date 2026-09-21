# -*- coding: utf-8 -*-
"""直接改页眉域的结果缓存（不用 Word 保存，避免 Word 重写丢 eastAsia/sz）
   流程：Word 只读取页数 -> python-docx 改 <w:t> 缓存 -> python-docx 保存"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import win32com.client as wc
import docx
from docx.oxml.ns import qn

W = r'C:/Users/15515/Desktop/09-文档资料/新建文件夹'
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']

# 1) 只读取页数
app = wc.gencache.EnsureDispatch('Word.Application'); app.Visible=False; app.DisplayAlerts=0
pages = {}
for nm in FILES:
    doc = app.Documents.Open(os.path.join(W,nm,nm+'.docx'), ReadOnly=True)
    doc.Repaginate(); pages[nm] = doc.ComputeStatistics(2); doc.Close(False)
app.Quit()
print("页数:", {k[:8]: v for k, v in pages.items()})

# 2) 改域缓存
def fix_cache(p, total):
    n = 0
    for para in p.paragraphs:
        runs = para._element.findall(qn('w:r'))
        instr = ''; mode = None
        for r in runs:
            fc = r.find(qn('w:fldChar'))
            if fc is not None:
                t = fc.get(qn('w:fldCharType'))
                if t == 'begin': mode = 'begin'; instr = ''
                elif t == 'separate': mode = 'sep'
                elif t == 'end': mode = None
                continue
            it = r.find(qn('w:instrText'))
            if it is not None:
                instr += (it.text or '')
                continue
            ts = r.findall(qn('w:t'))
            if mode == 'sep' and ts:
                if 'NUMPAGES' in instr.upper():
                    ts[0].text = str(total); n += 1
                elif 'PAGE' in instr.upper():
                    ts[0].text = '1'; n += 1
    return n

for nm in FILES:
    path = os.path.join(W, nm, nm+'.docx')
    d = docx.Document(path)
    tot = 0
    for s in d.sections:
        for h in (s.header, s.footer):
            if h is None: continue
            tot += fix_cache(h, pages[nm])
    d.save(path)
    d2 = docx.Document(path)
    txt = [p.text for s in d2.sections for p in s.header.paragraphs if p.text.strip()]
    print(f"{nm[:18]} 改缓存 {tot} 处  页眉={txt[:1]}")
