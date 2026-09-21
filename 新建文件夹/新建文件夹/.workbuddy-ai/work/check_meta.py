# -*- coding: utf-8 -*-
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from lxml import etree
W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'
FILES = [
    '再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
    '装配式施工质量管理问题及优化研究——以市政项目为例',
    'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例',
]
NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

def hdr_text(hdr):
    out = []
    for p in hdr.paragraphs:
        # 展开域
        xml = etree.tostring(p._element, encoding='unicode')
        codes = re.findall(r'<w:instrText[^>]*>(.*?)</w:instrText>', xml)
        out.append((p.text, codes))
    return out

for nm in FILES:
    d = Document(f'{W}\\{nm}\\{nm}.docx')
    print('=' * 70)
    print(nm[:24])
    print(f'  表格数={len(d.tables)}  内联图={len(d.inline_shapes)}  节数={len(d.sections)}')
    for si, s in enumerate(d.sections):
        print(f'  --- sec{si}: L={s.left_margin} R={s.right_margin} T={s.top_margin} B={s.bottom_margin}')
        for t, codes in hdr_text(s.header):
            print(f'      HDR {t!r} 域={codes}')
        for t, codes in hdr_text(s.footer):
            print(f'      FTR {t!r} 域={codes}')
    for ti, tb in enumerate(d.tables):
        first = ' | '.join(c.text.strip()[:14] for c in tb.rows[0].cells)
        print(f'  TBL{ti} rows={len(tb.rows)} cols={len(tb.columns)} :: {first}')
    # 真实图题/表题（独立段，短，居中）
    for p in d.paragraphs:
        t = p.text.strip()
        if re.match(r'^(图|表)\s*\d+[–\-—]\s*\d+', t):
            r = p.runs[0] if p.runs else None
            ea = None
            if r is not None:
                try:
                    rpr = r._element.rPr
                    if rpr is not None and rpr.rFonts is not None:
                        ea = rpr.rFonts.get(NS + 'eastAsia')
                except Exception:
                    pass
            print(f'   {"图/表题"} sz={r.font.size.pt if r and r.font.size else None} ea={ea} al={p.alignment} ls={p.paragraph_format.line_spacing} >> {t[:50]}')
    print()
