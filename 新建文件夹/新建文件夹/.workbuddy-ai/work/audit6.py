# -*- coding: utf-8 -*-
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn
BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']

def tbl_info(t, i):
    rows = len(t.rows); cols = len(t.columns)
    al = t.alignment
    st = t.style.name if t.style else None
    # 单元格字体
    szs = set()
    for r in t.rows[:3]:
        for c in r.cells[:3]:
            for p in c.paragraphs:
                for run in p.runs:
                    rPr = run._element.find(qn('w:rPr'))
                    if rPr is not None:
                        e = rPr.find(qn('w:sz'))
                        if e is not None: szs.add(e.get(qn('w:val')))
    return f"  表{i}: {rows}行x{cols}列 对齐={al} 样式={st} 字号={sorted(szs)}"

tpl = docx.Document(os.path.join(BASE,'炎黄职业技术学院毕业论文模板(1).docx'))
print("###### 模板表格 ######")
for i,t in enumerate(tpl.tables): print(tbl_info(t,i))
for sub in FILES:
    d = docx.Document(os.path.join(BASE, sub, sub+".docx"))
    print(f"\n###### {sub[:20]} ######")
    for i,t in enumerate(d.tables): print(tbl_info(t,i))
    # 图片
    imgs = d.element.body.findall('.//'+qn('a:blip'))
    print(f"  图片数={len(imgs)}")
    # 打印删除残留
    hits=[]
    for p in d.paragraphs:
        if '打印删除' in p.text or '（打印' in p.text: hits.append(p.text[:40])
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                if '打印删除' in c.text: hits.append(c.text[:40])
    print(f"  '打印删除'残留={len(hits)} {hits[:3]}")
