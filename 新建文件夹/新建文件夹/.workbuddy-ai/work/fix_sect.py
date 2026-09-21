# -*- coding: utf-8 -*-
"""修复：论文缺少模板的"封面节"（左3.00 右3.00 装订0.20）
   在封面最后一段(段23)末插入模板节1 的 sectPr，使其后的摘要节沿用 2.50/1.80/0.90"""
import sys, io, os, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn

BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']

tpl = docx.Document(os.path.join(BASE,'炎黄职业技术学院毕业论文模板(1).docx'))
tpl_sect = tpl.sections[0]._sectPr          # 3.00/3.00/0.20
print("模板节1 sectPr 片段:")
for ch in tpl_sect.iterchildren():
    if ch.tag.split('}')[-1] in ('pgSz','pgMar','type','cols','docGrid'):
        print("   ", ch.tag.split('}')[-1], dict(ch.attrib))

for sub in FILES:
    p_ = os.path.join(BASE, sub, sub+".docx")
    d = docx.Document(p_); ps = d.paragraphs
    # 定位：摘 要 段的前一个非空段（封面末行）
    ai = [i for i,p in enumerate(ps) if p.text.strip().startswith('摘 要：')][-1]
    j = ai - 1
    while j >= 0 and not ps[j].text.strip():
        j -= 1
    print(f"\n{sub[:20]}: 摘要段={ai} 封面末段={j} {ps[j].text[:30]!r}")
    pPr = ps[j]._element.get_or_add_pPr()
    if pPr.find(qn('w:sectPr')) is not None:
        print("   已有 sectPr，跳过"); continue
    new = copy.deepcopy(tpl_sect)
    # 清掉页眉页脚引用（论文页眉在正文节，封面节不需要）
    for tag in ('w:headerReference','w:footerReference'):
        for e in new.findall(qn(tag)):
            new.remove(e)
    pPr.append(new)
    d.save(p_)
    print("   已插入封面节 sectPr")
