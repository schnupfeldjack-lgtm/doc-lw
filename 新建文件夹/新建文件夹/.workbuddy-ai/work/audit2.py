# -*- coding: utf-8 -*-
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn
BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
tpl = docx.Document(os.path.join(BASE,"炎黄职业技术学院毕业论文模板(1).docx"))

def firstLine(p):
    pPr = p._element.find(qn('w:pPr'))
    if pPr is None: return None
    ind = pPr.find(qn('w:ind'))
    if ind is None: return None
    return (ind.get(qn('w:firstLine')), ind.get(qn('w:firstLineChars')),
            ind.get(qn('w:left')), ind.get(qn('w:hanging')))

print("###### 模板正文示例段落的首行缩进分布 ######")
from collections import Counter
c = Counter()
for i,p in enumerate(tpl.paragraphs):
    t=p.text.strip()
    if ('小4号宋体' in t) or ('小四号宋体' in t):
        fl = firstLine(p)
        c[fl]+=1
        print(f"  TPL[{i}] firstLine/left/hanging={fl}  {t[:40]!r}")
print("\n分布:", c)

print("\n###### 模板目录：域还是静态文本 ######")
for i,p in enumerate(tpl.paragraphs):
    if '目' in p.text and '录' in p.text[:6] or '\t' in p.text:
        xml = p._element.xml
        has_fld = 'fldChar' in xml or 'instrText' in xml or 'TOC' in xml
        ind = firstLine(p)
        print(f"  TPL[{i}] {p.text[:44]!r} 是域={has_fld} 缩进={ind}")
        if i>62: break
