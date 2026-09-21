# -*- coding: utf-8 -*-
"""精确审计：模板母段落 vs 论文对应段落，逐属性 diff（用 run 级 rPr，不用段落标记 rPr）"""
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn

BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
TPL  = os.path.join(BASE, "炎黄职业技术学院毕业论文模板(1).docx")

def props(p):
    """返回 (段落级属性dict, 首个run属性dict) —— run 级才是可见格式"""
    P, R = {}, {}
    pPr = p._element.find(qn('w:pPr'))
    if pPr is not None:
        sp = pPr.find(qn('w:spacing'))
        if sp is not None:
            for k, v in sp.attrib.items():
                P[k.split('}')[1]] = v
        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            for k, v in ind.attrib.items():
                P[k.split('}')[1]] = v
        for tag in ('w:jc','w:pageBreakBefore','w:keepNext','w:keepLines','w:outlineLvl','w:pStyle'):
            e = pPr.find(qn(tag))
            if e is not None:
                av = e.get(qn('w:val'))
                P[tag[2:]] = av if av is not None else 'Y'
    if p.runs:
        rPr = p.runs[0]._element.find(qn('w:rPr'))
        if rPr is not None:
            for tag in ('w:sz','w:szCs','w:b','w:i','w:vertAlign','w:color'):
                e = rPr.find(qn(tag))
                if e is not None:
                    av = e.get(qn('w:val'))
                    R[tag[2:]] = av if av is not None else 'Y'
            e = rPr.find(qn('w:rFonts'))
            if e is not None:
                for k, v in e.attrib.items():
                    R['font:'+k.split('}')[1]] = v
    return P, R

def show(name, p):
    P, R = props(p)
    print(f"  {name:<12} 段属性={P}")
    print(f"  {'':<12} run属性={R}")

tpl = docx.Document(TPL)
# 模板各类母段落
CATS = {
 '正文':      [i for i,p in enumerate(tpl.paragraphs) if '小4号宋体）' in p.text or '小4号宋体)' in p.text],
 '一级标题':  [i for i,p in enumerate(tpl.paragraphs) if '小3号黑体' in p.text and '加粗' in p.text],
 '二级标题':  [i for i,p in enumerate(tpl.paragraphs) if '4号黑体' in p.text],
 '三级标题':  [i for i,p in enumerate(tpl.paragraphs) if '小4号黑体' in p.text],
 '图表题':    [i for i,p in enumerate(tpl.paragraphs) if re.match(r'^\s*图\s*\d+', p.text) or re.match(r'^\s*表\s*\d+', p.text)],
 '结论正文':  [i for i,p in enumerate(tpl.paragraphs) if '小四号宋体，行距为固定值22磅' in p.text],
 '文献条目':  [i for i,p in enumerate(tpl.paragraphs) if re.match(r'^\[\d+\]', p.text.strip())],
 '目录条目':  [i for i,p in enumerate(tpl.paragraphs) if '\t' in p.text],
}
print("########## 模板母段落 ##########")
for c, idxs in CATS.items():
    print(f"\n-- {c}  模板段{idxs[:4]}")
    for i in idxs[:2]:
        show(f"TPL[{i}]", tpl.paragraphs[i])
