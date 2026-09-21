# -*- coding: utf-8 -*-
"""
1) 删除模板说明文字残留段落（注：/1．正文中公式…/2．正文各页…/3．为保证打印效果…）
2) 「结  论」段内的分页符移到其前一段，使结论从新页开始（与模板一致）
3) 「参 考 文 献」标题前补分页符（模板在结论页与文献页之间也换页）
"""
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
def q(t): return f"{{{W}}}{t}"

BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
DOCS = [
    (BASE + r"\再生混凝土在装配式建筑中的应用评价——以住宅项目为例\再生混凝土在装配式建筑中的应用评价——以住宅项目为例.docx", "论文一"),
    (BASE + r"\装配式施工质量管理问题及优化研究——以市政项目为例\装配式施工质量管理问题及优化研究——以市政项目为例.docx", "论文二"),
    (BASE + r"\BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例\BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例.docx", "论文三"),
]

NOTE_PAT = [r'^注：\s*$', r'^1．正文中公式', r'^2．正文各页', r'^3．为保证打印效果',
            r'^（本页为独立页', r'^（空\d行）\s*$', r'^说明[:：]', r'^例如[:：]']

def is_note(t):
    s = t.strip()
    return any(re.match(p, s) for p in NOTE_PAT)

def add_pagebreak_before(par):
    """在 par 之前插入一个只含分页符的空段落（沿用 par 的 pPr）"""
    pPr = par._p.find(q('pPr'))
    newp = OxmlElement('w:p')
    if pPr is not None:
        newp.append(pPr.__class__(pPr)) if False else newp.append(__import__('copy').deepcopy(pPr))
    r = OxmlElement('w:r')
    br = OxmlElement('w:br'); br.set(q('type'), 'page')
    r.append(br); newp.append(r)
    par._p.addprevious(newp)

def ensure_pagebreak_before(doc, idx):
    """确保 paragraphs[idx] 之前有分页符：优先放进前一段（若为空段），否则插入新空段"""
    tgt = doc.paragraphs[idx]._p
    prev = tgt.getprevious()
    if prev is not None and prev.tag == q('p'):
        # 前一段若为空（无 w:t 文本）则直接加分页符 run
        has_text = any((t.text or '').strip() for t in prev.iter(q('t')))
        if not has_text:
            r = OxmlElement('w:r')
            br = OxmlElement('w:br'); br.set(q('type'), 'page')
            r.append(br); prev.append(r)
            return '放入前一段'
    add_pagebreak_before(doc.paragraphs[idx])
    return '插入新空段'

for path, nm in DOCS:
    d = Document(path)
    # 1) 删除说明文字残留
    removed = []
    for p in list(d.paragraphs):
        if is_note(p.text):
            removed.append(p.text[:24])
            p._p.getparent().remove(p._p)
    # 2) 结论：分页符从标题段内移到标题之前
    ci = None
    for i, p in enumerate(d.paragraphs):
        if re.match(r'^结\s*论$', p.text.strip()):
            ci = i; break
    msg2 = '未找到'
    if ci is not None:
        par = d.paragraphs[ci]
        for r in list(par.runs):
            if r._element.findall(f".//{q('br')}"):
                r._element.getparent().remove(r._element)
        msg2 = ensure_pagebreak_before(d, ci)
    # 3) 参考文献前补分页符
    ri = None
    for i, p in enumerate(d.paragraphs):
        if re.match(r'^参\s*考\s*文\s*献$', p.text.strip()):
            ri = i; break
    msg3 = '未找到'
    if ri is not None:
        msg3 = ensure_pagebreak_before(d, ri)
    d.save(path)
    print(f"{nm}: 删除模板说明段 {len(removed)} 段 {removed}；结论分页符={msg2}；参考文献前分页符={msg3}")
print("完成")
