# -*- coding: utf-8 -*-
"""清理目录条目中重复的页码 run（retoc 误插造成）"""
import docx, os, re
BASE = r"C:/Users/15515/Desktop/09-文档资料/新建文件夹"
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']
for sub in FILES:
    p_ = os.path.join(BASE, sub, sub+".docx")
    d = docx.Document(p_); n=0
    for p in d.paragraphs:
        if '\t' not in p.text: continue
        ti = None
        for k,r in enumerate(p.runs):
            if '\t' in r.text: ti = k
        if ti is None: continue
        extra = p.runs[ti+1:]
        if not extra: continue
        for r in extra:
            r._element.getparent().remove(r._element)
        n += 1
        print(f"  清理 {p.text!r} -> ", end="")
    if n: d.save(p_)
    print(f"{sub[:22]} 清理 {n} 条；结果:")
    d2 = docx.Document(p_)
    for p in d2.paragraphs:
        if '\t' in p.text and re.sub(r'[\s.]','',p.text.split('\t')[0].strip()) in ('结论','致谢','参考文献'):
            print("    ", repr(p.text))
