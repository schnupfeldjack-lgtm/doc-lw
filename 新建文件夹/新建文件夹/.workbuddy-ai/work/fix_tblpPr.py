# -*- coding: utf-8 -*-
"""删除任务书表(TBL1)和中期检查表(TBL2)的浮动属性（tblpPr）
   原因：模板是浮动定位（Y=3157/2689 twips），但论文实际内容长，导致段落被挤到页底。
        改为按文档流渲染（删除 tblpPr）。"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn

BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']

for sub in FILES:
    p_ = os.path.join(BASE, sub, sub+".docx")
    d = docx.Document(p_)
    n = 0
    # TBL1 和 TBL2 是浮动表
    for ti in (1, 2):
        if ti >= len(d.tables): continue
        tbl = d.tables[ti]
        tblPr = tbl._element.find(qn('w:tblPr'))
        if tblPr is None: continue
        tblpPr = tblPr.find(qn('w:tblpPr'))
        if tblpPr is None: continue
        tblPr.remove(tblpPr); print(f"  {sub[:20]} TBL{ti} 移除 tblpPr")
        n += 1
    if n:
        d.save(p_)
        print(f"  {sub[:20]} 已保存 (改动 {n} 处)")
