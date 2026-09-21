# -*- coding: utf-8 -*-
"""替换问题参考文献条目（保持编号、格式不变）。所有替换条目均已网络核实。"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'

REPLACE = {
    '装配式施工质量管理问题及优化研究——以市政项目为例': {
        4:  '中建二局绵阳综合管廊项目主体结构全面完工[N/OL]. 腾讯新闻, 2026-02-25.',
        6:  '王伟. 城市综合管廊预制拼装施工技术[J]. 基层建设, 2019(3).',
        7:  '崔宇, 柴维伟. 浅论预制装配式建筑施工技术研究与应用[A]. 建筑科技发展论坛[C]. 中国智慧工程研究会, 2024.',
        8:  '姜莉. 装配式混凝土结构建筑设计与施工研究[J]. 工程建设与设计, 2023(13): 235-237.',
        12: '吕航, 闫晶, 李杨, 等. BIM技术在施工项目管理中的应用研究——以某图书馆为例[J]. 价值工程, 2019, 38(24): 212-214.',
    },
    'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例': {
        4:  '上海市房屋管理局. 保障房工程项目信息——闵行浦江镇基地召楼路以东S8-01地块市属保障房项目[EB/OL]. (2023-01-30)[2026-09-20]. https://fgj.sh.gov.cn/bzfgcxmxx/20230130/5f37fcbfa2e64ab4be1e854e68ec8f9e.html',
        5:  '李双双. BIM技术在装配式建筑中的应用探究[D]. 邯郸: 河北工程大学, 2017.',
        6:  '姜莉. 装配式混凝土结构建筑设计与施工研究[J]. 工程建设与设计, 2023(13): 235-237.',
        10: '吕航, 闫晶, 李杨, 等. BIM技术在施工项目管理中的应用研究——以某图书馆为例[J]. 价值工程, 2019, 38(24): 212-214.',
        11: '刘丹. 装配式建筑设计标准体系构建研究[D]. 哈尔滨: 东北林业大学, 2018.',
        15: '中华人民共和国住房和城乡建设部. 建筑信息模型施工应用标准: GB/T 51235-2017[S]. 北京: 中国建筑工业出版社, 2017.',
    },
}

def set_text(para, new):
    runs = para.runs
    if not runs:
        return False
    runs[0].text = new
    for r in runs[1:]:
        r.text = ''
    return True

for nm, mapping in REPLACE.items():
    p = f'{W}\\{nm}\\{nm}.docx'
    d = Document(p)
    done = []
    for para in d.paragraphs:
        t = para.text.strip()
        m = re.match(r'^\[(\d+)\]\s*(.*)$', t, re.S)
        if not m:
            continue
        idx = int(m.group(1))
        if idx in mapping:
            if set_text(para, f'[{idx}]  {mapping[idx]}'):
                done.append(idx)
    print(f'{nm[:22]} 替换 {sorted(done)}')
    d.save(p)

# 论文二：正文角标 [10] -> [5]（《装配式混凝土结构技术规程》应为 JGJ 1-2014，即文献[5]）
nm2 = '装配式施工质量管理问题及优化研究——以市政项目为例'
d = Document(f'{W}\\{nm2}\\{nm2}.docx')
n = 0
for para in d.paragraphs:
    t = para.text
    if re.match(r'^\[\d+\]', t.strip()):
        continue
    if '《装配式混' in t and '技术规程》' in t and '[10]' in t:
        for r in para.runs:
            if '[10]' in r.text:
                r.text = r.text.replace('[10]', '[5]')
                n += 1
                break
print(f'论文二 正文角标修正 {n} 处')
d.save(f'{W}\\{nm2}\\{nm2}.docx')

# 复核
print()
for nm in REPLACE:
    d = Document(f'{W}\\{nm}\\{nm}.docx')
    print('==', nm[:20])
    for para in d.paragraphs:
        t = para.text.strip()
        if re.match(r'^\[\d+\]', t):
            print('  ', t[:95])
