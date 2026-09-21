# -*- coding: utf-8 -*-
import sys, io, os, re, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import win32com.client as wc
from docx import Document

W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'
OUT = os.path.join(W, '.workbuddy-ai', 'work', 'out')
FILES = [
    ('论文一', '再生混凝土在装配式建筑中的应用评价——以住宅项目为例'),
    ('论文二', '装配式施工质量管理问题及优化研究——以市政项目为例'),
    ('论文三', 'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例'),
]
app = wc.gencache.EnsureDispatch('Word.Application')
app.Visible = False; app.DisplayAlerts = 0
stat = {}
for tag, nm in FILES:
    path = os.path.join(W, nm, nm + '.docx')
    doc = app.Documents.Open(path, ReadOnly=True)
    doc.Repaginate()
    pages = doc.ComputeStatistics(2)
    words = doc.ComputeStatistics(0)
    nt, ni = len(doc.Tables), len(doc.InlineShapes)
    doc.Close(False)
    d = Document(path)
    nref = len([p for p in d.paragraphs if re.match(r'^\[\d+\]', p.text.strip())])
    nfig = len([p for p in d.paragraphs if re.match(r'^图\s*\d+', p.text.strip())])
    ntab = len([p for p in d.paragraphs if re.match(r'^表\s*\d+', p.text.strip())])
    stat[tag] = dict(name=nm, pages=pages, words=words, tables=nt, pics=ni,
                     ref=nref, fig=nfig, tabcap=ntab)
    print(tag, pages, words, nt, ni, nref, nfig, ntab)
app.Quit()
json.dump(stat, open(os.path.join(OUT, 'stat.json'), 'w', encoding='utf-8'),
          ensure_ascii=False, indent=1)
