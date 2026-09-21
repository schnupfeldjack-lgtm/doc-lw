# -*- coding: utf-8 -*-
"""
按模板修复论文结构：
1) 页眉改为「右对齐制表位 8306 + 制表符」，杜绝页码两位数时换行
2) 补成与模板一致的 4 节：
     节1 表格/封面区  (1701/1701, 空页眉)
     节2 摘要页       (1418/1021, gutter510, 无页眉引用 -> 继承空页眉)
     节3 目录页       (1800/1800, 无页眉引用 -> 继承空页眉)
     节4 正文         (1800/1800, 有页眉 rId4)
"""
import re, copy, sys
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
def q(t): return f"{{{W}}}{t}"

BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
TPL = BASE + r"\炎黄职业技术学院毕业论文模板(1).docx"
DOCS = [
    (BASE + r"\再生混凝土在装配式建筑中的应用评价——以住宅项目为例\再生混凝土在装配式建筑中的应用评价——以住宅项目为例.docx", "论文一"),
    (BASE + r"\装配式施工质量管理问题及优化研究——以市政项目为例\装配式施工质量管理问题及优化研究——以市政项目为例.docx", "论文二"),
    (BASE + r"\BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例\BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例.docx", "论文三"),
]

# ---------- 1. 取模板两个 sectPr 母本 ----------
tpl = Document(TPL)
sect_abs = tpl.paragraphs[50]._p.find(q('pPr')).find(q('sectPr'))   # 摘要节
sect_toc = tpl.paragraphs[74]._p.find(q('pPr')).find(q('sectPr'))   # 目录节
assert sect_abs is not None and sect_toc is not None

def add_sectpr(par, src):
    pPr = par._p.find(q('pPr'))
    if pPr is None:
        pPr = etree.SubElement(par._p, q('pPr'))
        par._p.insert(0, pPr)
    old = pPr.find(q('sectPr'))
    if old is not None:
        pPr.remove(old)
    new = copy.deepcopy(src)
    rPr = pPr.find(q('rPr'))
    if rPr is not None:
        rPr.addnext(new)
    else:
        pPr.append(new)

def fix_header(doc):
    """页眉：空格串 -> 右对齐制表符"""
    n = 0
    for s in doc.sections:
        for p in s.header.paragraphs:
            if '炎黄职业技术学院毕业论文' not in p.text:
                continue
            if '\t' in p.text:
                continue
            runs = p.runs
            if len(runs) < 2:
                continue
            # 找到第一个空白 run（标题之后）
            idx = None
            for i in range(1, len(runs)):
                if runs[i].text.strip() == '':
                    idx = i
                    break
            if idx is None:
                continue
            sz = runs[idx]._element.find(q('rPr'))  # 沿用其 rPr
            # 删除标题之后、'第' 之前的所有空白 run
            j = idx
            while j < len(runs) and runs[j].text.strip() == '':
                runs[j]._element.getparent().remove(runs[j]._element)
                j += 1
            # 插入制表符 run
            from docx.oxml import OxmlElement
            r = OxmlElement('w:r')
            if sz is not None:
                r.append(copy.deepcopy(sz))
            t = OxmlElement('w:t')
            t.text = '\t'
            r.append(t)
            runs[0]._element.addnext(r)
            # pPr 加右对齐制表位 8306
            pPr = p._p.find(q('pPr'))
            if pPr is None:
                pPr = OxmlElement('w:pPr'); p._p.insert(0, pPr)
            tabs = pPr.find(q('tabs'))
            if tabs is None:
                tabs = OxmlElement('w:tabs'); pPr.append(tabs)
            for old in tabs.findall(q('tab')):
                tabs.remove(old)
            tab = OxmlElement('w:tab')
            tab.set(q('val'), 'right')
            tab.set(q('pos'), '8306')
            tabs.append(tab)
            n += 1
    return n

for path, nm in DOCS:
    d = Document(path)
    hp = fix_header(d)
    # 定位目录标题 与 第一章标题
    toc_idx = chap_idx = None
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if toc_idx is None and re.match(r'^目\s*录$', t):
            toc_idx = i
        if chap_idx is None and re.match(r'^\d+\s{1,2}\S', t) and '\t' not in t and len(t) < 40:
            chap_idx = i
    assert toc_idx is not None and chap_idx is not None, (nm, toc_idx, chap_idx)
    add_sectpr(d.paragraphs[toc_idx - 1], sect_abs)   # 摘要节结束于目录前一段
    add_sectpr(d.paragraphs[chap_idx], sect_toc)      # 目录节结束于第一章标题（与模板一致）
    d.save(path)
    print(f"{nm}: 页眉修复 {hp} 处；摘要节分节符插入于段落#{toc_idx-1}，目录节分节符插入于段落#{chap_idx}({d.paragraphs[chap_idx].text[:12]!r})")
print("完成")
