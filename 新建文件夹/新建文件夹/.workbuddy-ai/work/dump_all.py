# -*- coding: utf-8 -*-
"""完整 dump 模板：段落文本 + 样式 + 字号 + 字体 + 对齐 + 缩进 + 行距"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.shared import Pt

W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'
p = f'{W}\\炎黄职业技术学院毕业论文模板(1).docx'
d = Document(p)

print('########## 段落 ##########')
for i, para in enumerate(d.paragraphs):
    t = para.text
    pf = para.paragraph_format
    runs = para.runs
    if runs:
        r = runs[0]
        sz = r.font.size.pt if r.font.size else None
        fn = r.font.name
        b = r.font.bold
        ea = None
        try:
            rpr = r._element.rPr
            if rpr is not None and rpr.rFonts is not None:
                ea = rpr.rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
        except Exception:
            pass
    else:
        sz = fn = b = ea = None
    ls = pf.line_spacing
    ind = (pf.left_indent, pf.first_line_indent)
    print(f'{i:3d} | sty={para.style.name} | sz={sz} | fn={fn} | ea={ea} | b={b} | al={para.alignment} | ls={ls} | ind={ind}')
    print(f'      >> {t[:110]}')

print()
print('########## 表格 ##########')
for ti, tb in enumerate(d.tables):
    print(f'--- table {ti} style={tb.style.name if tb.style else None} rows={len(tb.rows)} cols={len(tb.columns)}')
    for ri, row in enumerate(tb.rows):
        for ci, cell in enumerate(row.cells):
            for para in cell.paragraphs:
                runs = para.runs
                if runs:
                    r = runs[0]
                    sz = r.font.size.pt if r.font.size else None
                    fn = r.font.name
                    b = r.font.bold
                else:
                    sz = fn = b = None
                print(f'   [{ri},{ci}] sz={sz} fn={fn} b={b} al={para.alignment} >> {para.text[:80]}')

print()
print('########## 节 ##########')
for si, s in enumerate(d.sections):
    print(f'--- section {si}: {s.page_width} x {s.page_height}  margins L={s.left_margin} R={s.right_margin} T={s.top_margin} B={s.bottom_margin}')
    hdr = s.header
    for para in hdr.paragraphs:
        print('    HDR >>', repr(para.text))
    ftr = s.footer
    for para in ftr.paragraphs:
        print('    FTR >>', repr(para.text))
