# -*- coding: utf-8 -*-
"""
按模板母段落精确值修正两处差异：
  1) 摘要段行距 -> line=560 / lineRule=exact（模板段31）
  2) 致谢正文首行缩进 -> firstLine=573（模板段136）
"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'
FILES = [
    '再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
    '装配式施工质量管理问题及优化研究——以市政项目为例',
    'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例',
]
PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr',
             'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs',
             'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct',
             'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
             'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents',
             'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
             'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr',
             'pPrChange']


def ppr(para):
    pPr = para._element.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        para._element.insert(0, pPr)
    return pPr


def sort_ppr(pPr):
    kids = list(pPr)
    kids.sort(key=lambda e: PPR_ORDER.index(e.tag.split('}')[1])
              if e.tag.split('}')[1] in PPR_ORDER else 999)
    for k in kids:
        pPr.append(k)


for nm in FILES:
    p = f'{W}\\{nm}\\{nm}.docx'
    d = Document(p)
    ps = d.paragraphs
    n1 = n2 = 0
    ack_body = False
    for para in ps:
        t = para.text.strip()
        # 1) 摘要
        if t.startswith('摘 要：'):
            pPr = ppr(para)
            s = pPr.find(qn('w:spacing'))
            if s is None:
                s = OxmlElement('w:spacing')
                pPr.append(s)
            s.set(qn('w:line'), '560')
            s.set(qn('w:lineRule'), 'exact')
            sort_ppr(pPr)
            n1 += 1
        # 2) 致谢正文（致谢标题之后的正文段）
        if t == '致  谢':
            ack_body = True
            continue
        if ack_body and t.startswith('本论文是在'):
            pPr = ppr(para)
            ind = pPr.find(qn('w:ind'))
            if ind is None:
                ind = OxmlElement('w:ind')
                pPr.append(ind)
            ind.set(qn('w:firstLine'), '573')
            if ind.get(qn('w:firstLineChars')):
                del ind.attrib[qn('w:firstLineChars')]
            sort_ppr(pPr)
            n2 += 1
            ack_body = False
    print(f'{nm[:18]} 摘要行距修正 {n1} 处；致谢正文缩进修正 {n2} 处')
    d.save(p)

# 复核
print()
for nm in FILES:
    d = Document(f'{W}\\{nm}\\{nm}.docx')
    for para in d.paragraphs:
        t = para.text.strip()
        if t.startswith('摘 要：'):
            pPr = para._element.find(qn('w:pPr'))
            s = pPr.find(qn('w:spacing'))
            print(f'{nm[:14]} 摘要 line={s.get(qn("w:line"))} lineRule={s.get(qn("w:lineRule"))}')
        if t.startswith('本论文是在'):
            pPr = para._element.find(qn('w:pPr'))
            ind = pPr.find(qn('w:ind'))
            print(f'{nm[:14]} 致谢正文 firstLine={ind.get(qn("w:firstLine")) if ind is not None else None}')
