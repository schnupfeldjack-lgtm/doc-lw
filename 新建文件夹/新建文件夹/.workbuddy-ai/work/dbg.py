# -*- coding: utf-8 -*-
import re, sys
from docx import Document
from docx.oxml.ns import qn
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
def q(t): return f"{{{W}}}{t}"
BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
p1 = BASE + r"\再生混凝土在装配式建筑中的应用评价——以住宅项目为例\再生混凝土在装配式建筑中的应用评价——以住宅项目为例.docx"
d = Document(p1)
# 正文过滤器命中的前 5 段
n = 0
print("=== 【10】过滤器命中的前5段 ===")
for p in d.paragraphs:
    t = p.text.strip()
    if (len(t) >= 60 and not re.match(r'^\d', t) and '\t' not in p.text
            and not t.startswith('[') and not re.match(r'^摘\s*要[：:]', t)
            and not t.startswith('Abstract') and not re.match(r'^\[\d+\]', t)):
        print(f"   {t[:50]!r}")
        n += 1
        if n >= 5: break
print("\n=== 结论标题后 8 段 ===")
for i, p in enumerate(d.paragraphs):
    if re.match(r'^结\s*论$', p.text.strip()):
        for j in range(i, min(i + 8, len(d.paragraphs))):
            print(f"   [{j}] {d.paragraphs[j].text[:50]!r}")
        break
print("\n=== 致谢标题后 8 段 ===")
for i, p in enumerate(d.paragraphs):
    if re.match(r'^致\s*谢$', p.text.strip()):
        for j in range(i, min(i + 8, len(d.paragraphs))):
            print(f"   [{j}] {d.paragraphs[j].text[:50]!r}")
        break
