# -*- coding: utf-8 -*-
import docx, os, copy
from docx.oxml.ns import qn
BASE = r"C:/Users/15515/Desktop/09-文档资料/新建文件夹"
sub = "再生混凝土在装配式建筑中的应用评价——以住宅项目为例"
p_ = os.path.join(BASE, sub, sub+".docx")
d = docx.Document(p_); ps = d.paragraphs
p = ps[145]
runs = p.runs
assert len(runs)==2 and runs[0].text.startswith('三个维度与下设指标'), [r.text[:20] for r in runs]
body = runs[1]                       # 原始整段正文 run（保留其格式）
body.text = ('三个维度之间的关系是：技术性能是前提，不达标就不能用；施工适应性是关键，'
             '决定了能否在工程上顺利推行；经济与环境效益是动力，决定了各方是否有积极性去推广。'
             '评价时先判断技术性能是否满足标准要求，再考察施工适应性和经济环境效益。')
ins = copy.deepcopy(body._element)
for t in ins.iter(qn('w:t')): t.text = '三个维度与下设指标共同构成三层结构的评价体系，如图4–1所示。'
tail = copy.deepcopy(body._element)
for t in tail.iter(qn('w:t')): t.text = '表4–1列出了评价指标体系的具体构成。'
runs[0]._element.getparent().remove(runs[0]._element)   # 删除误置于段首的 run
body._element.addnext(ins)
ins.addnext(tail)
d.save(p_)
d2 = docx.Document(p_)
print("段145:", d2.paragraphs[145].text)
