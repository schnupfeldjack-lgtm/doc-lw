# -*- coding: utf-8 -*-
"""
目录条目点号对齐模板原文：
  模板一级： 1.引言            （半角点，数字与文字间无空格）
  模板二级： 4.1品牌的适用...   （半角点，数字与文字间无空格）
现状：  一级 "1  绪论"  二级 "1．1  研究背景"
"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'
FILES = [
    '再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
    '装配式施工质量管理问题及优化研究——以市政项目为例',
    'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例',
]

for nm in FILES:
    p = f'{W}\\{nm}\\{nm}.docx'
    d = Document(p)
    n = 0
    for para in d.paragraphs:
        if '\t' not in para.text:
            continue
        t = para.text
        m = re.match(r'^(\d+)\s*[.．]?\s*(\d+)?\s*[.．]?\s*(.*?)(\t.*)$', t)
        if not m:
            continue
        a, b, rest, tail = m.group(1), m.group(2), m.group(3), m.group(4)
        if not rest:
            continue
        new = f'{a}.{rest}{tail}' if b is None else f'{a}.{b}{rest}{tail}'
        if new == t:
            continue
        # 改写第一个 run 的文本（目录条目一般 1~2 个 run）
        runs = para.runs
        if not runs:
            continue
        # 拼出全部 run 的合并文本再重排
        full = ''.join(r.text for r in runs)
        # 只替换编号部分（tab 之前）
        head, _, pg = full.partition('\t')
        head_new = f'{a}.{rest}' if b is None else f'{a}.{b}{rest}'
        if head == head_new:
            continue
        # 把新文本放进第一个 run，其余 run 清空
        runs[0].text = head_new + '\t' + pg
        for r in runs[1:]:
            r.text = ''
        n += 1
    print(f'{nm[:20]} 改写 {n} 条')
    d.save(p)

# 复核
print()
for nm in FILES:
    d = Document(f'{W}\\{nm}\\{nm}.docx')
    toc = [p.text for p in d.paragraphs if '\t' in p.text]
    print(f'{nm[:16]} 示例: {toc[:2]} ... {toc[-3:]}')
