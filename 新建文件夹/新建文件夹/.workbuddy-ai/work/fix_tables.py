# -*- coding: utf-8 -*-
"""
表格填写内容字号对齐模板：
  开题报告 / 任务书 / 中期检查表 (T0/T1/T2) -> 小四 12pt（模板示例即 12pt）
  封面年月数字 (T4 第0、2格) -> 三号 16pt 楷体_GB2312
    依据模板重要提示第2条：「本封面的填写内容一律用"三号楷体_GB2312"」
  封面 T3 已是 16pt 楷体_GB2312，保持不变
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'
FILES = [
    '再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
    '装配式施工质量管理问题及优化研究——以市政项目为例',
    'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例',
]

for nm in FILES:
    p = f'{W}\\{nm}\\{nm}.docx'
    d = Document(p)
    n1 = n2 = 0
    for ti in (0, 1, 2):
        for row in d.tables[ti].rows:
            for c in row.cells:
                for para in c.paragraphs:
                    for r in para.runs:
                        if not r.text.strip():
                            continue
                        if r.font.size is None or abs(r.font.size.pt - 12.0) > 0.01:
                            r.font.size = Pt(12)
                            n1 += 1
    # 封面年月数字
    for row in d.tables[4].rows:
        for ci, c in enumerate(row.cells):
            if ci not in (0, 2):
                continue
            for para in c.paragraphs:
                for r in para.runs:
                    if not r.text.strip():
                        continue
                    r.font.size = Pt(16)
                    rPr = r._element.get_or_add_rPr()
                    rf = rPr.find(qn('w:rFonts'))
                    if rf is None:
                        rf = OxmlElement('w:rFonts')
                        rPr.insert(0, rf)
                    rf.set(qn('w:eastAsia'), '楷体_GB2312')
                    n2 += 1
    print(f'{nm[:18]} 表格字号改12pt {n1} 处；封面年月数字改三号楷体 {n2} 处')
    d.save(p)
