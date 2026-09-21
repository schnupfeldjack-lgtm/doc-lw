# -*- coding: utf-8 -*-
"""修复：误复制出的重复"致  谢"段 -> 只保留 1 个标题，其余清空为空段落"""
import sys, io, os, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

W = r'C:/Users/15515/Desktop/09-文档资料/新建文件夹'
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']
for nm in FILES:
    path = os.path.join(W, nm, nm + '.docx')
    d = Document(path); ps = d.paragraphs
    idx = [i for i, p in enumerate(ps) if p.text.strip() == '致  谢']
    print('=' * 70); print(nm[:28], ' 致谢段位置:', idx)
    if len(idx) <= 1:
        print('   无需修复')
        continue
    keep = idx[0]
    for i in idx[1:]:
        p = ps[i]
        for r in p.runs:
            for t in r._element.findall(qn('w:t')):
                r._element.remove(t)
        print(f'   已清空重复致谢段 {i}')
    # 确认标题前空行数
    ps2 = d.paragraphs
    ki = [i for i, p in enumerate(ps2) if p.text.strip() == '致  谢'][0]
    n = 0; j = ki - 1
    while j >= 0 and not ps2[j].text.strip():
        n += 1; j -= 1
    print(f'   致  谢 前空行数 = {n}')
    if n < 2:
        src = ps2[ki - 1]._element if n > 0 else ps2[ki]._element
        for _ in range(2 - n):
            el = copy.deepcopy(src)
            for a in (qn('w14:paraId'),):
                if el.get(a) is not None: del el.attrib[a]
            for t in el.iter(qn('w:t')):
                t.text = ''
            ps2[ki]._element.addprevious(el)
        print(f'   补空行至 2')
    d.save(path)
print('完成')
