# -*- coding: utf-8 -*-
"""
终审：把模板"明文规定"逐条列出，对三篇论文实测取证。
输出每条规定的【模板要求 / 实测值 / 判定】。
"""
import re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
def q(t): return f"{{{W}}}{t}"

BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
TPL = BASE + r"\炎黄职业技术学院毕业论文模板(1).docx"
DOCS = [
    (BASE + r"\再生混凝土在装配式建筑中的应用评价——以住宅项目为例\再生混凝土在装配式建筑中的应用评价——以住宅项目为例.docx", "论文一"),
    (BASE + r"\装配式施工质量管理问题及优化研究——以市政项目为例\装配式施工质量管理问题及优化研究——以市政项目为例.docx", "论文二"),
    (BASE + r"\BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例\BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例.docx", "论文三"),
]

def twips(v):
    return None if v is None else int(v) / 20.0

def par_fmt(par):
    """取段落的 (字体, 字号pt, 加粗, 行距, 段前pt, 段后pt, 缩进pt, 对齐)"""
    pPr = par._p.find(q('pPr'))
    sp = pPr.find(q('spacing')) if pPr is not None else None
    ind = pPr.find(q('ind')) if pPr is not None else None
    jc = pPr.find(q('jc')) if pPr is not None else None
    pm = pPr.find(q('rPr')) if pPr is not None else None
    run = None
    for r in par.runs:
        if r.text.strip():
            run = r; break
    if run is None and par.runs: run = par.runs[0]
    rr = run._element.find(q('rPr')) if run is not None else None
    def pick(pr, tag):
        if pr is None: return None
        e = pr.find(q(tag))
        return e.get(q('val')) if e is not None else None
    ea = None
    for pr in (rr, pm):
        if pr is None: continue
        rf = pr.find(q('rFonts'))
        if rf is not None and rf.get(q('eastAsia')):
            ea = rf.get(q('eastAsia')); break
    ascii_f = None
    for pr in (rr, pm):
        if pr is None: continue
        rf = pr.find(q('rFonts'))
        if rf is not None and rf.get(q('ascii')):
            ascii_f = rf.get(q('ascii')); break
    sz = pick(rr, 'sz') or pick(pm, 'sz')
    bold = (rr is not None and rr.find(q('b')) is not None)
    line = pick(sp, 'spacing', 'line') if False else (sp.get(q('line')) if sp is not None else None)
    lrule = sp.get(q('lineRule')) if sp is not None else None
    before = sp.get(q('before')) if sp is not None else None
    after = sp.get(q('after')) if sp is not None else None
    first = ind.get(q('firstLine')) if ind is not None else None
    return dict(
        font=ea or ascii_f,
        sz=(int(sz)/2.0 if sz else None),
        bold=bold,
        line=(f"{int(line)/240.0:.2f}倍" if line and lrule == 'auto' else (f"{int(line)/20.0:.0f}磅(固定)" if line else None)),
        before=(twips(before) if before else 0),
        after=(twips(after) if after else 0),
        first=(twips(first) if first else 0),
        jc=(jc.get(q('val')) if jc is not None else 'left'),
    )

def fmt_s(d):
    return (f"{d['font']} {d['sz']}pt{' 加粗' if d['bold'] else ''} | 行距{d['line']} | "
            f"段前{d['before']}pt 段后{d['after']}pt | 首行缩进{d['first']}pt | {d['jc']}")

def find(d, pred, n=1):
    out = []
    for p in d.paragraphs:
        if pred(p.text):
            out.append(p)
            if len(out) >= n: break
    return out

def count_cn(s):
    return len(re.findall(r'[\u4e00-\u9fff]', s))

print("=" * 100)
print("模板明文规定 vs 三篇论文实测")
print("=" * 100)

tpl = Document(TPL)
docs = [(Document(p), nm) for p, nm in DOCS]

# ---------- 1 摘要 ----------
print("\n【1】摘要：标签+正文 小四号宋体加粗；正文行距1.5倍；字数 200~300")
for d, nm in docs:
    ps = find(d, lambda t: re.match(r'^摘\s*要[：:]', t.strip()))
    f = par_fmt(ps[0]) if ps else None
    body = ps[0].text.split('：', 1)[1] if ps else ''
    n = len(body.replace(' ', ''))
    print(f"  {nm}: {fmt_s(f) if f else '缺失'} | 正文字数={n} {'✓' if 200 <= n <= 300 else '★'}")
print(f"  模板母段落: {fmt_s(par_fmt(find(tpl, lambda t: t.startswith('摘 要'))[0]))}")

# ---------- 2 关键词 ----------
print("\n【2】关键词：标签小四宋体加粗；内容小四宋体")
for d, nm in docs:
    ps = find(d, lambda t: t.strip().startswith('关键词'))
    f = par_fmt(ps[0]) if ps else None
    print(f"  {nm}: {fmt_s(f) if f else '缺失'}")

# ---------- 3 Abstract ----------
print("\n【3】Abstract：小四号 Times New Roman 加粗；正文 1.5倍行距；200实词左右")
for d, nm in docs:
    ps = find(d, lambda t: t.strip().startswith('Abstract'))
    f = par_fmt(ps[0]) if ps else None
    body = ps[0].text.split('：', 1)[1] if ps else ''
    w = len(re.findall(r"[A-Za-z']+", body))
    print(f"  {nm}: {fmt_s(f) if f else '缺失'} | 实词数={w} {'✓' if 150 <= w <= 260 else '★'}")

# ---------- 4 Key words ----------
print("\n【4】Key words：小四号 Times New Roman 加粗")
for d, nm in docs:
    ps = find(d, lambda t: t.strip().startswith('Key words'))
    f = par_fmt(ps[0]) if ps else None
    print(f"  {nm}: {fmt_s(f) if f else '缺失'}")

# ---------- 5 目录标题 ----------
print("\n【5】目录标题：黑体小二号(18pt)加粗居中   （模板明文：黑体小二号字加粗居中）")
for d, nm in docs:
    ps = find(d, lambda t: re.match(r'^目\s*录$', t.strip()))
    f = par_fmt(ps[0]) if ps else None
    ok = f and f['sz'] == 18 and f['bold'] and f['jc'] == 'center' and '黑体' in (f['font'] or '')
    print(f"  {nm}: {fmt_s(f) if f else '缺失'} {'✓' if ok else '★'}")

# ---------- 6 目录条目 ----------
print("\n【6】目录条目：宋体小四号(12pt)，行距固定值28磅")
for d, nm in docs:
    ps = find(d, lambda t: '\t' in t and re.search(r'\t\d+\s*$', t))
    f = par_fmt(ps[0]) if ps else None
    ok = f and f['sz'] == 12 and f['line'] == '28磅(固定)'
    print(f"  {nm}: {fmt_s(f) if f else '缺失'} {'✓' if ok else '★'}")
print(f"  模板母段落: {fmt_s(par_fmt(find(tpl, lambda t: '\t' in t and t.strip().endswith('1'))[0]))}")

# ---------- 7 章标题 ----------
print("\n【7】章标题：小3号黑体(15pt)加粗，段前0.5行段后0.5行")
for d, nm in docs:
    ps = find(d, lambda t: re.match(r'^\d+\s{1,2}\S', t.strip()) and '\t' not in t, 99)
    vals = [par_fmt(p) for p in ps]
    bad = [v for v in vals if v['sz'] != 15 or not v['bold'] or '黑体' not in (v['font'] or '')]
    sp = sorted({(v['before'], v['after']) for v in vals})
    print(f"  {nm}: 共{len(ps)}章 | 字号/字体/加粗不符={len(bad)} | 段前段后取值集合={sp}")
    print(f"       示例: {fmt_s(vals[0])}")
print(f"  模板母段落(第1章): {fmt_s(par_fmt(find(tpl, lambda t: t.startswith('1  引言'))[0]))}")
print(f"  模板母段落(第2章): {fmt_s(par_fmt(find(tpl, lambda t: t.startswith('2  ×'))[0]))}")

# ---------- 8/9 二三级标题 ----------
print("\n【8】二级标题：4号黑体(14pt)加粗   【9】三级标题：小4号黑体(12pt)不加粗")
for d, nm in docs:
    h2 = find(d, lambda t: re.match(r'^\d+．\d+', t.strip()) and '\t' not in t, 99)
    h3 = find(d, lambda t: re.match(r'^\d+\.\d+\.\d+', t.strip()) and '\t' not in t, 99)
    f2 = [par_fmt(p) for p in h2]; f3 = [par_fmt(p) for p in h3]
    b2 = [v for v in f2 if v['sz'] != 14 or not v['bold'] or '黑体' not in (v['font'] or '')]
    b3 = [v for v in f3 if v['sz'] != 12 or v['bold'] or '黑体' not in (v['font'] or '')]
    print(f"  {nm}: 二级{len(h2)}个 不符{len(b2)} | 三级{len(h3)}个 不符{len(b3)}"
          + (f" | 二级示例 {fmt_s(f2[0])}" if f2 else ""))

# ---------- 10 正文 ----------
print("\n【10】正文：小4号宋体(12pt)，1.5倍行距，首行缩进2字符(24pt)")
for d, nm in docs:
    ps = [p for p in d.paragraphs if len(p.text.strip()) >= 60
          and not re.match(r'^\d', p.text.strip()) and '\t' not in p.text
          and not p.text.strip().startswith('[')
          and not re.match(r'^摘\s*要[：:]', p.text.strip())
          and not p.text.strip().startswith('Abstract')
          and not p.text.strip().startswith('Key words')
          and not re.match(r'^\[\d+\]', p.text.strip())]
    vals = [par_fmt(p) for p in ps]
    bad = [v for v in vals if v['sz'] != 12 or '宋体' not in (v['font'] or '') or v['line'] != '1.50倍' or abs(v['first'] - 24) > 1]
    print(f"  {nm}: 正文段{len(ps)}个，不符{len(bad)}" + (f" | 示例 {fmt_s(vals[0])}" if vals else ""))
print(f"  模板母段落: {fmt_s(par_fmt(find(tpl, lambda t: t.startswith('×××××××××（小4号宋体）'))[0]))}")

# ---------- 11 图题表题 ----------
print("\n【11】图/表题：5号宋体(10.5pt)；编号按章分组、用 en dash（图3–5）")
for d, nm in docs:
    figs = find(d, lambda t: re.match(r'^图\s*\d+', t.strip()), 99)
    tbls = find(d, lambda t: re.match(r'^表\s*\d+', t.strip()), 99)
    allc = figs + tbls
    fv = [par_fmt(p) for p in allc]
    bad = [v for v in fv if v['sz'] != 10.5 or '宋体' not in (v['font'] or '')]
    nodash = [p.text[:20] for p in allc if '–' not in p.text and '-' not in p.text]
    print(f"  {nm}: 图题{len(figs)} 表题{len(tbls)} | 字号/字体不符{len(bad)}"
          + (f" | 示例 {fmt_s(fv[0])}" if fv else "") + (f" | 非en dash编号={nodash}" if nodash else ""))
    print(f"       图题示例文本: {[p.text[:26] for p in figs[:3]]}")

# ---------- 12 表格内文字 ----------
print("\n【12】表格内文字：模板注1「正文中公式、图与表的字体一律用5号宋体」")
for d, nm in docs:
    sizes = {}
    for ti, tb in enumerate(d.tables):
        if ti < 5: continue   # 跳过开题报告/任务书/中期检查/封面等模板原有表
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if p.text.strip():
                        v = par_fmt(p)
                        sizes.setdefault((v['font'], v['sz']), 0)
                        sizes[(v['font'], v['sz'])] += 1
    print(f"  {nm}: 自定义表格字号分布 {sizes}")

# ---------- 13 结论 ----------
print("\n【13】结论：标题四号宋体(14pt)加粗居中；正文小四宋体，固定值22磅")
for d, nm in docs:
    ps = find(d, lambda t: re.match(r'^结\s*论$', t.strip()))
    f = par_fmt(ps[0]) if ps else None
    idx = next((i for i, p in enumerate(d.paragraphs) if p._p is ps[0]._p), None) if ps else None
    bodyf = None
    if idx is not None:
        for p in d.paragraphs[idx + 1: idx + 6]:
            if len(p.text.strip()) > 30:
                bodyf = par_fmt(p); break
    ok = f and f['sz'] == 14 and f['bold'] and f['jc'] == 'center' and '宋体' in (f['font'] or '')
    ok2 = bodyf and bodyf['sz'] == 12 and bodyf['line'] == '22磅(固定)'
    print(f"  {nm}: 标题 {fmt_s(f) if f else '缺失'} {'✓' if ok else '★'}")
    print(f"       正文 {fmt_s(bodyf) if bodyf else '缺失'} {'✓' if ok2 else '★'}")

# ---------- 14 参考文献 ----------
print("\n【14】参考文献：标题四号宋体加粗居中；条目五号宋体(10.5pt)固定值22磅；≥5篇")
for d, nm in docs:
    ps = find(d, lambda t: re.match(r'^参\s*考\s*文\s*献$', t.strip()))
    f = par_fmt(ps[0]) if ps else None
    items = find(d, lambda t: re.match(r'^\[\d+\]', t.strip()), 99)
    fv = [par_fmt(p) for p in items]
    bad = [v for v in fv if v['sz'] != 10.5 or '宋体' not in (v['font'] or '')]
    ok = f and f['sz'] == 14 and f['bold'] and f['jc'] == 'center'
    print(f"  {nm}: 标题 {fmt_s(f) if f else '缺失'} {'✓' if ok else '★'} | 条目{len(items)}篇 {'✓' if len(items) >= 5 else '★'} | 条目格式不符{len(bad)}"
          + (f" | 示例 {fmt_s(fv[0])}" if fv else ""))

# ---------- 15 致谢 ----------
print("\n【15】致谢：标题四号宋体(14pt)加粗左对齐；正文小四宋体固定值22磅")
for d, nm in docs:
    ps = find(d, lambda t: re.match(r'^致\s*谢$', t.strip()))
    f = par_fmt(ps[0]) if ps else None
    idx = next((i for i, p in enumerate(d.paragraphs) if p._p is ps[0]._p), None) if ps else None
    bodyf = None
    if idx is not None:
        for p in d.paragraphs[idx + 1: idx + 6]:
            if len(p.text.strip()) > 30:
                bodyf = par_fmt(p); break
    ok = f and f['sz'] == 14 and f['bold'] and f['jc'] in ('left', None)
    print(f"  {nm}: 标题 {fmt_s(f) if f else '缺失'} {'✓' if ok else '★'} | 正文 {fmt_s(bodyf) if bodyf else '缺失'}")

# ---------- 16 颜色 / 角注 / 页眉 ----------
print("\n【16】全文颜色黑色；正文有角注；摘要/关键词/目录无角注；页眉阿拉伯数字")
for d, nm in docs:
    body = d.element.body
    cols = [c.get(q('val')) for c in body.findall(f".//{q('color')}")]
    nonblack = [c for c in cols if c and c.upper() != '000000' and c.upper() != 'AUTO']
    txt = '\n'.join(p.text for p in d.paragraphs)
    cites = re.findall(r'\[\d+[，,~\-–\d]*\]', txt)
    # 摘要段
    ab = find(d, lambda t: t.strip().startswith('摘 要'))
    kw = find(d, lambda t: t.strip().startswith('关键词'))
    ab_cite = bool(re.search(r'\[\d+\]', ab[0].text)) if ab else False
    kw_cite = bool(re.search(r'\[\d+\]', kw[0].text)) if kw else False
    print(f"  {nm}: 非黑run={len(nonblack)} | 正文角注数={len(cites)} | 摘要含角注={ab_cite} 关键词含角注={kw_cite}")

print("\n" + "=" * 100)
