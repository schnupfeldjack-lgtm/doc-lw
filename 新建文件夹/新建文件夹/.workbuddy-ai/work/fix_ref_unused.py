# -*- coding: utf-8 -*-
"""删除正文中无对应角注的文献条目，并重编号使角注与文献表严格对应"""
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

W = r'C:/Users/15515/Desktop/09-文档资料/新建文件夹'
JOBS = [
    ('装配式施工质量管理问题及优化研究——以市政项目为例', [10, 11, 12]),
    ('BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例', [7, 12]),
]
for nm, drop in JOBS:
    path = os.path.join(W, nm, nm + '.docx')
    d = Document(path)
    ps = d.paragraphs
    ref_idx = [i for i, p in enumerate(ps) if re.match(r'^\[\d+\]', p.text.strip())]
    keep_old = []
    for i in ref_idx:
        n = int(re.match(r'^\[(\d+)\]', ps[i].text.strip()).group(1))
        if n not in drop:
            keep_old.append((i, n))
    mapping = {old: new for new, (i, old) in enumerate(keep_old, 1)}
    print('=' * 72); print(nm[:30])
    print('  删除条目:', drop, ' 剩余:', len(keep_old))
    print('  编号映射:', mapping)
    # 1) 正文角注按映射替换（跳过文献条目段）
    refset = set(ref_idx)
    cnt = 0
    for i, p in enumerate(ps):
        if i in refset or '\t' in p.text:
            continue
        for r in p.runs:
            m = re.fullmatch(r'\[(\d+)\]', r.text)
            if m:
                old = int(m.group(1))
                if old in mapping:
                    if mapping[old] != old:
                        r.text = f'[{mapping[old]}]'
                        cnt += 1
                else:
                    print(f'    !! 段{i} 角注[{old}] 无映射（将被删除的文献）')
    print(f'  正文角注重编号: {cnt} 处')
    # 2) 删除被弃用的文献条目段
    for i in ref_idx:
        n = int(re.match(r'^\[(\d+)\]', ps[i].text.strip()).group(1))
        if n in drop:
            el = ps[i]._element
            el.getparent().remove(el)
    # 3) 文献条目重编号
    ps2 = d.paragraphs
    ref2 = [p for p in ps2 if re.match(r'^\[\d+\]', p.text.strip())]
    assert len(ref2) == len(keep_old), (len(ref2), len(keep_old))
    for p, (i, old) in zip(ref2, keep_old):
        for r in p.runs:
            if re.match(r'^\[\d+\]', r.text):
                r.text = re.sub(r'^\[\d+\]', f'[{mapping[old]}]', r.text)
                break
    d.save(path)
    print(f'  已保存。剩余 {len(ref2)} 条')
