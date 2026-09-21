# -*- coding: utf-8 -*-
"""把模板正文全文（不截断）打印出来，用于逐条提取格式明文规定"""
import sys
from docx import Document
from docx.oxml.ns import qn
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
def q(t): return f"{{{W}}}{t}"

p = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹\炎黄职业技术学院毕业论文模板(1).docx"
d = Document(p)
print("########## 正文段落全文 ##########")
for i, par in enumerate(d.paragraphs):
    t = par.text
    if t.strip():
        print(f"[{i:3d}] {t}")
print("\n########## 表格内容 ##########")
for ti, tb in enumerate(d.tables):
    print(f"\n----- 表 {ti} （行 {len(tb.rows)} x 列 {len(tb.columns)}）-----")
    for ri, row in enumerate(tb.rows):
        cells = []
        seen = set()
        for c in row.cells:
            if id(c._tc) in seen: continue
            seen.add(id(c._tc))
            cells.append(c.text.replace('\n', ' ').strip())
        print(f"  R{ri}: {cells}")
print("\n########## 页脚 ##########")
for i, s in enumerate(d.sections):
    for hp in [s.footer, s.even_page_footer, s.first_page_footer]:
        for pp in hp.paragraphs:
            if pp.text.strip():
                print(f"  节{i} 页脚: {pp.text!r}")
print("\n########## 样式表（styles.xml 里的字号/字体定义）##########")
import zipfile
from lxml import etree
z = zipfile.ZipFile(p)
root = etree.fromstring(z.read('word/styles.xml'))
for st in root.iter(q('style')):
    sid = st.get(q('styleId')); nm = st.find(q('name'))
    rpr = st.find(q('rPr'))
    info = ''
    if rpr is not None:
        rf = rpr.find(q('rFonts'))
        sz = rpr.find(q('sz'))
        b = rpr.find(q('b')) is not None
        info = f"字体={rf.get(q('eastAsia')) or rf.get(q('ascii')) if rf is not None else None} sz={sz.get(q('val')) if sz is not None else None} b={b}"
    print(f"  styleId={sid} name={nm.get(q('val')) if nm is not None else None} {info}")
