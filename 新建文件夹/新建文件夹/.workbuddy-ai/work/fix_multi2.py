# -*- coding: utf-8 -*-
"""按模板明文：两篇及以上文献论述同一观点时角注写 [4，5] / [6~8]
   逐处人工判定后合并（不同观点/不同分句的保持独立）"""
import docx, os, copy, re
from docx.oxml.ns import qn

BASE = r"C:/Users/15515/Desktop/09-文档资料/新建文件夹"
P1 = "再生混凝土在装配式建筑中的应用评价——以住宅项目为例"
P3 = "BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例"

def path(sub): return os.path.join(BASE, sub, sub+".docx")

def merge_runs(p, run_idx):
    """把 run_idx 列表中的多个 [n] run 合并成 [a，b，c]，保留第一个 run 的格式"""
    runs = p.runs
    nums = []
    for k in run_idx:
        m = re.fullmatch(r'\[(\d+)\]', runs[k].text.strip())
        assert m, runs[k].text
        nums.append(m.group(1))
    first = runs[run_idx[0]]
    first.text = '[' + '，'.join(nums) + ']'
    for k in run_idx[1:]:
        runs[k]._element.getparent().remove(runs[k]._element)
    print(f"    合并 -> {first.text}")

# ---------- 论文一 ----------
d = docx.Document(path(P1)); ps = d.paragraphs
print(P1)
for i, grp in [(172,[1,2]), (179,[1,2,3]), (189,[3,4])]:
    runs = ps[i].runs
    chk = [runs[k].text for k in grp]
    assert all(re.fullmatch(r'\[\d+\]', c) for c in chk), (i, chk)
    print(f"  段{i} 原: {''.join(chk)}")
    merge_runs(ps[i], grp)
d.save(path(P1)); print("  已保存")

# ---------- 论文三 ----------
d = docx.Document(path(P3)); ps = d.paragraphs
print(P3)
# 段144 合并 [9][10]
runs = ps[144].runs
print(f"  段144 原: {runs[1].text}{runs[2].text}")
merge_runs(ps[144], [1,2])

# 段104 拆句：...作了系统规定[3]。预制装配式...提供了参考[12]。《装配式混凝土结构技术规程》...
p = ps[104]; runs = p.runs
r1, r2, r3 = runs[1], runs[2], runs[3]
assert r1.text=='[3]' and r2.text=='[12]', (r1.text, r2.text)
# 1) 把 [12] 移到 r3 之后
r3._element.addnext(r2._element)
# 2) 在 [3] 之后插入正文格式的句号
body = copy.deepcopy(r3._element)          # 复制正文 run 的格式
for t in body.iter(qn('w:t')): t.text = '。'
r1._element.addnext(body)
d.save(path(P3)); print("  已保存")
print("  段104 新:", ps[104].text)
