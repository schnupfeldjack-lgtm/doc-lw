# -*- coding: utf-8 -*-
"""1) 结论/参考文献/致谢标题前补到 2 个空行（模板明文"（空2行）"）
   2) 去掉致谢标题的强制分页，与模板母段落一致"""
import sys, io, os, re, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

W = r'C:/Users/15515/Desktop/09-文档资料/新建文件夹'
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']
TARGET = {'结  论': 2, '参 考 文 献': 2, '致  谢': 2}

def clean(el):
    for a in (qn('w14:paraId'), '{http://schemas.microsoft.com/office/word/2010/wordml}paraId',
              '{http://schemas.microsoft.com/office/word/2010/wordml}textId'):
        if el.get(a) is not None:
            del el.attrib[a]
    return el

for nm in FILES:
    path = os.path.join(W, nm, nm + '.docx')
    d = Document(path); ps = d.paragraphs
    print('=' * 70); print(nm[:28])
    # 先去掉致谢强制分页
    for p in ps:
        if p.text.strip() == '致  谢':
            pPr = p._element.find(qn('w:pPr'))
            pb = pPr.find(qn('w:pageBreakBefore')) if pPr is not None else None
            if pb is not None:
                pPr.remove(pb); print('   已去掉 致  谢 的 pageBreakBefore')
    ps = d.paragraphs
    for i, p in enumerate(ps):
        t = p.text.strip()
        if t in TARGET:
            need = TARGET[t]
            n = 0; j = i - 1
            while j >= 0 and not ps[j].text.strip():
                n += 1; j -= 1
            if n < need:
                src = ps[i - 1]._element if n > 0 else ps[i]._element
                for _ in range(need - n):
                    el = clean(copy.deepcopy(src))
                    ps[i]._element.addprevious(el)
                print(f'   {t} 前空行 {n} -> {need}（补 {need-n} 个）')
            else:
                print(f'   {t} 前空行 {n}（已达标）')
    d.save(path)
print('完成')
