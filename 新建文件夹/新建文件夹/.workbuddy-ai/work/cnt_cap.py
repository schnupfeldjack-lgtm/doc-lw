# -*- coding: utf-8 -*-
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'
for nm in ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
           '装配式施工质量管理问题及优化研究——以市政项目为例',
           'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']:
    d = Document(os.path.join(W, nm, nm + '.docx'))
    print('='*60); print(nm[:26])
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if re.match(r'^(图|表)\s*\d+', t) and len(t) < 45 and p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            print(f'  题注 i={i}: {t}')
    print('  表格总数=', len(d.tables))
    for ti, tb in enumerate(d.tables):
        first = tb.rows[0].cells[0].text.strip()[:24] if tb.rows and tb.rows[0].cells else ''
        print(f'    表[{ti}] {len(tb.rows)}行x{len(tb.columns)}列  首格: {first}')
