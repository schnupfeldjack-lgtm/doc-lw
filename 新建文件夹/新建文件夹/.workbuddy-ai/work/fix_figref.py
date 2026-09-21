# -*- coding: utf-8 -*-
"""给论文一的 图4–1 补一句正文引用（模板注1：图与表应设置在文章中首次提到处附近）"""
import docx, os, copy
from docx.oxml.ns import qn
BASE = r"C:/Users/15515/Desktop/09-文档资料/新建文件夹"
sub = "再生混凝土在装配式建筑中的应用评价——以住宅项目为例"
p_ = os.path.join(BASE, sub, sub+".docx")
d = docx.Document(p_); ps = d.paragraphs
tgt = ps[145]
assert '表4–1列出了评价指标体系的具体构成' in tgt.text, tgt.text[-60:]
anchor = None
for r in tgt.runs:
    if '表4–1列出了评价指标体系的具体构成' in r.text:
        anchor = r
assert anchor is not None
new = copy.deepcopy(anchor._element)     # 复制正文 run 格式（非上标）
for t in new.iter(qn('w:t')): t.text = '三个维度与下设指标共同构成三层结构的评价体系，如图4–1所示。'
anchor._element.addprevious(new)
d.save(p_)
d2 = docx.Document(p_); print("新段145尾部:", d2.paragraphs[145].text[-120:])
