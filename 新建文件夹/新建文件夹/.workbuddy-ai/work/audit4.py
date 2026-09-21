# -*- coding: utf-8 -*-
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn
BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']

def sect_map(d, tag):
    print(f"\n===== {tag} =====")
    body = d.element.body
    pi = 0
    marks = []
    for ch in body.iterchildren():
        if ch.tag == qn('w:p'):
            # 段落内的 sectPr 表示"本节到此结束"
            pPr = ch.find(qn('w:pPr'))
            if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
                marks.append(('段末', pi))
            pi += 1
        elif ch.tag == qn('w:tbl'):
            pi += 1
        elif ch.tag == qn('w:sectPr'):
            marks.append(('body末', pi))
    print("  sectPr 位置:", marks)
    for i, s in enumerate(d.sections):
        print(f"  节{i}: 左{s.left_margin.cm:.2f} 右{s.right_margin.cm:.2f} 装订{s.gutter.cm:.2f} "
              f"类型{s.start_type} 页眉不同首页={s.different_first_page_header_footer}")

tpl = docx.Document(os.path.join(BASE,"炎黄职业技术学院毕业论文模板(1).docx"))
sect_map(tpl, "模板")
# 模板每节开头的内容
print("\n  模板段落预览(前20):")
for i in range(20):
    t = tpl.paragraphs[i].text.strip()
    if t: print(f"    TPL[{i}] {t[:56]}")
for sub in FILES:
    d = docx.Document(os.path.join(BASE, sub, sub+".docx"))
    sect_map(d, sub[:22])
