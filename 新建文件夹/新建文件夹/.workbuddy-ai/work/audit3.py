# -*- coding: utf-8 -*-
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn
BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']
def sect_info(d, tag):
    print(f"\n===== {tag} =====")
    for i, s in enumerate(d.sections):
        print(f" 节{i}: 页 {s.page_width.cm:.2f}x{s.page_height.cm:.2f}cm  "
              f"上{s.top_margin.cm:.2f} 下{s.bottom_margin.cm:.2f} 左{s.left_margin.cm:.2f} 右{s.right_margin.cm:.2f} "
              f"装订线{s.gutter.cm:.2f} 页眉距{s.header_distance.cm:.2f} 页脚距{s.footer_distance.cm:.2f}")
        hdr = s.header
        for p in hdr.paragraphs:
            xml = p._element.xml
            fld = 'PAGE' in xml or 'NUMPAGES' in xml
            tabs = p._element.find(qn('w:pPr'))
            tb = None
            if tabs is not None:
                t = tabs.find(qn('w:tabs'))
                if t is not None: tb = [(c.get(qn('w:val')), c.get(qn('w:pos'))) for c in t]
            print(f"    页眉: {p.text!r}  域={'有' if fld else '无'}  制表位={tb}  对齐={p.alignment}")
        ft = s.footer
        for p in ft.paragraphs:
            print(f"    页脚: {p.text!r}")

tpl = docx.Document(os.path.join(BASE,"炎黄职业技术学院毕业论文模板(1).docx"))
sect_info(tpl, "模板")
for sub in FILES:
    d = docx.Document(os.path.join(BASE, sub, sub+".docx"))
    sect_info(d, sub[:22])
