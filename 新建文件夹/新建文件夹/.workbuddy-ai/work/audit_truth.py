# -*- coding: utf-8 -*-
"""全量独立审计：不用预设类别清单，直接比对"格式指纹"
   论文每个段落的格式指纹必须在模板段落指纹集合中出现，否则列出"""
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx
from docx.oxml.ns import qn

BASE = r"C:\Users\15515\Desktop\09-文档资料\新建文件夹"
TPL  = os.path.join(BASE, "炎黄职业技术学院毕业论文模板(1).docx")
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']

def fp(p):
    """段落格式指纹"""
    r = p._element.find(qn('w:pPr'))
    d = {}
    if r is not None:
        def g(tag, attr=None):
            e = r.find(qn(tag))
            if e is None: return None
            return e.get(qn(attr)) if attr else (e.get(qn('w:val')) or '')
        d['style'] = g('w:pStyle', 'w:val') or ''
        sp = r.find(qn('w:spacing'))
        if sp is not None:
            for k in ('w:line','w:lineRule','w:before','w:after','w:beforeLines','w:afterLines'):
                d[k] = sp.get(qn(k))
        ind = r.find(qn('w:ind'))
        if ind is not None:
            for k in ('w:firstLine','w:firstLineChars','w:left','w:right','w:hanging','w:leftChars'):
                d[k] = ind.get(qn(k))
        d['jc'] = (g('w:jc','w:val') or '')
        d['pb']  = '1' if r.find(qn('w:pageBreakBefore')) is not None else ''
        d['kn']  = '1' if r.find(qn('w:keepNext')) is not None else ''
        d['kl']  = '1' if r.find(qn('w:keepLines')) is not None else ''
        d['outline'] = g('w:outlineLvl','w:val') or ''
        rpr = r.find(qn('w:rPr'))
        if rpr is not None:
            e = rpr.find(qn('w:sz'));   d['sz']  = e.get(qn('w:val')) if e is not None else None
            e = rpr.find(qn('w:rFonts'))
            if e is not None:
                d['ea'] = e.get(qn('w:eastAsia'))
                d['as'] = e.get(qn('w:ascii'))
                d['cs'] = e.get(qn('w:cs'))
            d['b'] = '1' if rpr.find(qn('w:b')) is not None else ''
    # run 级（取第一个 run）
    if p.runs:
        rpr = p.runs[0]._element.find(qn('w:rPr'))
        if rpr is not None:
            e = rpr.find(qn('w:sz'))
            if e is not None and d.get('sz') is None: d['sz'] = e.get(qn('w:val'))
            e = rpr.find(qn('w:rFonts'))
            if e is not None:
                if not d.get('ea'): d['ea'] = e.get(qn('w:eastAsia'))
                if not d.get('as'): d['as'] = e.get(qn('w:ascii'))
    return tuple(sorted(f"{k}={v}" for k, v in d.items() if v not in (None, '', '0')))

tpl = docx.Document(TPL)
tpl_fps = {}
for i, p in enumerate(tpl.paragraphs):
    tpl_fps.setdefault(fp(p), []).append(i)
print(f"模板段落 {len(tpl.paragraphs)} 个，不同格式指纹 {len(tpl_fps)} 种\n")

for sub in FILES:
    d = docx.Document(os.path.join(BASE, sub, sub+".docx"))
    ps = d.paragraphs
    print("="*78); print(sub)
    miss = {}
    for i, p in enumerate(ps):
        t = p.text.strip()
        f = fp(p)
        if f not in tpl_fps:
            miss.setdefault(f, []).append(i)
    print(f"  段落 {len(ps)} 个，指纹未在模板中出现的有 {len(miss)} 种 / {sum(len(v) for v in miss.values())} 段")
    for f, idxs in sorted(miss.items(), key=lambda x: -len(x[1])):
        print(f"\n  --- 出现 {len(idxs)} 次  指纹: {' '.join(f)}")
        for i in idxs[:3]:
            print(f"      段[{i}] {ps[i].text[:60]!r}")
