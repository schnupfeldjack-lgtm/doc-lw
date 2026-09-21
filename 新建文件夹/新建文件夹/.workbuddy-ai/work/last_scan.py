# -*- coding: utf-8 -*-
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
W = r'C:/Users/15515/Desktop/09-文档资料/新建文件夹'
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']
BAD = ['×××', 'XXX', 'xx', '张三', '李四', '指导教师：王', '学号：20']
for nm in FILES:
    d = Document(os.path.join(W, nm, nm + '.docx'))
    print('=' * 66); print(nm[:26])
    hit = []
    for p in d.paragraphs:
        for b in BAD:
            if b in p.text:
                hit.append(('段', b, p.text.strip()[:50]))
    for ti, tb in enumerate(d.tables):
        for row in tb.rows:
            for c in row.cells:
                for b in BAD:
                    if b in c.text:
                        hit.append((f'表{ti}', b, c.text.strip()[:50]))
    print('  模板占位符残留:', len(hit))
    for h in hit[:8]:
        print('    ', h)
    # 封面/任务书个人信息格
    t0, t2 = d.tables[0], d.tables[2]
    def cells(t):
        out = []
        for row in t.rows:
            for c in row.cells:
                out.append(c.text.strip())
        return out
    c0, c2 = cells(t0), cells(t2)
    print('  表0:', [x for x in c0 if x][:14])
    print('  表2:', [x for x in c2 if x][:10])
