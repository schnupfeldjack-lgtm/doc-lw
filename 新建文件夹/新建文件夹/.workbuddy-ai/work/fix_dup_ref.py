# -*- coding: utf-8 -*-
"""论文三：消除重复文献 [2]/[7]、[3]/[12]
   正文角标 [7]->[2]（GB/T 51212）, [12]->[3]（GB/T 51231）
   空出的 [7][12] 换成真实可查的扩展文献
"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

W = r'C:\Users\15515\Desktop\09-文档资料\新建文件夹'
nm = 'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例'
p = f'{W}\\{nm}\\{nm}.docx'
d = Document(p)

# 1) 正文角标替换（跳过文献条目本身）
n7 = n12 = 0
for para in d.paragraphs:
    t = para.text.strip()
    if re.match(r'^\[\d+\]', t):
        continue
    for r in para.runs:
        if '[7]' in r.text:
            r.text = r.text.replace('[7]', '[2]')
            n7 += 1
        if '[12]' in r.text:
            r.text = r.text.replace('[12]', '[3]')
            n12 += 1
print(f'正文角标 [7]->[2] {n7} 处, [12]->[3] {n12} 处')

# 2) 替换空出的条目
NEW = {
    7:  '杨婷, 李晓娟, 陈伟斌. 装配式建筑质量影响因素的识别和控制[J]. 哈尔滨商业大学学报(自然科学版), 2021(3): 328-331.',
    12: '崔宇, 柴维伟. 浅论预制装配式建筑施工技术研究与应用[A]. 建筑科技发展论坛[C]. 中国智慧工程研究会, 2024.',
}
for para in d.paragraphs:
    m = re.match(r'^\[(\d+)\]\s*(.*)$', para.text.strip(), re.S)
    if not m:
        continue
    idx = int(m.group(1))
    if idx in NEW and para.runs:
        para.runs[0].text = f'[{idx}]  {NEW[idx]}'
        for r in para.runs[1:]:
            r.text = ''

d.save(p)

# 复核：查重
d = Document(p)
refs = {}
for para in d.paragraphs:
    m = re.match(r'^\[(\d+)\]\s*(.*)$', para.text.strip(), re.S)
    if m:
        refs[int(m.group(1))] = m.group(2).strip()
print(f'\n文献条目 {len(refs)} 条')
seen = {}
dup = []
for k in sorted(refs):
    key = re.sub(r'[\s,.:;\[\]/]', '', refs[k])
    if key in seen:
        dup.append((seen[key], k))
    else:
        seen[key] = k
print('重复:', dup if dup else '无')
for k in sorted(refs):
    print(f'  [{k}] {refs[k][:88]}')
