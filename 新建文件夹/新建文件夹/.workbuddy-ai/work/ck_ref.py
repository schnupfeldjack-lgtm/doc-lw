# -*- coding: utf-8 -*-
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

W = r'C:/Users/15515/Desktop/09-文档资料/新建文件夹'
TPL = os.path.join(W, '炎黄职业技术学院毕业论文模板(1).docx')
FILES = ['再生混凝土在装配式建筑中的应用评价——以住宅项目为例',
         '装配式施工质量管理问题及优化研究——以市政项目为例',
         'BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例']
EMU = 12700.0

def pinfo(p):
    pPr = p._element.find(qn('w:pPr'))
    d = {}
    if pPr is not None:
        s = pPr.find(qn('w:spacing'))
        if s is not None:
            d['line'] = s.get(qn('w:line')); d['rule'] = s.get(qn('w:lineRule'))
            d['before'] = s.get(qn('w:before')); d['after'] = s.get(qn('w:after'))
        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            for k in ('firstLine', 'hanging', 'left'):
                v = ind.get(qn('w:' + k))
                if v: d[k] = f'{int(v)/20:.2f}磅'
        for k in ('keepNext', 'keepLines', 'pageBreakBefore'):
            if pPr.find(qn('w:' + k)) is not None: d[k] = 'Y'
    r = p.runs[0] if p.runs else None
    if r is not None:
        d['sz'] = r.font.size.pt if r.font.size else None
        rpr = r._element.rPr
        if rpr is not None and rpr.rFonts is not None:
            d['ea'] = rpr.rFonts.get(qn('w:eastAsia'))
        d['b'] = r.font.bold
    d['al'] = str(p.alignment).split()[0] if p.alignment is not None else None
    return d

print('############ 模板 ############')
d = Document(TPL)
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t.startswith('参 考 文 献') or re.match(r'^\[\d+\]', t) or '空2行' in t:
        print(f'  TPL[{i}] {t[:38]:<40} {pinfo(p)}')

for nm in FILES:
    print('=' * 78); print(nm[:30])
    d = Document(os.path.join(W, nm, nm + '.docx'))
    ps = d.paragraphs
    # 1 标题
    for i, p in enumerate(ps):
        if p.text.strip() == '参 考 文 献':
            print(f'  标题段[{i}] {pinfo(p)}')
            # 前两段是否空行（空2行）
            for j in (i-1, i-2, i-3):
                if j >= 0:
                    print(f'     前段[{j}] 文本={ps[j].text.strip()[:20]!r}')
    # 2 文献条目
    refs = [(i, p) for i, p in enumerate(ps) if re.match(r'^\[\d+\]', p.text.strip())]
    print(f'  条目数={len(refs)}  首条[{refs[0][0]}] {pinfo(refs[0][1])}')
    print(f'                     末条[{refs[-1][0]}] {pinfo(refs[-1][1])}')
    # 缩进一致性
    fls = set()
    for i, p in refs:
        pi = pinfo(p); fls.add((pi.get('sz'), pi.get('line'), pi.get('rule'), pi.get('firstLine')))
    print(f'  条目(sz,line,rule,firstLine)取值集合: {fls}')
    # 3 编号连续性
    nums = [int(re.match(r'^\[(\d+)\]', p.text.strip()).group(1)) for i, p in refs]
    ok = nums == list(range(1, len(nums) + 1))
    print(f'  编号连续1..N: {"OK" if ok else "NG"}  {nums[:6]}...{nums[-3:]}')
    # 4 正文角注集合 vs 文献序号
    body_marks = set()
    for p in ps:
        if re.match(r'^\[\d+\]', p.text.strip()):
            continue
        if '\t' in p.text:      # 目录条目
            continue
        for m in re.finditer(r'\[(\d+)\]', p.text):
            body_marks.add(int(m.group(1)))
    missing = sorted(set(nums) - body_marks)
    extra = sorted(body_marks - set(nums))
    print(f'  正文角注引用集合 vs 文献序号: 未被引用={missing}  引用了不存在的={extra}')
    # 5 题名/摘要/关键词/目录 中是否有角注
    bad = []
    for i, p in enumerate(ps):
        t = p.text.strip()
        if i > 130: break
        if re.search(r'\[\d+\]', t) and not re.match(r'^\[\d+\]', t):
            if ('\t' in p.text or t.startswith(('摘 要', '关键词', 'Abstract', 'Key words'))
                or t.startswith('目') or len(t) < 30):
                bad.append((i, t[:40]))
    print(f'  题名/摘要/关键词/目录中的角注: {len(bad)} 处 {bad[:3]}')
