# -*- coding: utf-8 -*-
from pathlib import Path
import difflib, re, subprocess, sys, zipfile
from docx import Document
from docx.oxml.ns import qn
import pdfplumber

def norm(s):return re.sub(r"\s+","",s).replace("–","-").replace("—","-")
def body_idx(d):
    xs=[i for i,p in enumerate(d.paragraphs) if norm(p.text) in ("1引言","1绪论") and not (p.style and p.style.name.lower().startswith("toc"))]
    if not xs: raise RuntimeError("body heading missing")
    return xs[-1]
def ref_idx(d):return next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="参考文献")
def body_text(d):return "\n".join(p.text for p in d.paragraphs[body_idx(d):ref_idx(d)])
def han(d):return sum("\u4e00"<=c<="\u9fff" for c in body_text(d))
def ptext(pdf,i):return subprocess.check_output(["pdftotext","-f",str(i),"-l",str(i),"-layout",str(pdf),"-"],text=True,errors="ignore")
def total_pages(pdf):
    s=subprocess.check_output(["pdfinfo",str(pdf)],text=True,errors="ignore")
    return int(re.search(r"^Pages:\s+(\d+)",s,re.M).group(1))
def title_top(pdf,page,needle):
    with pdfplumber.open(pdf) as f:
        words=f.pages[page-1].extract_words()
    hits=[w for w in words if any(ch in w["text"] for ch in needle)]
    return min((w["top"] for w in hits),default=9999)

def compare_unique(cur,other,label):
    cp=[norm(p.text) for p in cur.paragraphs[body_idx(cur):ref_idx(cur)] if len(norm(p.text))>=80]
    op=[norm(p.text) for p in other.paragraphs[body_idx(other):ref_idx(other)] if len(norm(p.text))>=80]
    best=0;pair=None
    for a in cp:
        for b in op:
            r=difflib.SequenceMatcher(None,a,b,autojunk=False).ratio()
            if r>best:best=r;pair=(a[:40],b[:40])
    a=norm(body_text(cur));b=norm(body_text(other))
    sm=difflib.SequenceMatcher(None,a,b,autojunk=False)
    longest=sm.find_longest_match(0,len(a),0,len(b)).size
    print("SIMILARITY",label,"max_para",round(best,3),"longest_exact",longest,"pair",pair)
    if best>=0.58 or longest>=90:
        raise AssertionError(f"与{label}存在过高正文相似: {best:.3f}, exact={longest}")

def main():
    docx=Path(sys.argv[1]);pdf=Path(sys.argv[2]);repo=Path(sys.argv[3])
    d=Document(docx);hc=han(d)
    assert 9500<=hc<=11000,hc
    full="\n".join(p.text for p in d.paragraphs)
    assert "待填写" not in full
    refs=[p.text.strip() for p in d.paragraphs if re.match(r"^\[\d+\]",p.text.strip())]
    assert len(refs)==15,len(refs)
    nums=[int(re.match(r"^\[(\d+)\]",x).group(1)) for x in refs]
    assert nums==list(range(1,16)),nums
    assert len(d.inline_shapes)>=5,len(d.inline_shapes)
    # 页眉/页码结构
    for s in d.sections[:-1]:
        hx=s.header._element.xml
        assert " PAGE " not in hx and "炎黄职业技术学院毕业论文" not in hx
    hx=d.sections[-1].header._element.xml
    assert "炎黄职业技术学院毕业论文" in hx and " PAGE " in hx
    pg=d.sections[-1]._sectPr.find(qn("w:pgNumType"))
    assert pg is not None and pg.get(qn("w:start"))=="1"
    # 图片二进制真实嵌入
    with zipfile.ZipFile(docx) as z:
        media=[x for x in z.namelist() if x.startswith("word/media/")]
        png=[x for x in media if x.endswith(".png")]
        assert len(png)>=5,(len(png),png)
    total=total_pages(pdf);texts=[ptext(pdf,i) for i in range(1,total+1)];np=[norm(x) for x in texts]
    # 前置：开题/任务/中期/封面/中摘/英摘/目录2页
    assert "开题报告" in np[0]
    assert "毕业论文任务书" in np[1]
    assert "中期检查" in np[2]
    assert "毕业论文" in np[3] and "装配式施工质量管理问题及优化研究" in np[3]
    assert "摘要" in np[4]
    assert "Abstract" in texts[5]
    assert "目录" in np[6]
    bphys=next(i for i,t in enumerate(np,1) if "第1页" in t and "1引言" in t)
    assert bphys==9,f"正文物理起始页应为9，实际{bphys}"
    for i in range(1,bphys):
        assert not re.search(r"第\s*\d+\s*页\s*共\s*\d+\s*页",texts[i-1]),f"前置第{i}页出现正文页码"
    body_pages=total-bphys+1
    assert f"第1页共{body_pages}页" in np[bphys-1],np[bphys-1][:100]
    assert f"第{body_pages}页共{body_pages}页" in np[-1],np[-1][:100]
    # 每一页都应有实际内容，防止空白页
    for i,t in enumerate(np,1):
        assert len(t)>20,f"第{i}页疑似空白"
    # 参考文献、致谢独立页且靠上
    rphys=next(i for i,t in enumerate(np,1) if "参考文献" in t and "[1]" in t)
    aphys=next(i for i,t in enumerate(np,1) if "致谢" in t and "指导老师" in t)
    rt=title_top(pdf,rphys,"参考文献")
    at=title_top(pdf,aphys,"致谢")
    assert rt<230,rt
    assert at<230,at
    # 与前两篇正文做实质相似性检查
    p1=repo/"新建文件夹/新建文件夹/BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例/BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例_最终版.docx"
    p2=repo/"新建文件夹/新建文件夹/再生混凝土在装配式建筑中的应用评价——以住宅项目为例/再生混凝土在装配式建筑中的应用评价——以住宅项目为例_最终优化版.docx"
    if p1.exists():compare_unique(d,Document(p1),"BIM论文")
    if p2.exists():compare_unique(d,Document(p2),"再生混凝土论文")
    print("QA_OK","han",hc,"refs",len(refs),"physical_pages",total,"body_start",bphys,"body_pages",body_pages,"ref_page",rphys,"ack_page",aphys,"ref_top",rt,"ack_top",at)

if __name__=="__main__":main()
