# -*- coding: utf-8 -*-
"""再生混凝土论文最终质量检查。"""
from pathlib import Path
import difflib, re, subprocess, sys, zipfile
from docx import Document
from docx.oxml.ns import qn
import pdfplumber

ROOT=Path(__file__).resolve().parents[2]
BIM=ROOT/"新建文件夹/新建文件夹/BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例/BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例_最终版.docx"

def norm(s):return re.sub(r"\s+","",s).replace("–","-").replace("—","-")
def body_idx(d):
    return [i for i,p in enumerate(d.paragraphs)
            if norm(p.text).startswith("1绪论") and not (p.style and p.style.name.lower().startswith("toc"))][-1]
def ref_idx(d):return next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="参考文献")
def ack_idx(d):return next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="致谢")
def han(d):
    return sum("\u4e00"<=c<="\u9fff" for p in d.paragraphs[body_idx(d):ref_idx(d)] for c in p.text)
def pdftext(pdf,i):
    return subprocess.check_output(["pdftotext","-f",str(i),"-l",str(i),"-layout",str(pdf),"-"],
                                   text=True,errors="ignore")
def pagecount(pdf):
    s=subprocess.check_output(["pdfinfo",str(pdf)],text=True,errors="ignore")
    return int(re.search(r"^Pages:\s+(\d+)",s,re.M).group(1))

def similarity_check(newd):
    if not BIM.exists():
        print("WARN BIM final missing, similarity check skipped");return
    old=Document(BIM)
    def paras(d):
        s=body_idx(d);e=ref_idx(d)
        return [norm(p.text) for p in d.paragraphs[s:e] if len(norm(p.text))>=60]
    a=paras(newd);b=paras(old)
    # paragraph-level maximum
    max_ratio=0;max_pair=("","")
    for x in a:
        for y in b:
            r=difflib.SequenceMatcher(None,x,y,autojunk=False).ratio()
            if r>max_ratio:max_ratio=r;max_pair=(x[:80],y[:80])
    # whole body longest common contiguous text
    aa="".join(a);bb="".join(b)
    m=difflib.SequenceMatcher(None,aa,bb,autojunk=False).find_longest_match()
    print("SIMILARITY","max_para",round(max_ratio,3),"longest_common",m.size)
    if max_ratio>=0.78 or m.size>=100:
        raise AssertionError(f"与BIM论文存在过高文本相似: para={max_ratio:.3f}, common={m.size}, pair={max_pair}")

def superscript_ratio(d):
    total=sup=0
    pat=re.compile(r"\[(?:\d+)(?:\s*[,，、;；\-–]\s*\d+)*\]")
    for p in d.paragraphs[body_idx(d):ref_idx(d)]:
        pos=0
        for r in p.runs:
            for _ in pat.finditer(r.text or ""):
                total+=1
                rp=r._r.find(qn("w:rPr"))
                va=rp.find(qn("w:vertAlign")) if rp is not None else None
                if va is not None and va.get(qn("w:val"))=="superscript":sup+=1
    return total,sup

def main():
    docx=Path(sys.argv[1]);pdf=Path(sys.argv[2])
    d=Document(docx)
    hc=han(d)
    assert 9500<=hc<=11000,f"正文汉字数={hc}"
    alltxt="\n".join(p.text for p in d.paragraphs)
    for bad in ("0.5396","0.2970","0.1634","综合评价 1.0000","AHP判断矩阵","评价等级的划分方面"):
        assert bad not in alltxt,f"残留旧主观评分内容: {bad}"
    assert "案例本身不能替代预制构件" in alltxt or "不能替代预制构件" in alltxt
    assert "现浇" in alltxt and "装配式" in alltxt
    # 不允许重复独立结论
    assert sum(1 for p in d.paragraphs if norm(p.text)=="结论")==0
    # 四幅原创图都实际嵌入
    assert len(d.inline_shapes)>=4,f"inline_shapes={len(d.inline_shapes)}"
    with zipfile.ZipFile(docx) as z:
        media=[n for n in z.namelist() if n.startswith("word/media/")]
        assert len(media)>=5,media  # 含校徽
    # 参考文献数量与中英文结构
    refs=[p.text.strip() for p in d.paragraphs[ref_idx(d)+1:ack_idx(d)] if re.match(r"^\[\d+\]",p.text.strip())]
    assert len(refs)==17,f"refs={len(refs)}"
    assert sum(any("a"<=c.lower()<="z" for c in r) for r in refs)>=8
    total_cite,sup=superscript_ratio(d)
    assert total_cite>10 and sup==total_cite,f"角标上标 {sup}/{total_cite}"

    similarity_check(d)

    total=pagecount(pdf);texts=[pdftext(pdf,i) for i in range(1,total+1)];nts=[norm(t) for t in texts]
    # 前4页严格独立
    assert "开题报告" in nts[0]
    assert "毕业论文任务书" in nts[1]
    assert "中期检查指导表" in nts[2]
    assert "毕业论文" in nts[3] and "再生混凝土" in nts[3]
    assert "摘要" in nts[4]
    assert "Abstract" in texts[5]
    assert "目录" in nts[6]
    bphys=next(i for i,t in enumerate(nts,1) if "炎黄职业技术学院毕业论文" in t and "第1页" in t and "1绪论" in t)
    assert bphys==9,f"正文物理起始页应为9，实际{bphys}"
    # 前置页没有正文页码
    for i in range(1,bphys):
        assert not re.search(r"第\s*\d+\s*页\s*共\s*\d+\s*页",texts[i-1]),f"前置第{i}页出现正文页码"
    body_pages=total-bphys+1
    assert f"第1页共{body_pages}页" in nts[bphys-1],nts[bphys-1][:150]
    assert f"第{body_pages}页共{body_pages}页" in nts[-1],nts[-1][:150]
    # 参考文献、致谢各自起页；致谢标题不能掉到页面中下部
    rphys=next(i for i,t in enumerate(nts,1) if "参考文献" in t and any("[1]" in texts[i-1] or "海然" in texts[i-1] for _ in [0]))
    aphys=next(i for i,t in enumerate(nts,1) if "致谢" in t and "论文" in t)
    with pdfplumber.open(pdf) as pf:
        words=pf.pages[aphys-1].extract_words()
        cand=[w for w in words if w["text"] in ("致","谢","致谢") or "致" in w["text"] or "谢" in w["text"]]
        assert cand,"无法识别致谢标题"
        top=min(w["top"] for w in cand)
        assert top<230,f"致谢标题过低 top={top}"
    print("QA_OK","han",hc,"pages",total,"body_start",bphys,"body_pages",body_pages,
          "refs",len(refs),"citations",f"{sup}/{total_cite}","ref_page",rphys,"ack_page",aphys,"ack_top",round(top,1))

if __name__=="__main__":main()
