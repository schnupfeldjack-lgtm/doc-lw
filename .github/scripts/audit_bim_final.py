# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
import re
ROOT=Path(__file__).resolve().parents[2]
DOC=ROOT/"新建文件夹/新建文件夹/BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例/BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例_炎黄模板严格版_图片已补全.docx"
d=Document(DOC)
ps=d.paragraphs
anchors={}
for i,p in enumerate(ps):
    t=p.text.strip().replace(" ","")
    if t.startswith("1绪论") and "body" not in anchors: anchors["body"]=i
    if t in ("结论","结论与展望","结论与展望") or t.startswith("结论"): anchors.setdefault("conclusion",i)
    if t=="参考文献": anchors.setdefault("ref",i)
    if t=="致谢": anchors.setdefault("ack",i)
print("paragraphs",len(ps),"sections",len(d.sections),"anchors",anchors)
for k,v in anchors.items():
    print("\nANCHOR",k,v)
    for j in range(max(0,v-8),min(len(ps),v+8)):
        p=ps[j]
        print(j,repr(p.text), "style=",p.style.name if p.style else None, "pageBreakBefore=",p.paragraph_format.page_break_before)
start=anchors.get("body",0)
end=anchors.get("conclusion",anchors.get("ref",len(ps)))
body=''.join(p.text for p in ps[start:end])
han=len(re.findall(r'[\u4e00-\u9fff]',body))
chars=len(re.sub(r'\s+','',body))
print("\nBODY han",han,"nonspace",chars)
print("inline_shapes",len(d.inline_shapes),"tables",len(d.tables))
for si,s in enumerate(d.sections):
    print("SECTION",si,"start",s.start_type,"header_distance",s.header_distance,"footer_distance",s.footer_distance,
          "top",s.top_margin,"bottom",s.bottom_margin,"left",s.left_margin,"right",s.right_margin,
          "header_link",s.header.is_linked_to_previous,"footer_link",s.footer.is_linked_to_previous)
    print(" header text:",[p.text for p in s.header.paragraphs])
