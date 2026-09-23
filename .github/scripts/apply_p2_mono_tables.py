# -*- coding: utf-8 -*-
from pathlib import Path
import shutil, re
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
PAPER=ROOT/"新建文件夹/新建文件夹/装配式施工质量管理问题及优化研究——以市政项目为例"
SRC=PAPER/"装配式施工质量管理问题及优化研究——以市政项目为例_最终优化版.docx"
OUT=PAPER/"装配式施工质量管理问题及优化研究——以市政项目为例_黑白图表终版.docx"
FIG=ROOT/"新建文件夹/新建文件夹/.workbuddy-ai/work/figs_p2_mono"

def norm(s):return re.sub(r"\s+","",s).replace("–","-").replace("—","-")

def set_eastasia(run,font="宋体",size=9.5,bold=False):
    run.font.name=font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),font)
    run.font.size=Pt(size);run.bold=bold

def replace_picture(doc,prefix,img):
    cap=next((p for p in doc.paragraphs if norm(p.text).startswith(norm(prefix))),None)
    if cap is None:raise RuntimeError("找不到图题 "+prefix)
    prev=cap._p.getprevious()
    if prev is not None and prev.tag==qn("w:p") and prev.xpath(".//w:drawing"):
        prev.getparent().remove(prev)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(img),width=Cm(12.5))
    cap._p.addprevious(p._p)
    cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent=Pt(0);cap.paragraph_format.space_before=Pt(0);cap.paragraph_format.space_after=Pt(6)
    for r in cap.runs:
        if r.text:set_eastasia(r,"宋体",10.5,False)

def borders(table):
    pr=table._tbl.tblPr
    bd=pr.first_child_found_in("w:tblBorders")
    if bd is None:bd=OxmlElement("w:tblBorders");pr.append(bd)
    # 纯黑白，不填色；外框稍粗、内部细线
    for edge,sz in (("top","10"),("bottom","10"),("left","6"),("right","6"),("insideH","4"),("insideV","4")):
        el=bd.find(qn("w:"+edge))
        if el is None:el=OxmlElement("w:"+edge);bd.append(el)
        el.set(qn("w:val"),"single");el.set(qn("w:sz"),sz);el.set(qn("w:space"),"0");el.set(qn("w:color"),"000000")

def clear_shading(cell):
    pr=cell._tc.get_or_add_tcPr()
    shd=pr.find(qn("w:shd"))
    if shd is not None:pr.remove(shd)

def margins(cell):
    pr=cell._tc.get_or_add_tcPr(); mar=pr.first_child_found_in("w:tcMar")
    if mar is None:mar=OxmlElement("w:tcMar");pr.append(mar)
    for k,v in (("top",70),("start",90),("bottom",70),("end",90)):
        e=mar.find(qn("w:"+k))
        if e is None:e=OxmlElement("w:"+k);mar.append(e)
        e.set(qn("w:w"),str(v));e.set(qn("w:type"),"dxa")

def no_split(row):
    pr=row._tr.get_or_add_trPr(); e=OxmlElement("w:cantSplit");pr.append(e)

def repeat_header(row):
    pr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");pr.append(e)

def style_table(t,widths):
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;borders(t);repeat_header(t.rows[0])
    for ri,row in enumerate(t.rows):
        no_split(row)
        for ci,cell in enumerate(row.cells):
            clear_shading(cell);margins(cell);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width=Cm(widths[ci])
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or ci==0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    if r.text:set_eastasia(r,"宋体",9.2,ri==0)

def caption_before(doc,title,anchor):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True
    r=p.add_run(title);set_eastasia(r,"宋体",10.5,False)
    anchor.addprevious(p._p)

def add_table(doc,heading,title,headers,rows,widths):
    anchor=next((p for p in doc.paragraphs if norm(p.text)==norm(heading)),None)
    if anchor is None:raise RuntimeError("找不到表格插入锚点 "+heading)
    t=doc.add_table(rows=1,cols=len(headers))
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for data in rows:
        cells=t.add_row().cells
        for i,v in enumerate(data):cells[i].text=v
    style_table(t,widths);anchor._p.addprevious(t._tbl);caption_before(doc,title,t._tbl)

def main():
    shutil.copy2(SRC,OUT);d=Document(OUT)
    for prefix,name in [("图1-1","fig1_1.png"),("图2-1","fig2_1.png"),("图3-1","fig3_1.png"),("图4-1","fig4_1.png"),("图5-1","fig5_1.png")]:
        replace_picture(d,prefix,FIG/name)
    add_table(d,"4  质量缺陷的传播机制分析","表3-1  案例工程关键风险与前置控制点",
        ["环节","主要风险","早期信号","前置控制点"],
        [["工厂端","尺寸偏差、预埋件偏位、批次性缺陷","首件复测波动、试拼不顺畅","首件确认、模具/工装复核、异常批次加严"],
         ["物流端","运输碰损、连接面与止水构造受损","棱角破损、表面裂纹、保护缺失","装车清单、运输保护、到场状态对照"],
         ["拼装端","轴线/标高偏差连续累积","接缝宽度趋势异常、收口空间缩小","统一基准、连续测量、趋势预警"],
         ["接口端","灌浆/连接内部缺陷、接缝防水失控","材料或过程记录异常、封闭前状态不确定","隐蔽前确认、必要检测、复验后封闭"],
         ["信息端","编码不一致、资料与实体无法对应","异常历史查不到、责任边界模糊","统一编码、证据包、问题闭环"]],[2.0,3.8,3.8,4.0])
    add_table(d,"6  优化方案的实施评价与研究边界","表5-1  G0—G5质量门检查要点与放行条件",
        ["质量门","核心检查点","必备证据","放行条件"],
        [["G0 设计冻结门","图纸版本、构件编号、接口条件","冻结清单、会审/确认记录","生产依赖信息一致且无关键待定项"],
         ["G1 工厂放行门","首件、尺寸、预埋件、连接面、标识","首件记录、检验记录、出厂资料","构件达到可交付状态"],
         ["G2 进场验收门","身份对应、运输损伤、保护状态","到场验收单、关键影像","运输后状态未发生影响安装的变化"],
         ["G3 拼装几何门","轴线、标高、接缝、累计趋势","连续测量与复测记录","当前偏差与累计趋势均受控"],
         ["G4 隐蔽接口门","连接、防水、基层、过程状态","隐蔽验收、旁站/检测记录","证据充分且异常已复验关闭"],
         ["G5 数据闭环门","实体位置、问题整改、资料归档","闭环单、归档清单、构件身份链","实体与质量信息均形成闭环"]],[2.3,3.7,3.6,4.0])
    add_table(d,"6．2  实施顺序：先做一个标准段，再扩展到全线","表6-1  优化方案实施评价指标",
        ["评价维度","建议指标","指标含义","数据来源"],
        [["前序拦截","前序关闭问题占比","问题是否在进入下一阶段前被发现并关闭","质量门/问题台账"],
         ["几何稳定","二次调整构件数、偏差趋势","G3对累计偏差的控制程度","连续测量记录"],
         ["接口质量","隐蔽前整改与复验情况","连接、防水问题是否在封闭前处理","隐蔽验收/检测记录"],
         ["闭环效率","问题关闭时长、逾期数","异常从登记到完成复验的处理效率","问题闭环台账"],
         ["追溯完整","随机构件资料可调取率","设计—生产—交付—安装—整改记录是否贯通","构件档案/编码系统"]],[2.1,3.3,5.0,3.2])
    d.save(OUT)
    d2=Document(OUT)
    print("MONO_APPLY_OK","pics",len(d2.inline_shapes),"tables",len(d2.tables),"size",OUT.stat().st_size)

if __name__=="__main__":main()
