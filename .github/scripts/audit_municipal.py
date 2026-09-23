# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
import re
ROOT=Path(__file__).resolve().parents[2]
P=ROOT/"新建文件夹/新建文件夹/装配式施工质量管理问题及优化研究——以市政项目为例/装配式施工质量管理问题及优化研究——以市政项目为例_最终优化版.docx"
d=Document(P)
ps=d.paragraphs
print("sections",len(d.sections),"paras",len(ps),"inline_shapes",len(d.inline_shapes),"tables",len(d.tables))
for i,s in enumerate(d.sections):
    print("SECTION",i,"start",s.start_type,"header_link",s.header.is_linked_to_previous,"footer_link",s.footer.is_linked_to_previous,"header", [p.text for p in s.header.paragraphs])
# body/ref indices
def norm(s): return re.sub(r"\s+","",s).replace("–","-").replace("—","-")
body_idxs=[i for i,p in enumerate(ps) if norm(p.text) in ("1引言","1绪论") and not (p.style and p.style.name.startswith("toc"))]
print("body_idxs",body_idxs)
ref_idx=next((i for i,p in enumerate(ps) if norm(p.text)=="参考文献"),None)
ack_idx=next((i for i,p in enumerate(ps) if norm(p.text)=="致谢"),None)
print("ref_idx",ref_idx,"ack_idx",ack_idx)
if body_idxs and ref_idx:
    b=body_idxs[-1]
    txt=''.join(p.text for p in ps[b:ref_idx])
    han=len(re.findall(r'[\u4e00-\u9fff]',txt))
    print("body_han",han)
    cites=sorted(set(int(x) for x in re.findall(r'\[(\d+)\]',txt)))
    print("cites",cites)
# refs
if ref_idx is not None:
    print("REFERENCES")
    for i,p in enumerate(ps[ref_idx+1:ack_idx if ack_idx is not None else len(ps)], start=ref_idx+1):
        t=p.text.strip()
        if t:
            print(i,repr(t))
