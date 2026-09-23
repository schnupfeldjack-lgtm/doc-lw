# -*- coding: utf-8 -*-
from pathlib import Path
import re, subprocess, sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def norm(s): return re.sub(r"\s+","",s).replace("–","-").replace("—","-")

def pdf_pages(pdf):
    info=subprocess.check_output(["pdfinfo",str(pdf)],text=True,errors="ignore")
    total=int(re.search(r"^Pages:\s+(\d+)",info,re.M).group(1))
    out=[]
    for i in range(1,total+1):
        out.append(subprocess.check_output(["pdftotext","-f",str(i),"-l",str(i),"-layout",str(pdf),"-"],text=True,errors="ignore"))
    return out

def body_idx(d):
    xs=[i for i,p in enumerate(d.paragraphs) if norm(p.text)=="1引言" and not (p.style and p.style.name.lower().startswith("toc"))]
    return xs[-1]

def level(t):
    n=norm(t)
    if re.match(r"^\d+\.\d+\.\d+",n):return 3
    if re.match(r"^\d+[．.]\d+",n):return 2
    if re.match(r"^\d+",n):return 1
    return None

def replace_last_number(p,num):
    nodes=p._p.xpath(".//w:t")
    for n in reversed(nodes):
        txt=n.text or ""
        if re.search(r"\d+\s*$",txt):
            n.text=re.sub(r"\d+\s*$",str(num),txt);return
    if nodes: nodes[-1].text=(nodes[-1].text or "")+str(num)

def add_text(p,text):
    r=OxmlElement("w:r");rp=OxmlElement("w:rPr")
    fs=OxmlElement("w:rFonts")
    for a in ("w:ascii","w:hAnsi","w:eastAsia"):fs.set(qn(a),"宋体")
    rp.append(fs);rp.append(OxmlElement("w:b"))
    sz=OxmlElement("w:sz");sz.set(qn("w:val"),"21");rp.append(sz)
    sz2=OxmlElement("w:szCs");sz2.set(qn("w:val"),"21");rp.append(sz2)
    r.append(rp);t=OxmlElement("w:t");t.text=text;r.append(t);p.append(r)

def add_page_field(p):
    r=OxmlElement("w:r");f=OxmlElement("w:fldChar");f.set(qn("w:fldCharType"),"begin");r.append(f);p.append(r)
    r=OxmlElement("w:r");it=OxmlElement("w:instrText");it.set("{http://www.w3.org/XML/1998/namespace}space","preserve");it.text=" PAGE ";r.append(it);p.append(r)
    r=OxmlElement("w:r");f=OxmlElement("w:fldChar");f.set(qn("w:fldCharType"),"separate");r.append(f);p.append(r)
    add_text(p,"1")
    r=OxmlElement("w:r");f=OxmlElement("w:fldChar");f.set(qn("w:fldCharType"),"end");r.append(f);p.append(r)

def rebuild_header(h,total):
    root=h._element
    for c in list(root):root.remove(c)
    p=OxmlElement("w:p");pp=OxmlElement("w:pPr")
    jc=OxmlElement("w:jc");jc.set(qn("w:val"),"left");pp.append(jc)
    tabs=OxmlElement("w:tabs");tb=OxmlElement("w:tab");tb.set(qn("w:val"),"right");tb.set(qn("w:pos"),"8760");tabs.append(tb);pp.append(tabs)
    bdr=OxmlElement("w:pBdr");bt=OxmlElement("w:bottom")
    for k,v in (("w:val","single"),("w:sz","6"),("w:space","1"),("w:color","auto")):bt.set(qn(k),v)
    bdr.append(bt);pp.append(bdr)
    sp=OxmlElement("w:spacing");sp.set(qn("w:before"),"0");sp.set(qn("w:after"),"0");pp.append(sp)
    p.append(pp)
    add_text(p,"炎黄职业技术学院毕业论文")
    r=OxmlElement("w:r");r.append(OxmlElement("w:tab"));p.append(r)
    add_text(p,"第 ");add_page_field(p);add_text(p,f" 页  共 {total} 页")
    root.append(p)

def main():
    docx=Path(sys.argv[1]);pdf=Path(sys.argv[2])
    d=Document(docx); pages=pdf_pages(pdf); np=[norm(x) for x in pages]
    bphys=next((i for i,t in enumerate(np,1) if "1引言" in t and "市政综合管廊具有线路长" in t),None)
    if bphys is None: raise RuntimeError("找不到正文第1页")
    bi=body_idx(d)
    ti=next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="目录")
    ri=next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="参考文献")
    ai=next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="致谢")
    ci=next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="结论")
    # 重排目录末尾特殊条目为：结论、参考文献、致谢
    toc_special={}
    for p in d.paragraphs[ti+1:bi]:
        n=norm(p.text)
        for k in ("结论","参考文献","致谢"):
            if n.startswith(k):toc_special[k]=p
    if len(toc_special)==3:
        normal=[p for p in d.paragraphs[ti+1:bi] if all(p._element is not x._element for x in toc_special.values()) and p.text.strip()]
        anchor=normal[-1]
        cur=anchor._element
        for k in ("结论","参考文献","致谢"):
            el=toc_special[k]._element
            cur.addnext(el);cur=el
    # 正文一、二级标题 + 特殊页映射
    headings=[]
    for p in d.paragraphs[bi:ri]:
        if level(p.text) in (1,2):headings.append((norm(p.text),p.text.strip()))
    headings += [("结论","结  论"),("参考文献","参 考 文 献"),("致谢","致  谢")]
    mapping={}
    for key,raw in headings:
        for phys in range(bphys,len(pages)+1):
            if key in np[phys-1]:
                mapping[key]=phys-bphys+1;break
    # 更新目录缓存页码
    for p in d.paragraphs[ti+1:bi]:
        n=norm(p.text)
        matches=[(k,v) for k,v in mapping.items() if n.startswith(k)]
        if matches:
            k,v=max(matches,key=lambda kv:len(kv[0]));replace_last_number(p,v)
    body_pages=len(pages)-bphys+1
    rebuild_header(d.sections[-1].header,body_pages)
    d.save(docx)
    print("TOC_HEADER_OK","physical_total",len(pages),"body_physical",bphys,"body_pages",body_pages,"mapping",mapping)

if __name__=="__main__":main()
