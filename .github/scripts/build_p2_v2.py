# -*- coding: utf-8 -*-
from pathlib import Path
import shutil, sys
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "新建文件夹/新建文件夹/.workbuddy-ai/work"
sys.path.insert(0, str(WORK))
from fill_docx import fill

BASE = ROOT / "新建文件夹/新建文件夹"
PAPER = BASE / "装配式施工质量管理问题及优化研究——以市政项目为例"
TPL = BASE / "炎黄职业技术学院毕业论文模板(1).docx"
OUT = PAPER / "装配式施工质量管理问题及优化研究——以市政项目为例_最终优化版.docx"
TITLE = "装配式施工质量管理问题及优化研究——以市政项目为例"

FIG_DIR = WORK / "figs_p2_v2"
FIGURES = {
    "fig1_1": str(FIG_DIR / "fig1_1.png"),
    "fig2_1": str(FIG_DIR / "fig2_1.png"),
    "fig3_1": str(FIG_DIR / "fig3_1.png"),
    "fig4_1": str(FIG_DIR / "fig4_1.png"),
    "fig5_1": str(FIG_DIR / "fig5_1.png"),
}

def replace_picture_before_caption(doc, caption_prefix, img_path):
    cap = None
    for p in doc.paragraphs:
        if p.text.strip().replace(" ", "").startswith(caption_prefix.replace(" ", "")):
            cap = p
            break
    if cap is None:
        raise RuntimeError("找不到图题: " + caption_prefix)
    prev = cap._p.getprevious()
    if prev is not None and prev.tag.endswith("}p") and prev.xpath(".//w:drawing"):
        prev.getparent().remove(prev)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.add_run().add_picture(str(img_path), width=Cm(13.0))
    cap._p.addprevious(p._p)

def main():
    PAPER.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TPL, OUT)
    fill(
        str(PAPER),
        OUT.name,
        TITLE,
        [str(WORK/"p2_v2_part1.md"), str(WORK/"p2_v2_part2.md"), str(WORK/"p2_v2_part3.md")],
        str(WORK/"p2_v2_meta.md"),
        FIGURES,
        target_chars_per_page=1050,
    )
    d = Document(OUT)
    for prefix, name in [
        ("图1–1", "fig1_1.png"),
        ("图2–1", "fig2_1.png"),
        ("图3–1", "fig3_1.png"),
        ("图4–1", "fig4_1.png"),
        ("图5–1", "fig5_1.png"),
    ]:
        replace_picture_before_caption(d, prefix, FIG_DIR/name)
    d.save(OUT)
    print("BUILT", OUT, OUT.stat().st_size, "pictures", len(d.inline_shapes))

if __name__ == "__main__":
    main()
