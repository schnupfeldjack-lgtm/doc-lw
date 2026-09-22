# -*- coding: utf-8 -*-
"""
重建炎黄模板严格版并只修复正文两张图片。
母版：原始 _按模板排版.docx
目标：
- 开题报告/任务书/中期检查/封面各自独立一页
- 中文摘要、英文摘要、目录、正文分页保持已确认的严格版结构
- 正文页眉严格按炎黄模板
- 保留此前内容优化/章节整理
- 图3-1、图4-1使用已提取的原论文图片，修复错误的图片关系
"""
from pathlib import Path
import copy
import re
import shutil
import tempfile
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
BASE_DIR = ROOT / "新建文件夹/新建文件夹"
PAPER_DIR = BASE_DIR / "BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例"
SRC = PAPER_DIR / "BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例_按模板排版.docx"
TPL = BASE_DIR / "炎黄职业技术学院毕业论文模板(1).docx"
FIG3 = BASE_DIR / ".workbuddy-ai/work/figs_p3/fig3_1.png"
FIG4 = BASE_DIR / ".workbuddy-ai/work/figs_p3/fig4_1.png"
OUT = PAPER_DIR / "BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例_炎黄模板严格版_图片已补全.docx"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"

def text_nodes(p):
    return p._p.xpath(".//w:t")

def ptext(p):
    return "".join(n.text or "" for n in text_nodes(p)).strip()

def norm(s):
    return re.sub(r"\s+", "", s).replace("–", "-").replace("—", "-")

def set_para_text(p, new_text, start_node=0):
    nodes = text_nodes(p)
    if not nodes:
        p.add_run(new_text)
        return
    nodes[start_node].text = new_text
    for n in nodes[start_node + 1:]:
        n.text = ""

def set_abstract_cn(p, body):
    nodes = text_nodes(p)
    if len(nodes) >= 4:
        nodes[0].text = "摘"
        nodes[1].text = " "
        nodes[2].text = "要："
        nodes[3].text = body
        for n in nodes[4:]:
            n.text = ""
    else:
        set_para_text(p, "摘 要：" + body)

def set_abstract_en(p, body):
    nodes = text_nodes(p)
    if len(nodes) >= 2:
        nodes[0].text = "Abstract："
        nodes[1].text = body
        for n in nodes[2:]:
            n.text = ""
    else:
        set_para_text(p, "Abstract：" + body)

def move_block_after(anchor_p, block_ps):
    anchor = anchor_p._p
    for bp in block_ps:
        el = bp._p
        anchor.addnext(el)
        anchor = el

def delete_paragraph(p):
    if p is None or p._element is None:
        return
    parent = p._element.getparent()
    if parent is not None:
        parent.remove(p._element)

def optimize_content(src, dst):
    doc = Document(src)
    P = {i: p for i, p in enumerate(doc.paragraphs)}

    cn = (
        "BIM协同设计是提升装配式建筑结构设计质量的重要技术手段。本文以上海闵行浦江镇召楼路以东S8-01市属保障房项目为案例，"
        "围绕装配整体式剪力墙体系和装配式建造特点，从设计协同、构件深化、管线综合和设计变更四个维度分析BIM协同设计对结构设计质量的作用机制，"
        "并结合参数化构件库、多专业模型整合、碰撞检查、节点优化和设计—生产数据贯通等应用环节进行案例分析。研究认为，BIM协同设计有助于前移设计问题发现关口，"
        "降低构件尺寸、预埋件定位及专业碰撞等问题的发生风险，提升设计成果的一致性和可实施性。针对协同深度不足、数据传递不连续、人员能力不均衡等问题，"
        "提出加强全专业协同、推进设计—生产一体化、完善实施标准与人员培训等优化建议。"
    )
    en = (
        "BIM-based collaborative design is an important approach to improving structural design quality in prefabricated buildings. "
        "This paper takes the S8-01 municipal affordable housing project east of Zhaolou Road in Pujiang Town, Minhang District, Shanghai, as a case study. "
        "Focusing on the assembled integral shear-wall system and the characteristics of prefabricated construction, the study analyzes how BIM collaborative design affects structural design quality "
        "from four dimensions: design coordination, component detailing, MEP coordination, and design change management. "
        "The analysis is linked to practical applications including parametric component libraries, multidisciplinary model integration, clash detection, joint optimization, and design-to-production data exchange. "
        "The findings indicate that BIM collaborative design can move problem detection to earlier design stages, reduce the risk of dimensional errors, inaccurate embedded-part positioning, and cross-disciplinary clashes, "
        "and improve the consistency and constructability of design deliverables. In response to insufficient collaboration depth, discontinuous data transfer, and uneven staff capabilities, "
        "the paper proposes stronger multidisciplinary coordination, closer design-production integration, clearer implementation standards, and targeted training."
    )
    set_abstract_cn(P[25], cn)
    set_abstract_en(P[30], en)

    replacements = {
        71: "上海闵行浦江镇基地召楼路以东S8-01市属保障房项目是国家“十三五”课题示范项目、住建部装配式科技示范项目和上海市装配式建筑示范项目。项目总建筑面积约11万平方米，采用装配整体式剪力墙体系，装配式建筑比例达到100%，单体预制率处于较高水平[4]。该项目具备规模大、标准化程度高和装配式应用场景较完整等特点，为分析BIM协同设计与结构设计质量之间的关系提供了具有代表性的案例。本文据此从结构设计质量视角开展分析。",
        73: "本研究的目的有三个层面。第一，梳理BIM协同设计影响装配式结构设计质量的作用机理，明确BIM在不同设计环节中的作用路径。第二，基于S8-01项目公开资料，对BIM协同设计在大型装配式住宅项目中的应用表现进行案例分析。第三，结合分析中发现的协同、数据和人员等问题，提出具有可操作性的优化对策。",
        76: "国外对BIM相关理论和工程应用的研究起步较早，进入21世纪后，BIM在欧美建筑行业的设计协同、施工管理和信息交付等场景中持续推广[5]。在装配式建筑领域，相关研究主要集中在设计协同、构件深化、预制生产管理和施工模拟等方面。有研究认为，BIM与装配式建造相结合，有利于打通设计、生产与施工之间的信息链路，是建筑工业化的重要技术支撑[6]。",
        80: "本文主要开展四部分工作。第一，梳理BIM技术与协同设计的相关理论，明确BIM的内涵、协同机制和结构设计质量的评价要点。第二，介绍S8-01项目的基本情况、结构体系和BIM协同设计实施方案。第三，从设计协同、构件深化、管线综合和设计变更四个维度分析BIM协同设计影响结构设计质量的作用机制。第四，基于公开项目资料开展案例分析，并针对应用中存在的问题提出优化对策。",
        81: "技术路线上，本文采用文献分析法与案例分析法相结合的方法。文献分析用于确定分析维度和评价依据；案例分析以S8-01项目公开资料为主要依据，对BIM在协同设计、构件深化、管线综合和变更管理等环节的应用进行归纳，并结合相关研究进行对照分析。",
        107: "项目于2017年1月开工，计划于2018年9月竣工[4]。根据公开资料，项目BIM应用主要覆盖深化设计、碰撞检查和节点优化等环节，这些应用为设计问题的提前识别和专业协调提供了技术支持。",
        153: "在尺寸精度方面，参数化构件库和BIM模型的三维可视化使设计师能够直观检查构件尺寸及其空间关系，有助于减少二维图纸中尺寸标注不一致和构件冲突等问题。从S8-01项目公开的BIM应用资料看，深化设计环节通过参数化建模和模型校核加强了预制构件尺寸控制，为降低尺寸偏差风险提供了技术条件。",
        166: "从S8-01项目公开资料所反映的实施过程看，BIM协同设计能够将部分专业冲突和设计问题前移至设计阶段处理，从而降低施工阶段出现设计变更和图纸不一致的风险。由于公开资料未披露完整的变更次数、错误率及一次成品率等原始统计数据，本文不对改善幅度作定量判断，而是从作用机制和应用过程对其质量改善作用进行分析。",
        174: "S8-01项目具有较完整的装配式住宅BIM应用场景。鉴于公开资料未提供完整的项目内部质量统计数据，本节以项目公开应用信息为基础，对BIM协同设计在设计问题前置识别、构件深化、协同效率和设计—生产衔接等方面的作用进行案例分析，不对缺乏原始数据支撑的效果作数值化推断。",
        175: "设计问题前置识别。通过多专业模型整合和碰撞检查，可以在设计阶段识别结构构件、机电管线、预留预埋和连接节点之间的冲突，并在构件生产或现场施工前完成协调。对装配式住宅而言，这种前置处理机制有助于降低因设计问题引发的现场开洞、返工和临时变更风险。",
        176: "预制构件深化质量控制。参数化构件库、精细化节点建模和设计—生产数据衔接，使构件尺寸、预埋件位置和连接节点等信息能够在模型中进行统一表达与复核，有利于减少信息传递过程中的遗漏和二次录入错误，提高深化设计成果的完整性和可实施性。",
        177: "设计过程衔接更加可控。BIM协同设计通常需要在前期投入更多建模和协调工作，但通过提前解决专业冲突，可以减少后续施工阶段的反复核对和临时协调。其价值应从项目全周期观察，而不能仅以设计阶段的直接工作量衡量。",
        178: "多专业协同效率得到改善。统一模型环境使建筑、结构、机电等专业可以围绕同一空间关系开展校核和沟通，减少传统二维图纸往返传递造成的信息滞后。在管线密集空间和节点复杂部位，这种可视化协同方式尤其有利于快速定位问题和形成一致的调整方案。",
        179: "需要指出的是，上述作用的实现程度受到项目团队BIM应用能力、软件平台成熟度、协同规则和模型质量等多种因素影响。同样的技术路径在不同项目中可能产生不同效果，因此本文将其作为案例经验进行分析，不作简单外推。",
        180: "从适配性角度看，保障房项目通常具有建设规模大、户型标准化程度高、成本控制要求严格等特点，这些特点与BIM参数化和标准化复用具有较高契合度。户型标准化有利于提高构件库和标准模型的复用效率；设计阶段提前识别问题，也有助于减少后续变更和返工带来的不确定性。",
        181: "同时需要看到，BIM投入主要发生在设计与建模阶段，而部分效益在施工和运维阶段才逐步体现。因而对BIM应用经济性的评价应采用全周期视角，并结合建模投入、变更减少、返工减少和运维使用等指标进行核算。S8-01项目公开资料未披露完整的投入产出数据，本文不对其经济收益作定量评价。",
        189: "尽管BIM协同设计具备上述质量改善作用，但在实际应用过程中仍然存在一些限制因素，主要表现在以下几个方面。",
        207: "综合上述各维度的分析，可以对BIM协同设计在S8-01项目中的应用成效作出定性判断。从设计质量角度看，BIM协同设计通过模型整合、碰撞检查和参数化深化，有助于减少专业冲突、提高深化设计准确性并降低变更风险；从管理效率角度看，统一的模型环境和数据贯通机制有助于提高信息传递效率、减少重复性工作。",
        208: "同时也应看到，BIM应用成效的充分发挥依赖团队能力、标准规范、软件工具和管理机制等条件。由于本文主要依据公开资料开展案例分析，缺乏项目内部的完整统计数据，因此结论主要反映作用路径与实践特征，后续仍需通过可量化指标进行验证。",
        223: "第一，BIM协同设计是提升装配式结构设计质量的重要技术手段。通过参数化构件库、多专业模型整合、碰撞检查和设计—生产数据贯通等环节的应用，能够在设计阶段提前识别和协调专业冲突及细部设计问题，有助于提升设计成果的一致性和一次交付质量。",
        232: "随着BIM技术和装配式建造持续发展，BIM协同设计仍有较大的深化空间。后续研究可进一步结合真实项目过程数据，对设计质量改善程度、协同效率和全周期经济性进行量化验证。",
    }
    for i, txt in replacements.items():
        set_para_text(P[i], txt)

    # 第5章口径
    set_para_text(P[172], "5  案例分析与优化对策")
    set_para_text(P[173], "5．1  案例分析")
    set_para_text(P[58], "5  案例分析与优化对策")
    set_para_text(P[59], "5．1  案例分析")

    # 异常三级标题归位与重编号
    heading_changes = {
        83: "2.1.1  BIM技术的常用软件与数据标准",
        97: "2.2.1  协同设计的常用工作模式",
        113: "3.2.1  结构专业BIM模型的构建要点",
        126: "3.3.1  BIM协同设计的实施流程",
        148: "4.1.1  协同环境下的专业配合机制",
        144: "4.2.1  深化设计的主要内容与技术要点",
        159: "4.4.1  设计变更的管理与控制",
        182: "5.1.1  BIM应用效果的量化评价指标",
        185: "5.1.2  BIM应用与项目各阶段的衔接",
        206: "5.1.3  应用成效的综合判断",
        198: "5.3.5  完善组织保障机制",
        202: "5.3.6  强化BIM应用与项目特点的适配",
    }
    for i, txt in heading_changes.items():
        set_para_text(P[i], txt)

    # 原有段落对象直接移动，不重建其格式
    move_block_after(P[89], [P[83], P[84], P[85]])
    move_block_after(P[93], [P[97], P[98], P[99]])
    move_block_after(P[143], [P[148], P[149], P[150]])
    move_block_after(P[151], [P[144], P[145], P[146], P[147]])
    move_block_after(P[166], [P[159], P[160], P[161], P[162]])
    move_block_after(P[187], [P[206], P[207], P[208]])
    move_block_after(P[218], [P[198], P[199], P[200], P[201], P[202], P[203], P[204], P[205]])
    move_block_after(P[230], [P[228], P[229]])

    # 删除多余三级标题和重复独立结论
    delete_paragraph(P[227])
    for i in [233, 234, 235, 236, 237, 238, 239]:
        delete_paragraph(P[i])

    doc.save(dst)

def update_toc_cache(doc):
    # 严格版最终页码（与此前逐页渲染结果一致）
    pages = {
        "1绪论": 9,
        "1．1研究背景": 9,
        "1．2研究目的与意义": 9,
        "1．3国内外研究现状": 10,
        "1．4研究内容与技术路线": 10,
        "2BIM协同设计与结构设计质量相关理论": 11,
        "2．1BIM技术的基本概念": 11,
        "2．2协同设计的内涵": 12,
        "2．3结构设计质量评价": 13,
        "2．4相关标准与规定": 13,
        "3闵行浦江镇S8-01保障房项目概况与BIM应用方案": 14,
        "3．1项目基本情况": 14,
        "3．2结构体系": 14,
        "3．3BIM协同设计实施方案": 15,
        "3．4碰撞检查与节点优化": 18,
        "3．5BIM模型精度与交付要求": 19,
        "4BIM协同设计对结构设计质量的影响分析": 19,
        "4．1设计协同对设计质量的影响机理": 19,
        "4．2对结构构件深化设计质量的影响": 20,
        "4．3对管线综合与净空优化的影响": 21,
        "4．4对设计变更与设计错误率的影响": 21,
        "5案例分析与优化对策": 23,
        "5．1案例分析": 23,
        "5．2主要问题": 25,
        "5．3优化对策": 26,
        "6结论与展望": 28,
        "6．1主要结论": 28,
        "6．2展望": 29,
        "参考文献": 29,
    }
    ps = doc.paragraphs
    # 仅处理目录区域：模板/成稿中目录在正文前
    for i in range(min(35, len(ps)), min(70, len(ps))):
        p = ps[i]
        raw = ptext(p)
        key = norm(raw)
        target = None
        for prefix, page in pages.items():
            if key.startswith(prefix):
                target = page
                break
        if target is None:
            continue
        nodes = text_nodes(p)
        # 优先替换最后一个纯数字文本节点（静态目录页码）
        replaced = False
        for n in reversed(nodes):
            if (n.text or "").strip().isdigit():
                n.text = str(target)
                replaced = True
                break
        if not replaced and nodes:
            # 若页码与标题在同一节点，替换末尾数字
            old = nodes[-1].text or ""
            nodes[-1].text = re.sub(r"\d+\s*$", str(target), old)

def make_pagebreak_p():
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p

def is_pagebreak_p(el):
    return el is not None and el.tag == qn("w:p") and bool(el.xpath(".//w:br[@w:type='page']"))

def ensure_pb_after(el):
    nxt = el.getnext()
    if not is_pagebreak_p(nxt):
        el.addnext(make_pagebreak_p())

def remove_pb_after(el):
    nxt = el.getnext()
    if is_pagebreak_p(nxt):
        nxt.getparent().remove(nxt)

def table_text(tbl_el):
    return "".join(tbl_el.xpath(".//w:t/text()")).replace(" ", "").replace("\u3000", "")

def paragraph_text_el(p_el):
    return "".join(p_el.xpath(".//w:t/text()")).replace(" ", "").replace("\u3000", "")

def rebuild_strict_layout(src, template, dst):
    doc = Document(src)
    td = Document(template)
    update_toc_cache(doc)

    # 英文摘要独立一页
    for p in doc.paragraphs:
        if ptext(p).startswith("Abstract"):
            p.paragraph_format.page_break_before = True
            break

    body = doc._element.body

    # 1 开题报告独立页：首个大表后分页
    first_tbl = next((e for e in body.iterchildren() if e.tag == qn("w:tbl")), None)
    if first_tbl is not None:
        ensure_pb_after(first_tbl)

    # 2 任务书独立页：任务书日期后分页
    for el in body.iterchildren():
        if el.tag == qn("w:p") and paragraph_text_el(el) == "2025年9月10日":
            ensure_pb_after(el)
            break

    # 避免任务书表后额外分页造成空白页
    for el in body.iterchildren():
        if el.tag == qn("w:tbl") and "本课题研究进度安排" in table_text(el):
            remove_pb_after(el)
            break

    # 3 中期检查独立页：包含第四阶段与签名的中期表后分页
    mid_tbl = None
    for el in body.iterchildren():
        if el.tag == qn("w:tbl"):
            tt = table_text(el)
            if "第一阶段" in tt and "第四阶段" in tt and "指导老师签名" in tt:
                mid_tbl = el
                break
    if mid_tbl is not None:
        ensure_pb_after(mid_tbl)
    # 与此前严格修正版一致：第三个顶层表就是中期检查表，再做一次确定性分页
    if len(doc.tables) > 2:
        ensure_pb_after(doc.tables[2]._element)

    # 4 把封面节的 sectPr 放到封面日期之后，避免节属性把封面挤坏
    cover_sect = None
    for p in doc.paragraphs:
        if "专业技术职务" in ptext(p):
            pPr = p._p.find(qn("w:pPr"))
            if pPr is not None:
                sp = pPr.find(qn("w:sectPr"))
                if sp is not None:
                    cover_sect = copy.deepcopy(sp)
                    pPr.remove(sp)
                    break

    if cover_sect is not None:
        cover_end = None
        for el in body.iterchildren():
            if el.tag == qn("w:tbl") and table_text(el) == "2025年9月":
                cover_end = el
                break
        if cover_end is not None:
            # 若下一段已经带封面sectPr则不重复；否则插入空段承载
            nxt = cover_end.getnext()
            found = False
            if nxt is not None and nxt.tag == qn("w:p"):
                pp = nxt.find(qn("w:pPr"))
                found = pp is not None and pp.find(qn("w:sectPr")) is not None
            if not found:
                p = OxmlElement("w:p")
                pPr = OxmlElement("w:pPr")
                pPr.append(cover_sect)
                p.append(pPr)
                cover_end.addnext(p)

    # 5 正文页边距/页眉距离沿用模板正文页
    if doc.sections and td.sections:
        cs = doc.sections[-1]
        ts = td.sections[-1]
        for attr in ("header_distance", "footer_distance", "top_margin", "bottom_margin", "left_margin", "right_margin"):
            try:
                setattr(cs, attr, getattr(ts, attr))
            except Exception:
                pass

        # 仅正文最后一节使用严格页眉；前置页保持无正文页眉
        hdr = cs.header
        hdr.is_linked_to_previous = False
        root = hdr._element
        for child in list(root):
            root.remove(child)

        p = OxmlElement("w:p")
        pPr = OxmlElement("w:pPr")

        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "left")
        pPr.append(jc)

        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:pos"), "8760")
        tabs.append(tab)
        pPr.append(tabs)

        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "auto")
        bdr.append(bottom)
        pPr.append(bdr)

        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "0")
        ind.set(qn("w:right"), "0")
        ind.set(qn("w:firstLine"), "0")
        pPr.append(ind)

        sp = OxmlElement("w:spacing")
        sp.set(qn("w:before"), "0")
        sp.set(qn("w:after"), "0")
        pPr.append(sp)
        p.append(pPr)

        def add_text(text):
            r = OxmlElement("w:r")
            rp = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), "宋体")
            fonts.set(qn("w:hAnsi"), "宋体")
            fonts.set(qn("w:eastAsia"), "宋体")
            rp.append(fonts)
            bold = OxmlElement("w:b")
            rp.append(bold)
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), "21")
            rp.append(sz)
            szc = OxmlElement("w:szCs")
            szc.set(qn("w:val"), "21")
            rp.append(szc)
            r.append(rp)
            t = OxmlElement("w:t")
            t.text = text
            r.append(t)
            p.append(r)

        def add_tab():
            r = OxmlElement("w:r")
            r.append(OxmlElement("w:tab"))
            p.append(r)

        def add_field(instr, cached):
            r1 = OxmlElement("w:r")
            fc1 = OxmlElement("w:fldChar")
            fc1.set(qn("w:fldCharType"), "begin")
            r1.append(fc1)
            p.append(r1)

            r2 = OxmlElement("w:r")
            ins = OxmlElement("w:instrText")
            ins.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            ins.text = f" {instr} "
            r2.append(ins)
            p.append(r2)

            r3 = OxmlElement("w:r")
            sep = OxmlElement("w:fldChar")
            sep.set(qn("w:fldCharType"), "separate")
            r3.append(sep)
            p.append(r3)

            add_text(str(cached))

            r4 = OxmlElement("w:r")
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            r4.append(end)
            p.append(r4)

        add_text("炎黄职业技术学院毕业论文")
        add_tab()
        add_text("第 ")
        add_field("PAGE", 9)
        add_text(" 页  共 ")
        add_field("NUMPAGES", 31)
        add_text(" 页")
        root.append(p)

    # 6 开启字段自动更新
    settings = doc.settings._element
    uf = settings.find(qn("w:updateFields"))
    if uf is None:
        uf = OxmlElement("w:updateFields")
        settings.append(uf)
    uf.set(qn("w:val"), "true")

    doc.save(dst)

def patch_images(docx_path):
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        with zipfile.ZipFile(docx_path) as z:
            z.extractall(td)

        docxml = td / "word/document.xml"
        relxml = td / "word/_rels/document.xml.rels"
        media = td / "word/media"
        media.mkdir(parents=True, exist_ok=True)

        dt = etree.parse(str(docxml))
        rt = etree.parse(str(relxml))
        root = dt.getroot()
        relroot = rt.getroot()

        ns = {"w": W, "r": R, "a": A}

        def el_text(el):
            return "".join(el.xpath(".//w:t/text()", namespaces=ns)).strip()

        def find_caption(prefix):
            for p in root.xpath("//w:body//w:p", namespaces=ns):
                if norm(el_text(p)).startswith(prefix):
                    return p
            raise RuntimeError("未找到图题：" + prefix)

        def prev_drawing(cap):
            cur = cap.getprevious()
            for _ in range(10):
                if cur is None:
                    break
                if cur.tag == f"{{{W}}}p" and cur.xpath(".//w:drawing", namespaces=ns):
                    return cur
                cur = cur.getprevious()
            raise RuntimeError("图题前没有图片段落：" + el_text(cap))

        nums = []
        for r in relroot:
            m = re.fullmatch(r"rId(\d+)", r.get("Id", ""))
            if m:
                nums.append(int(m.group(1)))
        next_id = max(nums or [0]) + 1

        def add_image(prefix, src_img, out_name):
            nonlocal next_id
            cap = find_caption(prefix)
            pic = prev_drawing(cap)
            rid = f"rId{next_id}"
            next_id += 1

            shutil.copy2(src_img, media / out_name)

            rel = etree.Element(f"{{{REL}}}Relationship")
            rel.set("Id", rid)
            rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
            rel.set("Target", "media/" + out_name)
            relroot.append(rel)

            blips = pic.xpath(".//a:blip", namespaces=ns)
            if not blips:
                raise RuntimeError("图片段落缺少a:blip")
            for blip in blips:
                blip.set(f"{{{R}}}embed", rid)
                link = f"{{{R}}}link"
                if link in blip.attrib:
                    del blip.attrib[link]
            return rid

        rid3 = add_image("图3-1BIM协同设计实施方案主要环节", FIG3, "fig3_1_repaired.png")
        rid4 = add_image("图4-1BIM协同设计对结构设计质量的影响机理", FIG4, "fig4_1_repaired.png")

        dt.write(str(docxml), encoding="UTF-8", xml_declaration=True, standalone="yes")
        rt.write(str(relxml), encoding="UTF-8", xml_declaration=True, standalone="yes")

        rebuilt = docx_path.with_suffix(".rebuilt.docx")
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as z:
            for p in td.rglob("*"):
                if p.is_file():
                    z.write(p, p.relative_to(td).as_posix())
        shutil.move(rebuilt, docx_path)

    # 二进制结构校验
    with zipfile.ZipFile(docx_path) as z:
        names = set(z.namelist())
        assert "word/media/fig3_1_repaired.png" in names
        assert "word/media/fig4_1_repaired.png" in names
        assert len(z.read("word/media/fig3_1_repaired.png")) > 100000
        assert len(z.read("word/media/fig4_1_repaired.png")) > 100000
        d = z.read("word/document.xml").decode("utf-8")
        r = z.read("word/_rels/document.xml.rels").decode("utf-8")
        assert f'r:embed="{rid3}"' in d and f'r:embed="{rid4}"' in d
        assert f'Id="{rid3}"' in r and "fig3_1_repaired.png" in r
        assert f'Id="{rid4}"' in r and "fig4_1_repaired.png" in r

def qa(docx_path):
    d = Document(docx_path)
    assert len(d.sections) == 4, f"节数量异常: {len(d.sections)}"
    all_text = "\n".join(ptext(p) for p in d.paragraphs)
    assert "5  案例分析与优化对策" in all_text
    assert "5．1  案例分析" in all_text
    assert "2.0.4" not in all_text
    assert "6.1.6" not in all_text

    # 英文摘要必须分页
    abs_en = next(p for p in d.paragraphs if ptext(p).startswith("Abstract"))
    assert abs_en.paragraph_format.page_break_before

    # 正文页眉
    hdr_xml = d.sections[-1].header._element.xml
    assert "炎黄职业技术学院毕业论文" in hdr_xml
    assert "PAGE" in hdr_xml and "NUMPAGES" in hdr_xml

    # 前置页分页标志检查
    body = d._element.body
    tbls = [e for e in body.iterchildren() if e.tag == qn("w:tbl")]
    assert tbls and is_pagebreak_p(tbls[0].getnext()), "开题报告后未分页"
    mid = None
    for e in tbls:
        tt = table_text(e)
        if "第一阶段" in tt and "第四阶段" in tt and "指导老师签名" in tt:
            mid = e
            break
    def pb_within(el, steps=5):
        cur = el.getnext()
        for _ in range(steps):
            if cur is None:
                return False
            if is_pagebreak_p(cur):
                return True
            cur = cur.getnext()
        return False
    assert mid is not None and pb_within(mid), "中期检查后未分页"

    print("QA_OK", docx_path, docx_path.stat().st_size, "sections", len(d.sections))

def main():
    for p in (SRC, TPL, FIG3, FIG4):
        if not p.exists():
            raise FileNotFoundError(p)
    PAPER_DIR.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        optimized = td / "optimized.docx"
        strict = td / "strict.docx"
        optimize_content(SRC, optimized)
        rebuild_strict_layout(optimized, TPL, strict)
        shutil.copy2(strict, OUT)

    patch_images(OUT)
    qa(OUT)
    print("FINAL_OUTPUT", OUT)
    print("FINAL_SIZE", OUT.stat().st_size)

if __name__ == "__main__":
    main()
