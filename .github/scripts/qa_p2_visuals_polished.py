# -*- coding: utf-8 -*-
from pathlib import Path
import re, subprocess, sys, zipfile
from docx import Document
from docx.oxml.ns import qn
import pdfplumber

def norm(s):return re.sub(r"\s+","",s).replace("–","-").replace("—","-")
def ptext(pdf,i):return subprocess.check_output(["pdftotext","-f",str(i),"-l",str(i),"-layout",str(pdf),"-"],text=True,errors="ignore")
def total_pages(pdf):
    s=subprocess.check_output(["pdfinfo",str(pdf)],text=True,errors="ignore")
    return int(re.search(r"^Pages:\s+(\d+)",s,re.M).group(1))
def body_idx(d):
    return [i for i,p in enumerate(d.paragraphs) if norm(p.text)=="1引言" and not (p.style and p.style.name.lower().startswith("toc"))][-1]
def ref_idx(d):return next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="参考文献")

def main():
    src=Path(sys.argv[1]);docx=Path(sys.argv[2]);pdf=Path(sys.argv[3])
    a=Document(src);d=Document(docx)
    # 正文段落与参考文献不改写
    src_body="\n".join(p.text for p in a.paragraphs[body_idx(a):ref_idx(a)])
    out_body="\n".join(p.text for p in d.paragraphs[body_idx(d):ref_idx(d)] if not norm(p.text).startswith(("表3-1","表5-1","表6-1")))
    # 新增表题之外，原段落应完整保留（只检查长文本序列包含）
    for p in a.paragraphs[body_idx(a):ref_idx(a)]:
        if p.text.strip() and p.text not in out_body:
            raise AssertionError("正文原段落缺失: "+p.text[:60])
    refs=[p.text.strip() for p in d.paragraphs if re.match(r"^\[\d+\]",p.text.strip())]
    assert len(refs)==15,len(refs)
    assert len(d.inline_shapes)>=5,len(d.inline_shapes)
    assert len(d.tables)>=8,len(d.tables)
    for title in ("表3-1  案例工程关键风险与前置控制点","表5-1  G0—G5质量门检查要点与放行条件","表6-1  优化方案实施评价指标"):
        assert any(norm(p.text)==norm(title) for p in d.paragraphs),title
    # 学校模板前置四页仍不计页码；正文从第1页起
    total=total_pages(pdf); texts=[ptext(pdf,i) for i in range(1,total+1)]; np=[norm(x) for x in texts]
    assert "开题报告" in np[0]
    assert "毕业论文任务书" in np[1]
    assert "中期检查" in np[2]
    assert "装配式施工质量管理问题及优化研究" in np[3]
    bphys=next(i for i,t in enumerate(np,1) if "1引言" in t and "市政综合管廊具有线路长" in t)
    assert bphys==9,bphys
    for i in range(1,bphys):
        assert not re.search(r"第\s*\d+\s*页\s*共\s*\d+\s*页",texts[i-1]),i
    body_pages=total-bphys+1
    assert f"第1页共{body_pages}页" in np[bphys-1],np[bphys-1][:120]
    assert f"第{body_pages}页共{body_pages}页" in np[-1],np[-1][:120]
    for i,t in enumerate(np,1):
        assert len(t)>20,f"第{i}页疑似空白"
    # 图片真实嵌入并足够大
    with zipfile.ZipFile(docx) as z:
        pngs=[x for x in z.namelist() if x.startswith("word/media/") and x.endswith(".png")]
        sizes=sorted([len(z.read(x)) for x in pngs],reverse=True)
        assert len(pngs)>=5,(len(pngs),pngs)
        assert sum(1 for s in sizes if s>120000)>=5,sizes[:10]
    # 参考文献和致谢仍靠页面上部
    rphys=next(i for i,t in enumerate(np[bphys-1:],bphys) if "参考文献" in t and "[1]" in t)
    aphys=next(i for i,t in enumerate(np[bphys-1:],bphys) if "致谢" in t and "指导老师" in t)
    with pdfplumber.open(pdf) as f:
        for pg,needle in ((rphys,"参考"),(aphys,"致")):
            words=f.pages[pg-1].extract_words()
            tops=[w["top"] for w in words if any(ch in w["text"] for ch in needle)]
            assert tops and min(tops)<230,(pg,tops[:5])
    print("VISUAL_QA_OK","physical_pages",total,"body_start",bphys,"body_pages",body_pages,"pics",len(d.inline_shapes),"tables",len(d.tables),"refs",len(refs))

if __name__=="__main__":main()
