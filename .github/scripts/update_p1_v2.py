# -*- coding: utf-8 -*-
"""根据第一次PDF渲染结果校准目录页码与正文页眉总页数。"""
from pathlib import Path
import copy, re, subprocess, sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def norm(s):
    return re.sub(r"\s+","",s).replace("–","-").replace("—","-")

def pdf_pages(pdf):
    info=subprocess.check_output(["pdfinfo",str(pdf)],text=True,errors="ignore")
    total=int(re.search(r"^Pages:\s+(\d+)",info,re.M).group(1))
    pages=[]
    for i in range(1,total+1):
        t=subprocess.check_output(["pdftotext","-f",str(i),"-l",str(i),"-layout",str(pdf),"-"],
                                  text=True,errors="ignore")
        pages.append(t)
    return pages

def body_index(d):
    xs=[i for i,p in enumerate(d.paragraphs)
        if norm(p.text).startswith("1绪论") and not (p.style and p.style.name.lower().startswith("toc"))]
    return xs[-1]

def ref_index(d):
    return next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="参考文献")

def level(text):
    n=norm(text)
    if re.match(r"^\d+\.\d+\.\d+",n):return 3
    if re.match(r"^\d+[．.]\d+",n):return 2
    if re.match(r"^\d+",n):return 1
    return None

def replace_last_number(p,num):
    nodes=p._p.xpath(".//w:t")
    for n in reversed(nodes):
        tx=n.text or ""
        if re.search(r"\d+\s*$",tx):
            n.text=re.sub(r"\d+\s*$",str(num),tx)
            return
    if nodes:
        nodes[-1].text=(nodes[-1].text or "")+str(num)

def add_text_run(p,text):
    r=OxmlElement("w:r");rp=OxmlElement("w:rPr")
    fonts=OxmlElement("w:rFonts")
    for a in ("w:ascii","w:hAnsi","w:eastAsia"):fonts.set(qn(a),"宋体")
    rp.append(fonts);rp.append(OxmlElement("w:b"))
    sz=OxmlElement("w:sz");sz.set(qn("w:val"),"21");rp.append(sz)
    sz2=OxmlElement("w:szCs");sz2.set(qn("w:val"),"21");rp.append(sz2)
    r.append(rp);t=OxmlElement("w:t");t.text=text;r.append(t);p.append(r)

def add_page_field(p):
    r=OxmlElement("w:r");fc=OxmlElement("w:fldChar");fc.set(qn("w:fldCharType"),"begin");r.append(fc);p.append(r)
    r=OxmlElement("w:r");it=OxmlElement("w:instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space","preserve");it.text=" PAGE ";r.append(it);p.append(r)
    r=OxmlElement("w:r");fc=OxmlElement("w:fldChar");fc.set(qn("w:fldCharType"),"separate");r.append(fc);p.append(r)
    add_text_run(p,"1")
    r=OxmlElement("w:r");fc=OxmlElement("w:fldChar");fc.set(qn("w:fldCharType"),"end");r.append(fc);p.append(r)

def rebuild_header(header,total):
    root=header._element
    for x in list(root):root.remove(x)
    p=OxmlElement("w:p");pPr=OxmlElement("w:pPr")
    jc=OxmlElement("w:jc");jc.set(qn("w:val"),"left");pPr.append(jc)
    tabs=OxmlElement("w:tabs");tab=OxmlElement("w:tab");tab.set(qn("w:val"),"right");tab.set(qn("w:pos"),"8760");tabs.append(tab);pPr.append(tabs)
    bdr=OxmlElement("w:pBdr");bottom=OxmlElement("w:bottom")
    for k,v in (("w:val","single"),("w:sz","6"),("w:space","1"),("w:color","auto")):bottom.set(qn(k),v)
    bdr.append(bottom);pPr.append(bdr)
    ind=OxmlElement("w:ind");ind.set(qn("w:left"),"0");ind.set(qn("w:right"),"0");ind.set(qn("w:firstLine"),"0");pPr.append(ind)
    sp=OxmlElement("w:spacing");sp.set(qn("w:before"),"0");sp.set(qn("w:after"),"0");pPr.append(sp)
    p.append(pPr)
    add_text_run(p,"炎黄职业技术学院毕业论文")
    r=OxmlElement("w:r");r.append(OxmlElement("w:tab"));p.append(r)
    add_text_run(p,"第 ");add_page_field(p);add_text_run(p,f" 页  共 {total} 页")
    root.append(p)

def main():
    docx=Path(sys.argv[1]);pdf=Path(sys.argv[2])
    pages=pdf_pages(pdf); npages=[norm(x) for x in pages]
    d=Document(docx)
    bi=body_index(d);ri=ref_index(d)
    bphys=None
    for i,t in enumerate(npages,1):
        if "炎黄职业技术学院毕业论文" in t and "第1页" in t and "1绪论" in t:
            bphys=i;break
    if bphys is None:
        # 兜底：跳过前置页后寻找正文首章
        for i,t in enumerate(npages[8:],9):
            if "1绪论" in t:
                bphys=i;break
    if bphys is None:raise RuntimeError("无法确定正文物理起始页")
    body_pages=len(pages)-bphys+1

    headings=[]
    for p in d.paragraphs[bi:ri]:
        if level(p.text) in (1,2):
            headings.append((norm(p.text),p.text.strip()))
    mapping={}
    for key,raw in headings:
        for phys in range(bphys,len(pages)+1):
            if key in npages[phys-1]:
                mapping[key]=phys-bphys+1;break

    # 参考文献和致谢
    for key in ("参考文献","致谢"):
        for phys in range(bphys,len(pages)+1):
            if key in npages[phys-1]:
                mapping[key]=phys-bphys+1;break

    ti=next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="目录")
    for p in d.paragraphs[ti+1:bi]:
        title=norm(p.text.split("\t")[0])
        matches=[(k,v) for k,v in mapping.items() if title.startswith(k)]
        if matches:
            k,v=max(matches,key=lambda kv:len(kv[0]))
            replace_last_number(p,v)

    rebuild_header(d.sections[-1].header,body_pages)
    d.save(docx)
    print("TOC_HEADER_UPDATED","physical",len(pages),"body_start",bphys,"body_pages",body_pages,"mapping",mapping)

if __name__=="__main__":
    main()
