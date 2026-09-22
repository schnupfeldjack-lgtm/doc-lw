# -*- coding: utf-8 -*-
from pathlib import Path
import re, subprocess, sys, zipfile
from docx import Document
from docx.oxml.ns import qn
import pdfplumber

def norm(s):return re.sub(r"\s+","",s).replace("–","-").replace("—","-")
def body_idx(d):
    return [i for i,p in enumerate(d.paragraphs) if norm(p.text)=="1绪论" and not (p.style and p.style.name.startswith("toc"))][-1]
def ref_idx(d):return next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="参考文献")
def han(d):
    return sum("\u4e00"<=c<="\u9fff" for p in d.paragraphs[body_idx(d):ref_idx(d)] for c in p.text)
def pdf_text(pdf,i):
    return subprocess.check_output(["pdftotext","-f",str(i),"-l",str(i),"-layout",str(pdf),"-"],text=True,errors="ignore")
def pages_count(pdf):
    s=subprocess.check_output(["pdfinfo",str(pdf)],text=True,errors="ignore")
    return int(re.search(r"^Pages:\s+(\d+)",s,re.M).group(1))
def special(d,name):
    return next(p for p in d.paragraphs if norm(p.text)==norm(name))

def main():
    docx=Path(sys.argv[1]);pdf=Path(sys.argv[2])
    d=Document(docx)
    hc=han(d)
    assert 9500<=hc<=11000,hc
    assert len(d.inline_shapes)>=2
    assert len(d.sections)>=5
    for s in d.sections[:-1]:
        hx=s.header._element.xml
        assert "PAGE" not in hx and "SECTIONPAGES" not in hx
    hx=d.sections[-1].header._element.xml
    assert "炎黄职业技术学院毕业论文" in hx and "PAGE" in hx and "SECTIONPAGES" in hx
    pg=d.sections[-1]._sectPr.find(qn("w:pgNumType"))
    assert pg is not None and pg.get(qn("w:start"))=="1"
    assert special(d,"参 考 文 献").paragraph_format.page_break_before
    ack=special(d,"致 谢");assert ack.paragraph_format.page_break_before

    with zipfile.ZipFile(docx) as z:
        names=set(z.namelist())
        assert "word/media/fig3_1_repaired.png" in names and len(z.read("word/media/fig3_1_repaired.png"))>100000
        assert "word/media/fig4_1_repaired.png" in names and len(z.read("word/media/fig4_1_repaired.png"))>100000

    total=pages_count(pdf);texts=[pdf_text(pdf,i) for i in range(1,total+1)];ntexts=[norm(x) for x in texts]
    bphys=next(i for i,t in enumerate(ntexts,1) if "第1页" in t and "1绪论" in t)
    assert bphys==9,f"正文应为物理第9页，实际{bphys}"
    for i in range(1,bphys):
        assert not re.search(r"第\s*\d+\s*页\s*共\s*\d+\s*页",texts[i-1]),f"前置第{i}页出现页码"
    body_pages=total-bphys+1
    assert f"第1页共{body_pages}页" in ntexts[bphys-1],ntexts[bphys-1][:120]
    assert f"第{body_pages}页共{body_pages}页" in ntexts[-1],ntexts[-1][:120]

    ack_phys=next(i for i,t in enumerate(ntexts,1) if "致谢" in t and "本论文是在指导老师" in t)
    with pdfplumber.open(pdf) as pf:
        words=pf.pages[ack_phys-1].extract_words()
        cand=[w for w in words if "致" in w["text"] or "谢" in w["text"]]
        if not cand:raise AssertionError("致谢标题坐标未识别")
        top=min(w["top"] for w in cand)
        assert top<230,f"致谢标题位置过低，top={top}"
    print("QA_OK","han",hc,"physical_pages",total,"body_start",bphys,"body_pages",body_pages,"ack_page",ack_phys,"ack_top",top)

if __name__=="__main__":main()
