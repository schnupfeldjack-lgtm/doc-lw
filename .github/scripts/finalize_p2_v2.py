# -*- coding: utf-8 -*-
from pathlib import Path
import copy, re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT / "新建文件夹/新建文件夹/装配式施工质量管理问题及优化研究——以市政项目为例/装配式施工质量管理问题及优化研究——以市政项目为例_最终优化版.docx"

def norm(s):
    return re.sub(r"\s+", "", s).replace("–","-").replace("—","-")

def body_idx(d):
    xs=[i for i,p in enumerate(d.paragraphs) if norm(p.text)=="1引言" and not (p.style and p.style.name.lower().startswith("toc"))]
    if not xs: raise RuntimeError("找不到正文首章")
    return xs[-1]

def ref_idx(d):
    return next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="参考文献")

def han_count(d):
    ps=d.paragraphs; s=body_idx(d); e=ref_idx(d)
    return sum("\u4e00"<=c<="\u9fff" for p in ps[s:e] for c in p.text)

def delete_p(p):
    el=p._element; par=el.getparent()
    if par is not None: par.remove(el)

def trim_body(d):
    targets=[
        "这种组织优势的前提，是构件规格、运输节拍和安装基准保持一致",
        "Yu等的研究显示，不同参与方对装配式项目质量缺陷的影响存在依赖关系",
        "信息化并不意味着必须建设复杂平台",
        "刘美霞等关于质量追溯智能预警的研究与汪丛军等关于产业协同平台的研究",
        "为了避免责任争议，进场验收应尽量在卸车或首次转运前完成关键影像记录",
        "装配式工程适合使用二维码、BIM、点云、机器视觉等数字技术",
    ]
    removed=[]
    for prefix in targets:
        if han_count(d)<=10800: break
        p=next((p for p in d.paragraphs[body_idx(d):ref_idx(d)] if norm(p.text).startswith(norm(prefix))),None)
        if p:
            n=sum("\u4e00"<=c<="\u9fff" for c in p.text)
            delete_p(p); removed.append((prefix[:18],n))
    n=han_count(d)
    print("BODY_HAN_AFTER_TRIM",n,"removed",removed)
    if not 9500<=n<=11000:
        raise RuntimeError(f"正文汉字数不在约1万字范围: {n}")

def remove_inline_pagebreak(p):
    for br in list(p._p.xpath(".//w:br[@w:type='page']")):
        par=br.getparent()
        if par is not None: par.remove(br)

def remove_empty_neighbors(p):
    cur=p._element.getprevious()
    while cur is not None and cur.tag==qn("w:p") and not "".join(cur.xpath(".//w:t/text()")).strip():
        pp=cur.find(qn("w:pPr"))
        if pp is not None and pp.find(qn("w:sectPr")) is not None: break
        prev=cur.getprevious(); cur.getparent().remove(cur); cur=prev
    cur=p._element.getnext()
    while cur is not None and cur.tag==qn("w:p") and not "".join(cur.xpath(".//w:t/text()")).strip():
        pp=cur.find(qn("w:pPr"))
        if pp is not None and pp.find(qn("w:sectPr")) is not None: break
        nxt=cur.getnext(); cur.getparent().remove(cur); cur=nxt

def findp(d,name):
    n=norm(name)
    return next((p for p in d.paragraphs if norm(p.text)==n),None)

def format_special_pages(d):
    for name,align in [("结  论",WD_ALIGN_PARAGRAPH.LEFT),("参 考 文 献",WD_ALIGN_PARAGRAPH.CENTER),("致  谢",WD_ALIGN_PARAGRAPH.LEFT)]:
        p=findp(d,name)
        if not p: raise RuntimeError("找不到标题 "+name)
        remove_inline_pagebreak(p)
        remove_empty_neighbors(p)
        p.paragraph_format.page_break_before=True
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(10)
        p.paragraph_format.first_line_indent=Pt(0)
        p.alignment=align
    ack=findp(d,"致  谢")
    ps=d.paragraphs
    ai=next(i for i,p in enumerate(ps) if p._element is ack._element)
    if ai+1<len(ps):
        p=ps[ai+1]
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(0)

def make_pb():
    p=OxmlElement("w:p");r=OxmlElement("w:r");br=OxmlElement("w:br");br.set(qn("w:type"),"page");r.append(br);p.append(r);return p

def is_pb(el):
    return el is not None and el.tag==qn("w:p") and bool(el.xpath(".//w:br[@w:type='page']"))

def ensure_pb_after(el):
    if not is_pb(el.getnext()): el.addnext(make_pb())

def table_text(el):
    return "".join(el.xpath(".//w:t/text()")).replace(" ","").replace("\u3000","")

def paragraph_text(el):
    return "".join(el.xpath(".//w:t/text()")).replace(" ","").replace("\u3000","")

def repair_front_paging(d):
    body=d._element.body
    tops=[e for e in body.iterchildren()]
    tbls=[e for e in tops if e.tag==qn("w:tbl")]
    if tbls: ensure_pb_after(tbls[0])  # 开题报告
    # 任务书：日期段后分页
    for el in tops:
        if el.tag==qn("w:p") and "2025年9月10日" in paragraph_text(el):
            ensure_pb_after(el); break
    # 中期检查表：明确独立页
    if len(d.tables)>2: ensure_pb_after(d.tables[2]._element)
    # 英文摘要、目录、正文均从新页开始
    for key in ("Abstract：","目  录","1  引言"):
        p=next((p for p in d.paragraphs if p.text.strip().startswith(key)),None)
        if p:
            p.paragraph_format.page_break_before=True
    # 封面节属性移到封面年月之后，防止封面被切裂
    cover_sect=None
    for p in d.paragraphs:
        if "专业技术职务" in p.text:
            pp=p._p.find(qn("w:pPr"))
            if pp is not None:
                sp=pp.find(qn("w:sectPr"))
                if sp is not None:
                    cover_sect=copy.deepcopy(sp);pp.remove(sp);break
    if cover_sect is not None:
        for tag in ("w:headerReference","w:footerReference"):
            for e in list(cover_sect.findall(qn(tag))): cover_sect.remove(e)
        cover_end=None
        for el in body.iterchildren():
            if el.tag==qn("w:tbl") and "2025" in table_text(el) and "9" in table_text(el):
                cover_end=el
        if cover_end is not None:
            nxt=cover_end.getnext()
            already=False
            if nxt is not None and nxt.tag==qn("w:p"):
                pp=nxt.find(qn("w:pPr")); already=pp is not None and pp.find(qn("w:sectPr")) is not None
            if not already:
                p=OxmlElement("w:p");pp=OxmlElement("w:pPr");pp.append(cover_sect);p.append(pp);cover_end.addnext(p)

def clear_hf(part):
    root=part._element
    for c in list(root): root.remove(c)
    root.append(OxmlElement("w:p"))

def add_text(p,text,bold=True):
    r=OxmlElement("w:r");rp=OxmlElement("w:rPr")
    fs=OxmlElement("w:rFonts")
    for a in ("w:ascii","w:hAnsi","w:eastAsia"): fs.set(qn(a),"宋体")
    rp.append(fs)
    if bold: rp.append(OxmlElement("w:b"))
    sz=OxmlElement("w:sz");sz.set(qn("w:val"),"21");rp.append(sz)
    sz2=OxmlElement("w:szCs");sz2.set(qn("w:val"),"21");rp.append(sz2)
    r.append(rp);t=OxmlElement("w:t");t.text=text;r.append(t);p.append(r)

def add_page_field(p):
    r=OxmlElement("w:r");f=OxmlElement("w:fldChar");f.set(qn("w:fldCharType"),"begin");r.append(f);p.append(r)
    r=OxmlElement("w:r");it=OxmlElement("w:instrText");it.set("{http://www.w3.org/XML/1998/namespace}space","preserve");it.text=" PAGE ";r.append(it);p.append(r)
    r=OxmlElement("w:r");f=OxmlElement("w:fldChar");f.set(qn("w:fldCharType"),"separate");r.append(f);p.append(r)
    add_text(p,"1")
    r=OxmlElement("w:r");f=OxmlElement("w:fldChar");f.set(qn("w:fldCharType"),"end");r.append(f);p.append(r)

def build_header(h,total=1):
    root=h._element
    for c in list(root):root.remove(c)
    p=OxmlElement("w:p");pp=OxmlElement("w:pPr")
    jc=OxmlElement("w:jc");jc.set(qn("w:val"),"left");pp.append(jc)
    tabs=OxmlElement("w:tabs");tb=OxmlElement("w:tab");tb.set(qn("w:val"),"right");tb.set(qn("w:pos"),"8760");tabs.append(tb);pp.append(tabs)
    bdr=OxmlElement("w:pBdr");bt=OxmlElement("w:bottom")
    for k,v in (("w:val","single"),("w:sz","6"),("w:space","1"),("w:color","auto")):bt.set(qn(k),v)
    bdr.append(bt);pp.append(bdr)
    sp=OxmlElement("w:spacing");sp.set(qn("w:before"),"0");sp.set(qn("w:after"),"0");pp.append(sp)
    p.append(pp)
    add_text(p,"炎黄职业技术学院毕业论文")
    r=OxmlElement("w:r");r.append(OxmlElement("w:tab"));p.append(r)
    add_text(p,"第 ");add_page_field(p);add_text(p,f" 页  共 {total} 页")
    root.append(p)

def fix_headers(d):
    if len(d.sections)<3: raise RuntimeError("节数量异常")
    for s in d.sections[:-1]:
        s.header.is_linked_to_previous=False;clear_hf(s.header)
        s.footer.is_linked_to_previous=False;clear_hf(s.footer)
        pg=s._sectPr.find(qn("w:pgNumType"))
        if pg is not None:s._sectPr.remove(pg)
    s=d.sections[-1]
    s.header.is_linked_to_previous=False;build_header(s.header,1)
    s.footer.is_linked_to_previous=False;clear_hf(s.footer)
    pg=s._sectPr.find(qn("w:pgNumType"))
    if pg is None:pg=OxmlElement("w:pgNumType");s._sectPr.append(pg)
    pg.set(qn("w:start"),"1")
    uf=d.settings._element.find(qn("w:updateFields"))
    if uf is None:uf=OxmlElement("w:updateFields");d.settings._element.append(uf)
    uf.set(qn("w:val"),"true")

def superscript_citations(d):
    pat=re.compile(r"(\[\d+\])")
    s=body_idx(d);e=ref_idx(d)
    count=0
    for p in d.paragraphs[s:e]:
        if not pat.search(p.text): continue
        if re.match(r"^\s*\d+(?:[．.]\d+)*\s",p.text): continue
        full=p.text
        # preserve first visible run formatting
        base=None
        for r in p.runs:
            rp=r._r.find(qn("w:rPr"))
            if rp is not None: base=copy.deepcopy(rp);break
        for el in list(p._p):
            if el.tag in (qn("w:r"),qn("w:hyperlink")):p._p.remove(el)
        for tok in pat.split(full):
            if not tok:continue
            r=OxmlElement("w:r")
            if base is not None:r.append(copy.deepcopy(base))
            rp=r.find(qn("w:rPr"))
            if rp is None:rp=OxmlElement("w:rPr");r.insert(0,rp)
            if pat.fullmatch(tok):
                va=OxmlElement("w:vertAlign");va.set(qn("w:val"),"superscript");rp.append(va);count+=1
            t=OxmlElement("w:t");t.text=tok;r.append(t);p._p.append(r)
    print("SUPERSCRIPT_CITATIONS",count)

def main():
    if not DOC.exists(): raise FileNotFoundError(DOC)
    d=Document(DOC)
    trim_body(d)
    format_special_pages(d)
    repair_front_paging(d)
    superscript_citations(d)
    fix_headers(d)
    text="\n".join(p.text for p in d.paragraphs)
    if "待填写" in text: raise RuntimeError("仍存在待填写")
    d.save(DOC)
    d2=Document(DOC)
    refs=[p.text.strip() for p in d2.paragraphs if re.match(r"^\[\d+\]",p.text.strip())]
    print("FINALIZE_OK","sections",len(d2.sections),"han",han_count(d2),"refs",len(refs),"pics",len(d2.inline_shapes),"size",DOC.stat().st_size)
    if len(refs)!=15: raise RuntimeError(f"参考文献数量错误: {len(refs)}")
    if len(d2.inline_shapes)<5: raise RuntimeError(f"图片数量不足: {len(d2.inline_shapes)}")

if __name__=="__main__":
    main()
