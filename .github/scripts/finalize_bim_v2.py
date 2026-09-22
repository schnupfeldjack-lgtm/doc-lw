# -*- coding: utf-8 -*-
from pathlib import Path
import copy, re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
PAPER = ROOT / "新建文件夹/新建文件夹/BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例"
SRC = PAPER / "BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例_炎黄模板严格版_图片已补全.docx"
OUT = PAPER / "BIM协同设计对建筑结构设计质量的影响研究——以装配式住宅项目为例_最终版.docx"

def norm(s):
    return re.sub(r"\s+", "", s).replace("–","-").replace("—","-")

def ptext(p):
    return p.text.strip()

def delete_p(p):
    el=p._element
    parent=el.getparent()
    if parent is not None:
        parent.remove(el)

def body_index(doc):
    idxs=[i for i,p in enumerate(doc.paragraphs) if norm(p.text)=="1绪论" and not (p.style and p.style.name.startswith("toc"))]
    if not idxs:
        raise RuntimeError("找不到正文首章")
    return idxs[-1]

def ref_index(doc):
    for i,p in enumerate(doc.paragraphs):
        if norm(p.text)=="参考文献":
            return i
    raise RuntimeError("找不到参考文献")

def han_count(doc):
    ps=doc.paragraphs
    s=body_index(doc); e=ref_index(doc)
    txt="".join(p.text for p in ps[s:e])
    return sum("\u4e00" <= ch <= "\u9fff" for ch in txt)

def heading_level(text):
    t=norm(text)
    if re.match(r"^\d+\.\d+\.\d+",t): return 3
    if re.match(r"^\d+[．.]\d+",t): return 2
    if re.match(r"^\d+",t): return 1
    return None

def find_body_heading(doc, target):
    t=norm(target)
    s=body_index(doc)
    for i,p in enumerate(doc.paragraphs[s:],start=s):
        if norm(p.text)==t:
            return i
    return None

def section_range(doc, target):
    i=find_body_heading(doc,target)
    if i is None:
        return None
    ps=doc.paragraphs
    lvl=heading_level(ps[i].text)
    j=i+1
    while j<len(ps):
        lv=heading_level(ps[j].text)
        if lv is not None and lv<=lvl:
            break
        if norm(ps[j].text) in ("参考文献","致谢"):
            break
        j+=1
    return i,j

def section_han(doc,target):
    rg=section_range(doc,target)
    if not rg:return 0
    i,j=rg
    return sum("\u4e00"<=ch<="\u9fff" for p in doc.paragraphs[i:j] for ch in p.text)

def remove_section(doc,target):
    rg=section_range(doc,target)
    if not rg:return False
    i,j=rg
    for p in list(doc.paragraphs[i:j]):
        delete_p(p)
    return True

def remove_paragraph_prefix(doc,prefix):
    s=body_index(doc); pfx=norm(prefix)
    for p in doc.paragraphs[s:]:
        if norm(p.text).startswith(pfx):
            delete_p(p); return True
    return False

def trim_body(doc):
    # 优先删除偏离论文核心、且与前后内容重复的扩展小节；核心案例分析、碰撞检查和质量影响分析保留。
    candidates=[
        "5.3.5  完善组织保障机制",
        "5.3.6  强化BIM应用与项目特点的适配",
        "5.1.2  BIM应用与项目各阶段的衔接",
        "2.2.1  协同设计的常用工作模式",
        "3.2.1  结构专业BIM模型的构建要点",
        "3.3.1  BIM协同设计的实施流程",
        "2.1.1  BIM技术的常用软件与数据标准",
        "5.1.1  BIM应用效果的量化评价指标",
    ]
    before=han_count(doc)
    removed=[]
    for h in candidates:
        cur=han_count(doc)
        if cur<=10500:break
        n=section_han(doc,h)
        if n and cur-n>=9500 and remove_section(doc,h):
            removed.append((h,n))
    # 6.2 保留研究局限与未来研究，删除偏离结构设计主题的运维展开段，必要时用于进一步收敛字数。
    for prefix in [
        "BIM的价值不仅体现在设计阶段，也应延伸至运维阶段",
        "要实现BIM在运维阶段的应用，需要在设计阶段就考虑运维的信息需求",
    ]:
        cur=han_count(doc)
        if cur<=10500:break
        # 先估算目标段字数
        target=None
        for p in doc.paragraphs[body_index(doc):ref_index(doc)]:
            if norm(p.text).startswith(norm(prefix)):
                target=p;break
        if target:
            n=sum("\u4e00"<=ch<="\u9fff" for ch in target.text)
            if cur-n>=9500:
                delete_p(target);removed.append((prefix,n))
    after=han_count(doc)
    print("正文汉字数",before,"->",after)
    print("删除的重复/偏题扩展:",removed)
    if not (9500<=after<=11000):
        raise RuntimeError(f"正文汉字数未落入约1万字范围: {after}")

def find_special(doc,name):
    n=norm(name)
    for p in doc.paragraphs:
        if norm(p.text)==n:return p
    return None

def remove_empty_neighbors(p):
    # 删除标题前后的连续空段，避免参考文献/致谢页出现大面积无意义空白
    cur=p._element.getprevious()
    while cur is not None and cur.tag==qn("w:p") and not "".join(cur.xpath(".//w:t/text()")).strip():
        prev=cur.getprevious()
        pp=cur.find(qn("w:pPr"))
        if pp is not None and pp.find(qn("w:sectPr")) is not None:
            break
        cur.getparent().remove(cur);cur=prev
    cur=p._element.getnext()
    while cur is not None and cur.tag==qn("w:p") and not "".join(cur.xpath(".//w:t/text()")).strip():
        nxt=cur.getnext()
        pp=cur.find(qn("w:pPr"))
        if pp is not None and pp.find(qn("w:sectPr")) is not None:
            break
        cur.getparent().remove(cur);cur=nxt

def format_special_pages(doc):
    ref=find_special(doc,"参 考 文 献")
    ack=find_special(doc,"致 谢")
    if not ref or not ack: raise RuntimeError("找不到参考文献或致谢标题")
    for p,align in [(ref,WD_ALIGN_PARAGRAPH.CENTER),(ack,WD_ALIGN_PARAGRAPH.LEFT)]:
        remove_empty_neighbors(p)
        p.paragraph_format.page_break_before=True
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(12)
        p.paragraph_format.first_line_indent=Pt(0)
        p.alignment=align
        for r in p.runs:
            if r.text:
                r.font.name="宋体";r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"宋体")
                r.font.size=Pt(14);r.bold=True
    # 致谢正文紧接标题，不再保留多余空白
    ps=doc.paragraphs
    ai=ps.index(ack)
    if ai+1<len(ps):
        body=ps[ai+1]
        body.paragraph_format.space_before=Pt(0)
        body.paragraph_format.space_after=Pt(0)
        body.paragraph_format.line_spacing=Pt(22)
        body.paragraph_format.first_line_indent=Pt(24)

def has_page_break(el):
    return bool(el.xpath(".//w:br[@w:type='page']"))

def remove_preceding_explicit_pagebreak(body_p):
    cur=body_p._element.getprevious()
    if cur is not None and cur.tag==qn("w:p") and has_page_break(cur):
        cur.getparent().remove(cur)

def split_before_body(doc):
    bi=body_index(doc);bp=doc.paragraphs[bi]
    bp.paragraph_format.page_break_before=False
    remove_preceding_explicit_pagebreak(bp)

    # 复制末节页面设置作为“前置部分最后一节”的节属性，但不复制页眉页脚。
    final_sect=doc._element.body.find(qn("w:sectPr"))
    if final_sect is None: raise RuntimeError("文档缺少末节sectPr")
    pre=copy.deepcopy(final_sect)
    for tag in ("w:headerReference","w:footerReference","w:pgNumType"):
        for el in list(pre.findall(qn(tag))):
            pre.remove(el)
    typ=pre.find(qn("w:type"))
    if typ is None:
        typ=OxmlElement("w:type");pre.insert(0,typ)
    typ.set(qn("w:val"),"nextPage")

    # 在正文前插入一个仅承载分节符的空段，正文由新节开始。
    carrier=OxmlElement("w:p");pPr=OxmlElement("w:pPr");pPr.append(pre);carrier.append(pPr)
    bp._element.addprevious(carrier)

def clear_header(header):
    root=header._element
    for c in list(root):root.remove(c)
    p=OxmlElement("w:p");root.append(p)

def clear_footer(footer):
    root=footer._element
    for c in list(root):root.remove(c)
    p=OxmlElement("w:p");root.append(p)

def add_text_run(p,text,bold=True):
    r=OxmlElement("w:r");rp=OxmlElement("w:rPr")
    fonts=OxmlElement("w:rFonts")
    for a in ("w:ascii","w:hAnsi","w:eastAsia"):fonts.set(qn(a),"宋体")
    rp.append(fonts)
    if bold:rp.append(OxmlElement("w:b"))
    sz=OxmlElement("w:sz");sz.set(qn("w:val"),"21");rp.append(sz)
    szc=OxmlElement("w:szCs");szc.set(qn("w:val"),"21");rp.append(szc)
    r.append(rp);t=OxmlElement("w:t");t.text=text;r.append(t);p.append(r)

def add_tab(p):
    r=OxmlElement("w:r");r.append(OxmlElement("w:tab"));p.append(r)

def add_field(p,instr,cached):
    r=OxmlElement("w:r");fc=OxmlElement("w:fldChar");fc.set(qn("w:fldCharType"),"begin");r.append(fc);p.append(r)
    r=OxmlElement("w:r");it=OxmlElement("w:instrText");it.set("{http://www.w3.org/XML/1998/namespace}space","preserve");it.text=f" {instr} ";r.append(it);p.append(r)
    r=OxmlElement("w:r");fc=OxmlElement("w:fldChar");fc.set(qn("w:fldCharType"),"separate");r.append(fc);p.append(r)
    add_text_run(p,str(cached))
    r=OxmlElement("w:r");fc=OxmlElement("w:fldChar");fc.set(qn("w:fldCharType"),"end");r.append(fc);p.append(r)

def build_body_header(header):
    root=header._element
    for c in list(root):root.remove(c)
    p=OxmlElement("w:p");pPr=OxmlElement("w:pPr")
    jc=OxmlElement("w:jc");jc.set(qn("w:val"),"left");pPr.append(jc)
    tabs=OxmlElement("w:tabs");tab=OxmlElement("w:tab");tab.set(qn("w:val"),"right");tab.set(qn("w:pos"),"8760");tabs.append(tab);pPr.append(tabs)
    bdr=OxmlElement("w:pBdr");bottom=OxmlElement("w:bottom")
    for k,v in [("w:val","single"),("w:sz","6"),("w:space","1"),("w:color","auto")]:bottom.set(qn(k),v)
    bdr.append(bottom);pPr.append(bdr)
    ind=OxmlElement("w:ind");ind.set(qn("w:left"),"0");ind.set(qn("w:right"),"0");ind.set(qn("w:firstLine"),"0");pPr.append(ind)
    sp=OxmlElement("w:spacing");sp.set(qn("w:before"),"0");sp.set(qn("w:after"),"0");pPr.append(sp)
    p.append(pPr)
    add_text_run(p,"炎黄职业技术学院毕业论文");add_tab(p);add_text_run(p,"第 ");add_field(p,"PAGE",1);add_text_run(p," 页  共 ");add_field(p,"SECTIONPAGES",1);add_text_run(p," 页")
    root.append(p)

def set_body_paging(doc_path):
    d=Document(doc_path)
    if len(d.sections)<5:
        raise RuntimeError(f"正文分节失败，当前节数={len(d.sections)}")
    # 前置全部不显示页眉页脚
    for s in d.sections[:-1]:
        s.header.is_linked_to_previous=False;clear_header(s.header)
        s.footer.is_linked_to_previous=False;clear_footer(s.footer)
        pg=s._sectPr.find(qn("w:pgNumType"))
        if pg is not None:s._sectPr.remove(pg)
    # 正文独立页眉，从1开始，只统计正文节页数
    s=d.sections[-1]
    s.header.is_linked_to_previous=False;build_body_header(s.header)
    s.footer.is_linked_to_previous=False;clear_footer(s.footer)
    pg=s._sectPr.find(qn("w:pgNumType"))
    if pg is None:pg=OxmlElement("w:pgNumType");s._sectPr.append(pg)
    pg.set(qn("w:start"),"1")
    # 自动更新字段
    uf=d.settings._element.find(qn("w:updateFields"))
    if uf is None:uf=OxmlElement("w:updateFields");d.settings._element.append(uf)
    uf.set(qn("w:val"),"true")
    d.save(doc_path)

def main():
    if not SRC.exists():raise FileNotFoundError(SRC)
    OUT.parent.mkdir(parents=True,exist_ok=True)
    d=Document(SRC)
    trim_body(d)
    format_special_pages(d)
    split_before_body(d)
    d.save(OUT)
    set_body_paging(OUT)
    # 最终结构复核
    d=Document(OUT)
    n=han_count(d)
    ack=find_special(d,"致 谢")
    print("FINAL",OUT,OUT.stat().st_size,"sections",len(d.sections),"han",n,"ack_pagebreak",ack.paragraph_format.page_break_before)
    assert 9500<=n<=11000
    assert len(d.sections)>=5

if __name__=="__main__":
    main()
