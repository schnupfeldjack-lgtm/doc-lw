# -*- coding: utf-8 -*-
"""对重写版做最终结构、分页和格式处理。"""
from pathlib import Path
import copy, re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
PAPER = ROOT / "新建文件夹/新建文件夹/再生混凝土在装配式建筑中的应用评价——以住宅项目为例"
DOCX = PAPER / "再生混凝土在装配式建筑中的应用评价——以住宅项目为例_最终优化版.docx"

def norm(s):
    return re.sub(r"\s+", "", s).replace("–","-").replace("—","-")

def delete_p(p):
    el=p._element
    par=el.getparent()
    if par is not None:
        par.remove(el)

def body_index(d):
    xs=[i for i,p in enumerate(d.paragraphs)
        if norm(p.text).startswith("1绪论") and not (p.style and p.style.name.lower().startswith("toc"))]
    if not xs: raise RuntimeError("找不到正文首章")
    return xs[-1]

def exact_index(d,text,start=0):
    n=norm(text)
    for i,p in enumerate(d.paragraphs[start:],start):
        if norm(p.text)==n:
            return i
    return None

def ref_index(d):
    i=exact_index(d,"参考文献")
    if i is None: i=exact_index(d,"参 考 文 献")
    if i is None: raise RuntimeError("找不到参考文献")
    return i

def ack_index(d):
    i=exact_index(d,"致谢")
    if i is None: i=exact_index(d,"致 谢")
    if i is None: raise RuntimeError("找不到致谢")
    return i

def han_count(d):
    ps=d.paragraphs
    s=body_index(d); e=ref_index(d)
    return sum("\u4e00"<=ch<="\u9fff" for p in ps[s:e] for ch in p.text)

def remove_duplicate_standalone_conclusion(d):
    # fill_docx 按模板额外生成独立“结论”。本论文已有第6章结论，删除独立重复部分。
    ps=d.paragraphs
    b=body_index(d)
    candidates=[i for i,p in enumerate(ps[b:],b) if norm(p.text)=="结论"]
    if not candidates:
        return
    i=candidates[-1]
    r=ref_index(d)
    if i<r:
        for p in list(d.paragraphs[i:r]):
            delete_p(p)

def remove_blank_neighbors(p):
    # 保留承载 sectPr 的空段；删除无意义空段，避免标题被推到页面中下部
    cur=p._element.getprevious()
    while cur is not None and cur.tag==qn("w:p"):
        txt="".join(cur.xpath(".//w:t/text()")).strip()
        pp=cur.find(qn("w:pPr"))
        has_sect=pp is not None and pp.find(qn("w:sectPr")) is not None
        has_break=bool(cur.xpath(".//w:br[@w:type='page']"))
        if txt or has_sect or has_break: break
        prev=cur.getprevious();cur.getparent().remove(cur);cur=prev
    cur=p._element.getnext()
    while cur is not None and cur.tag==qn("w:p"):
        txt="".join(cur.xpath(".//w:t/text()")).strip()
        pp=cur.find(qn("w:pPr"))
        has_sect=pp is not None and pp.find(qn("w:sectPr")) is not None
        has_break=bool(cur.xpath(".//w:br[@w:type='page']"))
        if txt or has_sect or has_break: break
        nxt=cur.getnext();cur.getparent().remove(cur);cur=nxt

def style_special_pages(d):
    ri=ref_index(d); ai=ack_index(d)
    ref=d.paragraphs[ri]; ack=d.paragraphs[ai]
    for p,align in ((ref,WD_ALIGN_PARAGRAPH.CENTER),(ack,WD_ALIGN_PARAGRAPH.LEFT)):
        remove_blank_neighbors(p)
        p.paragraph_format.page_break_before=True
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(12)
        p.paragraph_format.first_line_indent=Pt(0)
        p.alignment=align
        for r in p.runs:
            if r.text:
                r.font.name="宋体"
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"宋体")
                r.font.size=Pt(14);r.bold=True

def rebuild_extra_toc_entries(d):
    ps=d.paragraphs
    ti=exact_index(d,"目录")
    if ti is None: ti=exact_index(d,"目 录")
    bi=body_index(d)
    if ti is None or ti>=bi:return
    toc=ps[ti+1:bi]
    # 删除 fill_docx 自动附加但次序不对的“结论/致谢/参考文献”。
    for p in list(toc):
        n=norm(p.text.split("\t")[0])
        if n in ("结论","致谢","参考文献"):
            delete_p(p)
    # 重新获取列表，找到最后一个目录条目
    ps=d.paragraphs; ti=exact_index(d,"目录"); bi=body_index(d)
    entries=[p for p in ps[ti+1:bi] if "\t" in p.text]
    if not entries:return
    proto=entries[0]
    anchor=entries[-1]._element
    for label in ("参考文献","致谢"):
        new=copy.deepcopy(proto._element)
        # 清正文 run，保留 pPr；新建 label + tab + 0
        for tag in ("w:r","w:hyperlink","w:fldSimple"):
            for x in list(new.findall(qn(tag))): new.remove(x)
        def add_text(txt):
            r=OxmlElement("w:r")
            # 从 proto 第一文字 run复制格式
            for rr in proto._p.findall(qn("w:r")):
                if rr.find(qn("w:t")) is not None:
                    rp=rr.find(qn("w:rPr"))
                    if rp is not None:r.append(copy.deepcopy(rp))
                    break
            t=OxmlElement("w:t");t.text=txt;r.append(t);new.append(r)
        add_text(label)
        r=OxmlElement("w:r");r.append(OxmlElement("w:tab"));new.append(r)
        add_text("0")
        anchor.addnext(new);anchor=new


def remove_template_residue(d):
    markers=[
        "正文中公式、图与表的字体一律用5号宋体",
        "正文各页的格式请以此页为标准复制",
        "为保证打印效果",
        "说明：结论",
        "说明:结论",
        "参考文献著录规则",
        "毕业设计论文所列",
    ]
    s=body_index(d);e=ref_index(d)
    for p in list(d.paragraphs[s:e]):
        if any(m in p.text for m in markers):
            delete_p(p)

def ensure_front_pagination(d):
    # 中文摘要=物理第5页、英文摘要=第6页、目录=第7-8页、正文=第9页。
    ps=d.paragraphs
    # English abstract starts a new page.
    for p in ps:
        if p.text.strip().startswith("Abstract"):
            p.paragraph_format.page_break_before=True
            break
    # TOC title starts a new page.
    toc=None
    for p in d.paragraphs:
        if norm(p.text)=="目录":
            toc=p;break
    if toc is not None:
        toc.paragraph_format.page_break_before=True
    # Force second TOC page at Chapter 4.
    bi=body_index(d)
    ti=next((i for i,p in enumerate(d.paragraphs) if norm(p.text)=="目录"),None)
    if ti is not None:
        for p in d.paragraphs[ti+1:bi]:
            if norm(p.text.split("\t")[0]).startswith("4面向装配式住宅的证据分层评价"):
                p.paragraph_format.page_break_before=True
                break

def _has_page_break(el):
    return bool(el.xpath(".//w:br[@w:type='page']"))

def split_before_body(d):
    # Give正文 its own section. Remove page-break duplication first to avoid a blank page.
    bi=body_index(d); bp=d.paragraphs[bi]
    bp.paragraph_format.page_break_before=False
    prev=bp._element.getprevious()
    if prev is not None and prev.tag==qn("w:p") and _has_page_break(prev):
        pp=prev.find(qn("w:pPr"))
        has_sect=pp is not None and pp.find(qn("w:sectPr")) is not None
        if not has_sect:
            prev.getparent().remove(prev)

    final_sect=d._element.body.find(qn("w:sectPr"))
    if final_sect is None:
        raise RuntimeError("文档缺少末节节属性")
    pre=copy.deepcopy(final_sect)
    for tag in ("w:headerReference","w:footerReference","w:pgNumType"):
        for x in list(pre.findall(qn(tag))):
            pre.remove(x)
    typ=pre.find(qn("w:type"))
    if typ is None:
        typ=OxmlElement("w:type");pre.insert(0,typ)
    typ.set(qn("w:val"),"nextPage")
    carrier=OxmlElement("w:p");pPr=OxmlElement("w:pPr");pPr.append(pre);carrier.append(pPr)
    bp._element.addprevious(carrier)

def clear_header(h):
    root=h._element
    for x in list(root):root.remove(x)
    root.append(OxmlElement("w:p"))

def clear_footer(f):
    root=f._element
    for x in list(root):root.remove(x)
    root.append(OxmlElement("w:p"))

def add_text_run(p,text):
    r=OxmlElement("w:r");rp=OxmlElement("w:rPr")
    fonts=OxmlElement("w:rFonts")
    for a in ("w:ascii","w:hAnsi","w:eastAsia"): fonts.set(qn(a),"宋体")
    rp.append(fonts);rp.append(OxmlElement("w:b"))
    sz=OxmlElement("w:sz");sz.set(qn("w:val"),"21");rp.append(sz)
    sz2=OxmlElement("w:szCs");sz2.set(qn("w:val"),"21");rp.append(sz2)
    r.append(rp);t=OxmlElement("w:t");t.text=text;r.append(t);p.append(r)

def add_page_field(p):
    for typ in ("begin",):
        r=OxmlElement("w:r");fc=OxmlElement("w:fldChar");fc.set(qn("w:fldCharType"),typ);r.append(fc);p.append(r)
    r=OxmlElement("w:r");it=OxmlElement("w:instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space","preserve");it.text=" PAGE "
    r.append(it);p.append(r)
    r=OxmlElement("w:r");fc=OxmlElement("w:fldChar");fc.set(qn("w:fldCharType"),"separate");r.append(fc);p.append(r)
    add_text_run(p,"1")
    r=OxmlElement("w:r");fc=OxmlElement("w:fldChar");fc.set(qn("w:fldCharType"),"end");r.append(fc);p.append(r)

def build_header(h,total="00"):
    root=h._element
    for x in list(root):root.remove(x)
    p=OxmlElement("w:p");pPr=OxmlElement("w:pPr")
    jc=OxmlElement("w:jc");jc.set(qn("w:val"),"left");pPr.append(jc)
    tabs=OxmlElement("w:tabs");tab=OxmlElement("w:tab")
    tab.set(qn("w:val"),"right");tab.set(qn("w:pos"),"8760");tabs.append(tab);pPr.append(tabs)
    bdr=OxmlElement("w:pBdr");bottom=OxmlElement("w:bottom")
    for k,v in (("w:val","single"),("w:sz","6"),("w:space","1"),("w:color","auto")):bottom.set(qn(k),v)
    bdr.append(bottom);pPr.append(bdr)
    sp=OxmlElement("w:spacing");sp.set(qn("w:before"),"0");sp.set(qn("w:after"),"0");pPr.append(sp)
    ind=OxmlElement("w:ind");ind.set(qn("w:left"),"0");ind.set(qn("w:right"),"0");ind.set(qn("w:firstLine"),"0");pPr.append(ind)
    p.append(pPr)
    add_text_run(p,"炎黄职业技术学院毕业论文")
    r=OxmlElement("w:r");r.append(OxmlElement("w:tab"));p.append(r)
    add_text_run(p,"第 ");add_page_field(p);add_text_run(p,f" 页  共 {total} 页")
    root.append(p)

def set_body_pagination(d):
    if len(d.sections)<2:raise RuntimeError("文档分节异常")
    for s in d.sections[:-1]:
        s.header.is_linked_to_previous=False;clear_header(s.header)
        s.footer.is_linked_to_previous=False;clear_footer(s.footer)
        pg=s._sectPr.find(qn("w:pgNumType"))
        if pg is not None:s._sectPr.remove(pg)
    s=d.sections[-1]
    s.header.is_linked_to_previous=False;build_header(s.header)
    s.footer.is_linked_to_previous=False;clear_footer(s.footer)
    pg=s._sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg=OxmlElement("w:pgNumType");s._sectPr.append(pg)
    pg.set(qn("w:start"),"1")
    uf=d.settings._element.find(qn("w:updateFields"))
    if uf is None:
        uf=OxmlElement("w:updateFields");d.settings._element.append(uf)
    uf.set(qn("w:val"),"true")

def superscript_citations(d):
    pat=re.compile(r"\[(?:\d+)(?:\s*[,，、;；\-–]\s*\d+)*\]")
    start=body_index(d);end=ref_index(d)
    for p in d.paragraphs[start:end]:
        for run in list(p.runs):
            txt=run.text or ""
            if not pat.search(txt):continue
            rxml=run._r
            parent=rxml.getparent()
            pos=parent.index(rxml)
            rp=rxml.find(qn("w:rPr"))
            parts=[];last=0
            for m in pat.finditer(txt):
                if m.start()>last:parts.append((txt[last:m.start()],False))
                parts.append((m.group(),True));last=m.end()
            if last<len(txt):parts.append((txt[last:],False))
            parent.remove(rxml)
            for content,is_sup in reversed(parts):
                nr=OxmlElement("w:r")
                if rp is not None:nr.append(copy.deepcopy(rp))
                if is_sup:
                    nrp=nr.find(qn("w:rPr"))
                    if nrp is None:nrp=OxmlElement("w:rPr");nr.insert(0,nrp)
                    va=nrp.find(qn("w:vertAlign"))
                    if va is None:va=OxmlElement("w:vertAlign");nrp.append(va)
                    va.set(qn("w:val"),"superscript")
                t=OxmlElement("w:t");t.text=content
                if content.startswith(" ") or content.endswith(" "):
                    t.set("{http://www.w3.org/XML/1998/namespace}space","preserve")
                nr.append(t);parent.insert(pos,nr)

def main():
    if not DOCX.exists():raise FileNotFoundError(DOCX)
    d=Document(DOCX)
    remove_duplicate_standalone_conclusion(d)
    remove_template_residue(d)
    rebuild_extra_toc_entries(d)
    ensure_front_pagination(d)
    style_special_pages(d)
    superscript_citations(d)
    split_before_body(d)
    set_body_pagination(d)
    n=han_count(d)
    d.save(DOCX)
    print("FINALIZED",DOCX,DOCX.stat().st_size,"sections",len(d.sections),"han",n)
    if not 9000<=n<=11500:
        raise RuntimeError(f"正文汉字数不在目标区间: {n}")

if __name__=="__main__":
    main()
