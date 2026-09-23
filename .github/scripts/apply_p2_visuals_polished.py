# -*- coding: utf-8 -*-
from pathlib import Path
import shutil, re, copy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
PAPER=ROOT/"新建文件夹/新建文件夹/装配式施工质量管理问题及优化研究——以市政项目为例"
SRC=PAPER/"装配式施工质量管理问题及优化研究——以市政项目为例_最终优化版.docx"
OUT=PAPER/"装配式施工质量管理问题及优化研究——以市政项目为例_图表优化终版.docx"
FIG=ROOT/"新建文件夹/新建文件夹/.workbuddy-ai/work/figs_p2_polished"

def norm(s):
    return re.sub(r"\s+","",s).replace("–","-").replace("—","-")

def font_run(run, size=10.5, bold=False, color=None, font="宋体"):
    run.font.name=font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),font)
    run.font.size=Pt(size);run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)

def style_caption(p):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.keep_with_next=False
    for r in p.runs:
        if r.text: font_run(r,10.5,False,font="宋体")

def replace_picture(doc, caption_prefix, img):
    cap=None
    target=norm(caption_prefix)
    for p in doc.paragraphs:
        if norm(p.text).startswith(target):
            cap=p;break
    if cap is None: raise RuntimeError("找不到图题 "+caption_prefix)
    prev=cap._p.getprevious()
    # 删除紧邻图题的旧图片段落
    if prev is not None and prev.tag==qn("w:p") and prev.xpath(".//w:drawing"):
        prev.getparent().remove(prev)
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(img),width=Cm(13.0))
    cap._p.addprevious(p._p)
    style_caption(cap)

def set_cell_shading(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn("w:shd"))
    if shd is None:
        shd=OxmlElement("w:shd");tcPr.append(shd)
    shd.set(qn("w:fill"),fill);shd.set(qn("w:val"),"clear")

def set_cell_margins(cell,top=60,start=80,bottom=60,end=80):
    tc=cell._tc;tcPr=tc.get_or_add_tcPr()
    tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar=OxmlElement("w:tcMar");tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn("w:"+m))
        if node is None:
            node=OxmlElement("w:"+m);tcMar.append(node)
        node.set(qn("w:w"),str(v));node.set(qn("w:type"),"dxa")

def set_table_borders(table,color="A8B3BF"):
    tblPr=table._tbl.tblPr
    borders=tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders=OxmlElement("w:tblBorders");tblPr.append(borders)
    for edge,sz in (("top","8"),("bottom","8"),("left","4"),("right","4"),("insideH","4"),("insideV","4")):
        el=borders.find(qn("w:"+edge))
        if el is None:
            el=OxmlElement("w:"+edge);borders.append(el)
        el.set(qn("w:val"),"single");el.set(qn("w:sz"),sz);el.set(qn("w:space"),"0");el.set(qn("w:color"),color)

def no_split(row):
    trPr=row._tr.get_or_add_trPr()
    cant=OxmlElement("w:cantSplit");trPr.append(cant)

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr()
    hdr=OxmlElement("w:tblHeader");hdr.set(qn("w:val"),"true");trPr.append(hdr)

def style_body_table(table, widths=None):
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    set_table_borders(table)
    repeat_header(table.rows[0])
    for ri,row in enumerate(table.rows):
        no_split(row)
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri==0:set_cell_shading(cell,"DCE6F1")
            elif ri%2==0:set_cell_shading(cell,"F7F9FB")
            else:set_cell_shading(cell,"FFFFFF")
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.0
                if ri==0 or ci==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                else:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    if r.text:
                        font_run(r,9.5,ri==0 or ci==0,font="宋体")
        if widths:
            for ci,w in enumerate(widths):
                for cell in row.cells[ci:ci+1]:
                    cell.width=Cm(w)

def new_caption(doc,text,before_el):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True;p.paragraph_format.first_line_indent=Pt(0)
    r=p.add_run(text);font_run(r,10.5,False,font="宋体")
    before_el.addprevious(p._p)
    return p

def add_table_before(doc,heading_text,title,headers,rows,widths):
    # 防止重复
    if any(norm(p.text)==norm(title) for p in doc.paragraphs):
        return
    anchor=next((p for p in doc.paragraphs if norm(p.text)==norm(heading_text)),None)
    if anchor is None: raise RuntimeError("找不到插入锚点 "+heading_text)
    tbl=doc.add_table(rows=1,cols=len(headers))
    for i,h in enumerate(headers):tbl.rows[0].cells[i].text=h
    for data in rows:
        cells=tbl.add_row().cells
        for i,v in enumerate(data):cells[i].text=v
    style_body_table(tbl,widths)
    anchor._p.addprevious(tbl._tbl)
    new_caption(doc,title,tbl._tbl)
    return tbl

def preserve_text_sequence(src,out):
    a=[p.text for p in Document(src).paragraphs]
    b=[p.text for p in Document(out).paragraphs if not norm(p.text).startswith(("表3-1","表5-1","表6-1"))]
    # 新增表题之外，正文段落文本必须原样存在且顺序一致
    it=iter(b)
    for x in a:
        for y in it:
            if y==x:break
        else:
            raise RuntimeError("段落文本顺序校验失败: "+x[:50])

def main():
    if not SRC.exists():raise FileNotFoundError(SRC)
    shutil.copy2(SRC,OUT)
    d=Document(OUT)
    for prefix,name in [
        ("图1-1","fig1_1.png"),("图2-1","fig2_1.png"),("图3-1","fig3_1.png"),
        ("图4-1","fig4_1.png"),("图5-1","fig5_1.png")]:
        replace_picture(d,prefix,FIG/name)

    # 仅新增/优化正文表格；开题、任务书、中期检查、封面表格严格保留学校模板原貌。
    add_table_before(d,"4  质量缺陷的传播机制分析","表3-1  案例工程关键风险与前置控制点",
        ["环节","主要风险","早期信号","前置控制点"],
        [
        ["工厂端","尺寸偏差、预埋件偏位、批次性缺陷","首件复测波动、试拼不顺畅","首件确认、模具/工装复核、异常批次加严"],
        ["物流端","运输碰损、连接面与止水构造受损","棱角破损、表面裂纹、保护缺失","装车清单、运输保护、到场状态对照"],
        ["拼装端","轴线/标高偏差连续累积","接缝宽度趋势异常、收口空间缩小","统一基准、连续测量、趋势预警"],
        ["接口端","灌浆/连接内部缺陷、接缝防水失控","材料或过程记录异常、封闭前状态不确定","隐蔽前确认、必要检测、复验后封闭"],
        ["信息端","编码不一致、资料与实体无法对应","异常历史查不到、责任边界模糊","统一编码、证据包、问题闭环"],
        ],[2.0,3.8,3.8,4.0])

    add_table_before(d,"6  优化方案的实施评价与研究边界","表5-1  G0—G5质量门检查要点与放行条件",
        ["质量门","核心检查点","必备证据","放行条件"],
        [
        ["G0 设计冻结门","图纸版本、构件编号、接口条件","冻结清单、会审/确认记录","生产依赖信息一致且无关键待定项"],
        ["G1 工厂放行门","首件、尺寸、预埋件、连接面、标识","首件记录、检验记录、出厂资料","构件达到可交付状态"],
        ["G2 进场验收门","身份对应、运输损伤、保护状态","到场验收单、关键影像","运输后状态未发生影响安装的变化"],
        ["G3 拼装几何门","轴线、标高、接缝、累计趋势","连续测量与复测记录","当前偏差与累计趋势均受控"],
        ["G4 隐蔽接口门","连接、防水、基层、过程状态","隐蔽验收、旁站/检测记录","证据充分且异常已复验关闭"],
        ["G5 数据闭环门","实体位置、问题整改、资料归档","闭环单、归档清单、构件身份链","实体与质量信息均形成闭环"],
        ],[2.3,3.7,3.6,4.0])

    add_table_before(d,"6．2  实施顺序：先做一个标准段，再扩展到全线","表6-1  优化方案实施评价指标",
        ["评价维度","建议指标","指标含义","数据来源"],
        [
        ["前序拦截","前序关闭问题占比","问题是否在进入下一阶段前被发现并关闭","质量门/问题台账"],
        ["几何稳定","二次调整构件数、偏差趋势","G3对累计偏差的控制程度","连续测量记录"],
        ["接口质量","隐蔽前整改与复验情况","连接、防水问题是否在封闭前处理","隐蔽验收/检测记录"],
        ["闭环效率","问题关闭时长、逾期数","异常从登记到完成复验的处理效率","问题闭环台账"],
        ["追溯完整","随机构件资料可调取率","设计—生产—交付—安装—整改记录是否贯通","构件档案/编码系统"],
        ],[2.1,3.3,5.0,3.2])

    d.save(OUT)
    preserve_text_sequence(SRC,OUT)
    d2=Document(OUT)
    titles=[p.text.strip() for p in d2.paragraphs if norm(p.text).startswith(("表3-1","表5-1","表6-1"))]
    print("POLISH_OK","pics",len(d2.inline_shapes),"tables",len(d2.tables),"titles",titles,"size",OUT.stat().st_size)

if __name__=="__main__":main()
