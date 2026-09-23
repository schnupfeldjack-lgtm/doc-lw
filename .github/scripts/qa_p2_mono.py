# -*- coding: utf-8 -*-
from pathlib import Path
import re,subprocess,sys,zipfile
from docx import Document
from docx.oxml.ns import qn
import pdfplumber
from PIL import Image
import io

def norm(s):return re.sub(r"\s+","",s).replace("–","-").replace("—","-")
def ptext(pdf,i):return subprocess.check_output(["pdftotext","-f",str(i),"-l",str(i),"-layout",str(pdf),"-"],text=True,errors="ignore")
def pages(pdf):
    s=subprocess.check_output(["pdfinfo",str(pdf)],text=True,errors="ignore")
    return int(re.search(r"^Pages:\s+(\d+)",s,re.M).group(1))

def main():
    docx=Path(sys.argv[1]);pdf=Path(sys.argv[2]);d=Document(docx)
    assert len(d.inline_shapes)>=5
    assert len(d.tables)>=8
    refs=[p.text.strip() for p in d.paragraphs if re.match(r"^\[\d+\]",p.text.strip())]
    assert len(refs)==15,len(refs)
    # 所有正文新增表格禁止底色和有色边框
    for t in d.tables[5:]:
        xml=t._tbl.xml.lower()
        assert 'w:fill="dce6f1"' not in xml and 'w:fill="f7f9fb"' not in xml
        for bad in ("5b9bd5","70ad47","ed7d31","8064a2","2f8f9d","c94c4c"):
            assert bad not in xml
    # 图片二进制真实嵌入，且5张正文技术图必须是纯黑白/灰度，不允许出现彩色像素。
    with zipfile.ZipFile(docx) as z:
        pngs=[x for x in z.namelist() if x.startswith("word/media/") and x.endswith(".png")]
        blobs=[(x,z.read(x)) for x in pngs]
        big=sorted(blobs,key=lambda kv:len(kv[1]),reverse=True)[:5]
        assert len(big)==5 and all(len(b)>100000 for _,b in big),[(n,len(b)) for n,b in big]
        for n,b in big:
            im=Image.open(io.BytesIO(b)).convert("RGB")
            # 缩小后抽样，检查RGB三通道必须近似相等，即真正灰度。
            im.thumbnail((500,500))
            pix=list(im.getdata())
            colored=sum(1 for rr,gg,bb in pix if max(rr,gg,bb)-min(rr,gg,bb)>3)
            assert colored<=max(5,len(pix)//10000),(n,colored,len(pix))
    total=pages(pdf);texts=[ptext(pdf,i) for i in range(1,total+1)];np=[norm(x) for x in texts]
    bphys=next(i for i,t in enumerate(np,1) if "1引言" in t and "市政综合管廊具有线路长" in t)
    assert bphys==9,bphys
    for i in range(1,bphys):
        assert not re.search(r"第\s*\d+\s*页\s*共\s*\d+\s*页",texts[i-1]),i
    body_pages=total-bphys+1
    assert f"第1页共{body_pages}页" in np[bphys-1]
    assert f"第{body_pages}页共{body_pages}页" in np[-1]
    for i,t in enumerate(np,1):assert len(t)>20,f"blank {i}"
    # Word结构级检查：所有5张正文图必须是“嵌入型(inline)”而不是浮动锚点，
    # 图片段落必须居中，宽度不超过12.6cm。这样从结构上杜绝图片横向/纵向漂移。
    for label in ("图1-1","图2-1","图3-1","图4-1","图5-1"):
        cap=next((p for p in d.paragraphs if norm(p.text).startswith(norm(label))),None)
        assert cap is not None,label
        prev=cap._p.getprevious()
        assert prev is not None and prev.tag==qn("w:p"),label
        assert prev.xpath(".//wp:inline"),(label,"not-inline")
        assert not prev.xpath(".//wp:anchor"),(label,"floating-anchor")
        jc=prev.find(qn("w:pPr"))
        jc=jc.find(qn("w:jc")) if jc is not None else None
        assert jc is not None and jc.get(qn("w:val"))=="center",(label,"not-centered")
        ext=prev.xpath(".//wp:inline/wp:extent")
        assert ext,(label,"no-extent")
        cx=int(ext[0].get("cx"))
        assert cx<=4536000,(label,cx)  # 12.6cm
    print("MONO_QA_OK","pages",total,"body_start",bphys,"body_pages",body_pages,"pics",len(d.inline_shapes),"tables",len(d.tables),"refs",len(refs))

if __name__=="__main__":main()
