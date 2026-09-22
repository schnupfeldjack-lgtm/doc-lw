# -*- coding: utf-8 -*-
from pathlib import Path
import re, subprocess, sys
from docx import Document

def norm(s):
    return re.sub(r"\s+","",s).replace("–","-").replace("—","-")

def pdf_pages(pdf):
    info=subprocess.check_output(["pdfinfo",str(pdf)],text=True,errors="ignore")
    n=int(re.search(r"^Pages:\s+(\d+)",info,re.M).group(1))
    pages=[]
    for i in range(1,n+1):
        t=subprocess.check_output(["pdftotext","-f",str(i),"-l",str(i),"-layout",str(pdf),"-"],text=True,errors="ignore")
        pages.append(t)
    return pages

def level(t):
    n=norm(t)
    if re.match(r"^\d+\.\d+\.\d+",n):return 3
    if re.match(r"^\d+[．.]\d+",n):return 2
    if re.match(r"^\d+",n):return 1
    return None

def body_idx(doc):
    xs=[i for i,p in enumerate(doc.paragraphs) if norm(p.text)=="1绪论" and not (p.style and p.style.name.startswith("toc"))]
    return xs[-1]

def replace_page(p,num):
    nodes=p._p.xpath(".//w:t")
    for n in reversed(nodes):
        txt=n.text or ""
        if re.search(r"\d+\s*$",txt):
            n.text=re.sub(r"\d+\s*$",str(num),txt);return True
    # fallback: modify last run
    if p.runs:
        p.runs[-1].text=(p.runs[-1].text or "")+str(num);return True
    return False

def main():
    docx=Path(sys.argv[1]);pdf=Path(sys.argv[2])
    d=Document(docx);pages=pdf_pages(pdf);np=[norm(x) for x in pages]
    bphys=None
    for i,t in enumerate(np,1):
        if "第1页" in t and "1绪论" in t:
            bphys=i;break
    if bphys is None:raise RuntimeError("无法识别正文物理起始页")
    bi=body_idx(d);ri=next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="参考文献")
    heads=[]
    for p in d.paragraphs[bi:ri]:
        lv=level(p.text)
        if lv in (1,2):
            heads.append((norm(p.text),p.text.strip()))
    mapping={}
    for key,raw in heads:
        for phys in range(bphys,len(pages)+1):
            if key in np[phys-1]:
                mapping[key]=phys-bphys+1;break
        if key not in mapping:print("WARN heading page not found",raw)
    ti=next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=="目录")
    for p in d.paragraphs[ti+1:bi]:
        tn=norm(p.text)
        matches=[(k,v) for k,v in mapping.items() if tn.startswith(k)]
        if matches:
            k,v=max(matches,key=lambda x:len(x[0]))
            replace_page(p,v)
    d.save(docx)
    print("TOC_UPDATED body_physical",bphys,"mapping",mapping)

if __name__=="__main__":main()
